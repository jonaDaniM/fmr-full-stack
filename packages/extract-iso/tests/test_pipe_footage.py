import base64
import csv
import json
from decimal import Decimal
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree

import fitz
import pytest
from docx import Document

from iso_bom.pipe_cli import main as pipe_cli_main
from iso_bom.pipe_parser import parse_linear_feet, parse_pipe_iso_page
from iso_bom.pipe_pipeline import run_pipe_footage


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SUPPLIED_PACKAGE = PROJECT_ROOT / "templates" / "CreateFMR" / "IP-SMM30R127MMPP-K447.pdf"
PIPE_FOOTAGE_2_PACKAGE = (
    PROJECT_ROOT / "input" / "getPipe2" / "IP-SMM30R117MMPP-K447 combined.pdf"
)
PIPE_FOOTAGE_3_PACKAGE = (
    PROJECT_ROOT / "input" / "getPipe3" / "IP-SMM30R107MMPP-K447 combined.pdf"
)

ONE_PIXEL_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def _insert_cover(document: fitz.Document, iwp_number: str) -> None:
    page = document.new_page(width=612, height=792)
    page.insert_text((170, 35), "Installation Work Package - TEST", fontsize=12)
    page.insert_text((50, 75), "IWP Number:", fontsize=10)
    page.insert_text((130, 75), iwp_number, fontsize=10)


def _insert_iso(document: fitz.Document, drawing_number: str, revision: str, rows) -> None:
    page = document.new_page(width=900, height=700)
    page.insert_text((420, 35), "BILL OF MATERIALS", fontsize=9)
    for x, text in [
        (420, "PT"), (455, "DESCRIPTION"), (650, "NPD"),
        (710, "CMDTY"), (830, "QTY"),
    ]:
        page.insert_text((x, 60), text, fontsize=8)
    y = 90
    for row in rows:
        for x, text in zip((420, 455, 650, 710, 830), row):
            if text:
                page.insert_text((x, y), text, fontsize=8)
        y += 28
    page.insert_text((500, 610), "ISOMETRIC DRAWING NUMBER", fontsize=8)
    page.insert_text((790, 610), "REV", fontsize=8)
    page.insert_text((500, 635), drawing_number, fontsize=8)
    page.insert_text((795, 635), revision, fontsize=8)


def _insert_weld_log(document: fitz.Document) -> None:
    page = document.new_page(width=612, height=792)
    page.insert_text((50, 50), "SECTION 15117", fontsize=10)
    page.insert_text((50, 70), "FABRICATION OF METALLIC PIPE AND TUBING", fontsize=10)
    page.insert_text((50, 90), "ATTACHMENT A - WELD/BRAZE LOG SHEET", fontsize=10)


def _insert_pipe_support(document: fitz.Document) -> None:
    page = document.new_page(width=612, height=792)
    page.insert_text(
        (50, 50),
        "PIPE HANGERS AND SUPPORTS - ATTACHMENT C-1",
        fontsize=10,
    )


def _insert_image_attachment(document: fitz.Document) -> None:
    page = document.new_page(width=500, height=300)
    page.insert_image(page.rect, stream=ONE_PIXEL_PNG)


def _create_complete_package(path: Path) -> None:
    document = fitz.open()
    _insert_cover(document, "IWP-PIPE-100")
    _insert_iso(document, "ISO-001", "0", [
        ("1", "PIPE SCH 10S", "2", "P-001", "99.9'"),
    ])
    current_rows = [
        ("1", "PIPE SCH 10S", "2", "P-001", "10.0'"),
        ("2", "PIPE SCH 10S", "1", "P-002", "2.5'"),
        ("3", "PIPET SCH 10S 316 SS", "2X1", "FIT-001", "1"),
        ("4", "ELL 90 LR", "2", "FIT-002", "2"),
    ]
    _insert_iso(document, "ISO-001", "2", current_rows)
    _insert_iso(document, "ISO-001", "2", current_rows)
    _insert_iso(document, "ISO-002", "0", [
        ("1", "ELL 45 LR", "4", "FIT-003", "1"),
    ])
    _insert_weld_log(document)
    document.save(path)
    document.close()


def _document_text(path: Path) -> str:
    document = Document(path)
    values = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        values.extend(cell.text for row in table.rows for cell in row.cells)
    return "\n".join(values)


def test_parse_linear_feet_requires_foot_mark_and_preserves_decimal():
    assert parse_linear_feet("7'") == Decimal("7")
    assert parse_linear_feet("21.60'") == Decimal("21.60")
    assert parse_linear_feet("3.25\u2032") == Decimal("3.25")
    for invalid in ("", "7", "7 ft", "1' 6\""):
        with pytest.raises(ValueError):
            parse_linear_feet(invalid)


def test_pipe_page_reads_all_pipe_rows_and_excludes_pipet():
    document = fitz.open()
    _insert_iso(document, "ISO-MULTI", "1", [
        ("1", "PIPE SCH 10S", "2", "P-001", "0.1'"),
        ("2", "PIPE SCH 40", "4", "P-002", "0.2'"),
        ("3", "PIPET SCH 10S", "4X2", "F-001", "1"),
    ])
    parsed = parse_pipe_iso_page(document[0], 1, "package.pdf", "package.pdf")
    document.close()
    assert [item.row.point_number for item in parsed.pipe_measurements] == ["1", "2"]
    assert parsed.total_linear_feet == Decimal("0.3")
    assert parsed.precision == 1
    assert parsed.review_reasons == []


def test_complete_pipeline_creates_report_audit_and_exact_total(tmp_path):
    source = tmp_path / "input"
    output = tmp_path / "output"
    source.mkdir()
    _create_complete_package(source / "package.pdf")

    summary = run_pipe_footage(source, output)
    assert summary["status"] == "complete"
    assert summary["iwp_number"] == "IWP-PIPE-100"
    assert summary["iso_pages_detected"] == 4
    assert summary["unique_drawings"] == 2
    assert summary["selected_iso_drawings"] == 2
    assert summary["superseded_pages_ignored"] == 1
    assert summary["duplicate_pages_ignored"] == 1
    assert summary["pipe_bom_rows"] == 2
    assert summary["package_total_linear_feet"] == "12.5"
    assert summary["quarantine_entries"] == 0
    assert summary["estimated_time_saved_seconds"] == "50.66"
    assert summary["time_estimate_available"] is True
    assert float(summary["processing_seconds"]) > 0

    manifests = list((output / "analytics_runs").glob("*.json"))
    assert len(manifests) == 1
    manifest = json.loads(manifests[0].read_text())
    assert manifest["workflow"] == "pipe_footage"
    assert manifest["status"] == "complete"
    assert manifest["package_number"] == "IWP-PIPE-100"
    assert [item["drawing_number"] for item in manifest["isos"]] == [
        "ISO-001", "ISO-002",
    ]
    assert manifest["metrics"]["package_total_linear_feet"] == "12.5"
    assert manifest["metrics"]["estimated_time_saved_seconds"] == "50.66"
    assert [item["estimated_time_saved_seconds"] for item in manifest["isos"]] == [
        "8.00", "8.00",
    ]

    report = Path(summary["report"])
    assert report.name == "IWP-PIPE-100_PIPE_FOOTAGE.docx"
    text = _document_text(report)
    assert "ISO-001" in text and "12.5 LF" in text
    assert "ISO-002" in text and "0.0 LF" in text
    assert "PACKAGE TOTAL" in text

    document = Document(report)
    section = document.sections[0]
    assert section.page_width == 7772400
    assert section.page_height == 10058400
    assert section.left_margin == section.right_margin == 914400
    table_xml = document.tables[0]._tbl
    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    assert table_xml.find("w:tblPr/w:tblW", namespace).get(f"{{{namespace['w']}}}w") == "9360"
    assert [
        node.get(f"{{{namespace['w']}}}w")
        for node in table_xml.findall("w:tblGrid/w:gridCol", namespace)
    ] == ["6480", "960", "1920"]
    with ZipFile(report) as archive:
        styles = ElementTree.fromstring(archive.read("word/styles.xml"))
    title_style = styles.find("w:style[@w:styleId='Title']", namespace)
    subtitle_style = styles.find("w:style[@w:styleId='Subtitle']", namespace)
    assert title_style.find("w:pPr/w:pBdr", namespace) is None
    assert subtitle_style.find("w:pPr/w:numPr", namespace) is None

    with (output / "pipe_footage_audit.csv").open() as handle:
        audit = list(csv.DictReader(handle))
    assert [row["point_number"] for row in audit] == ["1", "2"]
    assert [row["linear_feet"] for row in audit] == ["10.0", "2.5"]
    assert all(row["status"] == "accepted" for row in audit)
    with (output / "pipe_footage_quarantine.csv").open() as handle:
        assert list(csv.DictReader(handle)) == []

    with pytest.raises(FileExistsError, match="already exists"):
        run_pipe_footage(source, output)
    overwritten = run_pipe_footage(source, output, overwrite=True)
    assert overwritten["status"] == "complete"


def test_review_required_report_withholds_total(tmp_path):
    source = tmp_path / "input"
    output = tmp_path / "output"
    source.mkdir()
    document = fitz.open()
    _insert_cover(document, "IWP-PIPE-REVIEW")
    _insert_iso(document, "ISO-BAD", "0", [
        ("1", "PIPE SCH 10S", "2", "P-001", "8"),
    ])
    document.new_page(width=612, height=792)
    document.save(source / "package.pdf")
    document.close()

    summary = run_pipe_footage(source, output)
    assert summary["status"] == "review_required"
    assert summary["package_total_linear_feet"] is None
    assert summary["quarantine_entries"] == 2
    assert summary["time_estimate_available"] is False
    assert summary["estimated_time_saved_seconds"] is None
    manifest_path = next((output / "analytics_runs").glob("*.json"))
    manifest = json.loads(manifest_path.read_text())
    assert manifest["status"] == "review_required"
    assert manifest["metrics"]["package_total_linear_feet"] is None
    assert len(manifest["issues"]) == 2
    text = _document_text(Path(summary["report"]))
    assert "REVIEW REQUIRED" in text
    assert "Package Total: WITHHELD" in text
    assert "invalid pipe quantity" in text
    report = Document(summary["report"])
    assert not any(
        cell.text == "PACKAGE TOTAL"
        for table in report.tables
        for row in table.rows
        for cell in row.cells
    )


def test_iwp_override_and_conflict_exit_codes(tmp_path, capsys):
    source = tmp_path / "input"
    source.mkdir()
    document = fitz.open()
    _insert_iso(document, "ISO-001", "0", [
        ("1", "PIPE SCH 10S", "2", "P-001", "5.0'"),
    ])
    document.save(source / "iso.pdf")
    document.close()

    complete_output = tmp_path / "complete"
    assert pipe_cli_main([
        "--input", str(source), "--output", str(complete_output),
        "--iwp-number", "IWP-MANUAL-1",
    ]) == 0
    capsys.readouterr()

    conflict_source = tmp_path / "conflict-input"
    conflict_source.mkdir()
    document = fitz.open()
    _insert_cover(document, "IWP-DETECTED")
    _insert_iso(document, "ISO-001", "0", [
        ("1", "PIPE SCH 10S", "2", "P-001", "5.0'"),
    ])
    document.save(conflict_source / "package.pdf")
    document.close()
    conflict_output = tmp_path / "conflict-output"
    assert pipe_cli_main([
        "--input", str(conflict_source), "--output", str(conflict_output),
        "--iwp-number", "IWP-OTHER",
    ]) == 2
    capsys.readouterr()
    assert not list(conflict_output.glob("*.docx"))


def test_missing_iwp_adds_one_package_quarantine_entry(tmp_path):
    source = tmp_path / "input"
    output = tmp_path / "output"
    source.mkdir()
    document = fitz.open()
    _insert_iso(document, "ISO-001", "0", [
        ("1", "PIPE SCH 10S", "2", "P-001", "5.0'"),
    ])
    document.save(source / "iso.pdf")
    document.close()

    summary = run_pipe_footage(source, output)
    assert summary["status"] == "blocked"
    assert summary["quarantine_entries"] == 1
    assert summary["time_estimate_available"] is False
    manifest_path = next((output / "analytics_runs").glob("*.json"))
    manifest = json.loads(manifest_path.read_text())
    assert manifest["status"] == "blocked"
    assert manifest["package_number"] == ""
    with (output / "pipe_footage_quarantine.csv").open() as handle:
        entries = list(csv.DictReader(handle))
    assert [entry["reason_code"] for entry in entries] == ["iwp_missing"]


def test_consecutive_image_support_views_use_bracketing_context(tmp_path):
    source = tmp_path / "input"
    output = tmp_path / "output"
    source.mkdir()
    document = fitz.open()
    _insert_cover(document, "IWP-IMAGE-RUN")
    _insert_pipe_support(document)
    _insert_image_attachment(document)
    _insert_image_attachment(document)
    _insert_pipe_support(document)
    _insert_iso(document, "ISO-001", "0", [
        ("1", "PIPE SCH 10S", "2", "P-001", "5.0'"),
    ])
    document.save(source / "package.pdf")
    document.close()

    summary = run_pipe_footage(source, output)
    assert summary["status"] == "complete"
    assert summary["pipe_support_pages_ignored"] == 2
    assert summary["attachment_image_pages_ignored"] == 2
    assert summary["quarantine_entries"] == 0


def test_readable_non_iso_page_is_ignored(tmp_path):
    source = tmp_path / "input"
    output = tmp_path / "output"
    source.mkdir()
    document = fitz.open()
    _insert_cover(document, "IWP-NON-ISO")
    notes = document.new_page(width=612, height=792)
    notes.insert_text(
        (50, 50),
        "Construction schedule, routing notes, and general package information only.",
        fontsize=10,
    )
    _insert_iso(document, "ISO-001", "0", [
        ("1", "PIPE SCH 10S", "2", "P-001", "5.0'"),
    ])
    document.save(source / "package.pdf")
    document.close()

    summary = run_pipe_footage(source, output)
    assert summary["status"] == "complete"
    assert summary["other_non_iso_pages_ignored"] == 1
    assert summary["quarantine_entries"] == 0


@pytest.mark.skipif(
    not SUPPLIED_PACKAGE.exists(),
    reason="Local supplied IWP package is unavailable",
)
def test_supplied_package_pipe_footage_regression(tmp_path):
    summary = run_pipe_footage(SUPPLIED_PACKAGE.parent, tmp_path / "output")
    assert summary["status"] == "complete"
    assert summary["pages_scanned"] == 85
    assert summary["iso_pages_detected"] == 51
    assert summary["selected_iso_drawings"] == 49
    assert summary["superseded_pages_ignored"] == 2
    assert summary["package_total_linear_feet"] == "852.6"


@pytest.mark.skipif(
    not PIPE_FOOTAGE_2_PACKAGE.exists(),
    reason="Local pipeFootage2 package is unavailable",
)
def test_pipe_footage_2_mixed_package_regression(tmp_path):
    summary = run_pipe_footage(
        PIPE_FOOTAGE_2_PACKAGE.parent,
        tmp_path / "output",
    )
    assert summary["status"] == "complete"
    assert summary["iwp_number"] == "IP-SMM30R117MMPP-K447"
    assert summary["pages_scanned"] == 34
    assert summary["cover_pages_detected"] == 1
    assert summary["iso_pages_detected"] == 16
    assert summary["selected_iso_drawings"] == 8
    assert summary["duplicate_pages_ignored"] == 8
    assert summary["weld_log_pages_ignored"] == 8
    assert summary["pipe_support_pages_ignored"] == 5
    assert summary["attachment_image_pages_ignored"] == 2
    assert summary["fmr_attachment_pages_ignored"] == 1
    assert summary["logistics_pages_ignored"] == 1
    assert summary["package_total_linear_feet"] == "182.3"
    assert summary["quarantine_entries"] == 0


@pytest.mark.skipif(
    not PIPE_FOOTAGE_3_PACKAGE.exists(),
    reason="Local pipeFootage3 package is unavailable",
)
def test_pipe_footage_3_consecutive_support_views_regression(tmp_path):
    summary = run_pipe_footage(
        PIPE_FOOTAGE_3_PACKAGE.parent,
        tmp_path / "output",
    )
    assert summary["status"] == "complete"
    assert summary["iwp_number"] == "IP-SMM30R107MMPP-K447"
    assert summary["pages_scanned"] == 53
    assert summary["iso_pages_detected"] == 26
    assert summary["selected_iso_drawings"] == 26
    assert summary["pipe_bom_rows"] == 29
    assert summary["attachment_image_pages_ignored"] == 2
    assert summary["package_total_linear_feet"] == "456.3"
    assert summary["quarantine_entries"] == 0
    report = Document(summary["report"])
    assert len(report.tables) == 2
    assert [len(table.rows) for table in report.tables] == [19, 10]
