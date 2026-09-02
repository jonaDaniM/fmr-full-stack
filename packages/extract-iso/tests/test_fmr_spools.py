import csv
import json
from pathlib import Path

import fitz
import pytest
from openpyxl import load_workbook

from iso_bom.fmr_cli import build_parser, main as fmr_main
from iso_bom.fmr_model import IsoPage
from iso_bom.fmr_pipeline import run_fmr
from iso_bom.fmr_workbook import MATERIAL_CAPACITY, MATERIAL_LAST_ROW, build_fmr_workbook
from iso_bom.model import BomRow
from iso_bom.spool_parser import extract_spool_markers


PROJECT_ROOT = Path(__file__).resolve().parents[1]
FMR_TEMPLATE = PROJECT_ROOT / "templates" / "FMR" / "blankFMR.xlsx"


def _insert_cover(document: fitz.Document, iwp_number: str = "IWP-SPOOL-100") -> None:
    page = document.new_page(width=612, height=792)
    page.insert_text((170, 35), "Installation Work Package - TEST", fontsize=12)
    page.insert_text((50, 75), "IWP Number:", fontsize=10)
    page.insert_text((130, 75), iwp_number, fontsize=10)


def _insert_iso(
    document: fitz.Document,
    drawing_number: str,
    revision: str,
    rows=(("1", "PIPE SCH 10S", "2", "PIPE-001", "10"),
          ("2", "ELL 90 LR", "2", "ELL-001", "1")),
) -> fitz.Page:
    page = document.new_page(width=900, height=700)
    page.insert_text((420, 35), "BILL OF MATERIALS", fontsize=9)
    for x, text in [
        (420, "PT"), (455, "DESCRIPTION"), (650, "NPD"),
        (710, "CMDTY"), (830, "QTY"),
    ]:
        page.insert_text((x, 60), text, fontsize=8)
    y = 90
    for row in rows:
        for x, text in zip((420, 455, 650, 710, 830), row):
            page.insert_text((x, y), text, fontsize=8)
        y += 28
    page.insert_text((500, 610), "ISOMETRIC DRAWING NUMBER", fontsize=8)
    page.insert_text((790, 610), "REV", fontsize=8)
    page.insert_text((500, 635), drawing_number, fontsize=8)
    page.insert_text((795, 635), revision, fontsize=8)
    return page


def _draw_red_spool(
    page: fitz.Page,
    value: str,
    rect: fitz.Rect,
    *,
    split: bool = False,
    separate_segments: bool = False,
) -> None:
    if separate_segments:
        page.draw_line(rect.top_left, rect.top_right, color=(1, 0, 0), width=1)
        page.draw_line(rect.top_right, rect.bottom_right, color=(1, 0, 0), width=1)
        page.draw_line(rect.bottom_right, rect.bottom_left, color=(1, 0, 0), width=1)
        page.draw_line(rect.bottom_left, rect.top_left, color=(1, 0, 0), width=1)
    else:
        page.draw_rect(rect, color=(1, 0, 0), width=1)
    baseline = rect.y0 + 20
    if split:
        middle = len(value) // 2
        page.insert_text((rect.x0 + 6, baseline), value[:middle], fontsize=12, color=(1, 0, 0))
        page.insert_text((rect.x0 + 6 + middle * 7, baseline), value[middle:], fontsize=12, color=(1, 0, 0))
    else:
        page.insert_text((rect.x0 + 6, baseline), value, fontsize=12, color=(1, 0, 0))


def _draw_blue_spool(page: fitz.Page, value: str, rect: fitz.Rect, *, split=False) -> None:
    page.draw_oval(rect, color=(0, 0, 0.63), width=1)
    baseline = (rect.y0 + rect.y1) / 2 + 5
    if split:
        middle = len(value) // 2
        page.insert_text((rect.x0 + 8, baseline), value[:middle], fontsize=10)
        page.insert_text((rect.x0 + 8 + middle * 5.5, baseline), value[middle:], fontsize=10)
    else:
        page.insert_text((rect.x0 + 8, baseline), value, fontsize=10)


def _save_package(path: Path, pages) -> None:
    document = fitz.open()
    _insert_cover(document)
    for drawing_number, revision, spool_values in pages:
        page = _insert_iso(document, drawing_number, revision)
        for index, value in enumerate(spool_values):
            _draw_red_spool(
                page, value, fitz.Rect(50 + index * 130, 220, 155 + index * 130, 254)
            )
    document.save(path)
    document.close()


def _quarantine_rows(output: Path):
    with (output / "fmr_quarantine.csv").open() as handle:
        return list(csv.DictReader(handle))


def test_spool_detector_supports_reference_shapes_and_review_candidates():
    document = fitz.open()
    page = document.new_page(width=900, height=700)
    _draw_red_spool(page, "00001", fitz.Rect(50, 90, 155, 124))
    _draw_red_spool(
        page, "123456", fitz.Rect(180, 90, 285, 124), split=True
    )
    _draw_red_spool(
        page, "42", fitz.Rect(310, 90, 415, 124), separate_segments=True
    )
    _draw_blue_spool(page, "S00166", fitz.Rect(50, 160, 155, 194))
    _draw_blue_spool(
        page, "A-17/B", fitz.Rect(180, 160, 285, 194), split=True
    )
    page.insert_text((50, 245), "777", fontsize=12, color=(1, 0, 0))
    page.insert_text((120, 245), "04", fontsize=12, color=(1, 0, 0))
    page.insert_text((170, 245), "01", fontsize=12, color=(1, 0, 0))

    markers, candidates = extract_spool_markers(page)

    assert sorted(marker.value for marker in markers) == [
        "00001", "123456", "A-17/B", "S00166",
    ]
    assert [(item.value, item.reason) for item in candidates] == [
        ("777", "unboxed_red_digit_candidate"),
    ]
    document.close()


def test_red_bolt_up_box_is_never_a_spool_or_review_candidate():
    document = fitz.open()
    page = document.new_page(width=900, height=700)
    rect = fitz.Rect(50, 90, 220, 145)
    page.draw_rect(rect, color=(1, 0, 0), width=1)
    page.insert_text((58, 112), "BU-D92A5C3E", fontsize=11, color=(1, 0, 0))
    page.insert_text((58, 134), "777", fontsize=11, color=(1, 0, 0))

    markers, candidates = extract_spool_markers(page)

    assert markers == []
    assert candidates == []
    document.close()


def test_spool_inclusive_fmr_writes_spools_then_complete_bom(tmp_path):
    source = tmp_path / "input"
    output = tmp_path / "output"
    source.mkdir()
    _save_package(source / "package.pdf", [("ISO-001", "0", ["0007"])])

    summary = run_fmr(
        source, output, FMR_TEMPLATE, include_spool_numbers=True
    )

    assert summary["spool_numbers_included"] is True
    assert summary["spool_numbers_only"] is False
    assert summary["spool_rows_written"] == 1
    assert summary["spool_review_entries"] == 0
    assert summary["pipe_rows_included"] is True
    assert summary["estimated_time_saved_seconds"] == "90.66"
    workbook = load_workbook(summary["workbook"])
    sheet = workbook["IWP-SPOOL-100(00)"]
    assert (sheet["B8"].value, sheet["C8"].value, sheet["D8"].value, sheet["E8"].value) == (
        "SPOOL 0007", None, 1, None,
    )
    assert sheet["B8"].number_format == "@"
    assert (sheet["B9"].value, sheet["E9"].value) == ("PIPE-001", "PIPE SCH 10S")
    assert (sheet["B10"].value, sheet["E10"].value) == ("ELL-001", "ELL 90 LR")
    workbook.close()
    manifest_path = next((output / "analytics_runs").glob("*.json"))
    manifest = json.loads(manifest_path.read_text())
    assert manifest["isos"][0]["spool_rows"] == 1
    assert manifest["isos"][0]["material_rows"] == 2
    assert manifest["isos"][0]["pipe_rows"] == 1


def test_spool_only_keeps_blank_sheet_and_quarantines_missing_number(tmp_path):
    source = tmp_path / "input"
    output = tmp_path / "output"
    source.mkdir()
    _save_package(source / "package.pdf", [("ISO-EMPTY", "0", [])])

    summary = run_fmr(
        source, output, FMR_TEMPLATE, spool_numbers_only=True
    )

    assert summary["status"] == "complete"
    assert summary["fmr_sheets_created"] == 1
    assert summary["pipe_rows_included"] is False
    assert summary["spool_rows_written"] == 0
    assert summary["spool_review_entries"] == 1
    workbook = load_workbook(summary["workbook"])
    sheet = workbook["IWP-SPOOL-100(00)"]
    assert all(sheet.cell(8, column).value is None for column in range(2, 12))
    workbook.close()
    rows = _quarantine_rows(output)
    assert [row["reason_code"] for row in rows] == ["no_spool_number_detected"]


def test_same_revision_copies_write_only_common_spools(tmp_path):
    source = tmp_path / "input"
    output = tmp_path / "output"
    source.mkdir()
    _save_package(source / "package.pdf", [
        ("ISO-DUP", "2", ["0001", "0002"]),
        ("ISO-DUP", "2", ["0001", "0003"]),
    ])

    summary = run_fmr(
        source, output, FMR_TEMPLATE, include_spool_numbers=True
    )

    assert summary["duplicate_pages_ignored"] == 1
    assert summary["spool_rows_written"] == 1
    assert summary["spool_review_entries"] == 2
    workbook = load_workbook(summary["workbook"])
    sheet = workbook["IWP-SPOOL-100(00)"]
    assert sheet["B8"].value == "SPOOL 0001"
    assert sheet["B9"].value == "PIPE-001"
    workbook.close()
    rows = _quarantine_rows(output)
    assert [row["reason_code"] for row in rows] == [
        "spool_number_duplicate_conflict",
        "spool_number_duplicate_conflict",
    ]
    assert "0002" in rows[0]["reason_detail"]
    assert "0003" in rows[1]["reason_detail"]


def test_fmr_cli_material_modes_are_mutually_exclusive():
    for first, second in [
        ("--pipe-only", "--spool-numbers"),
        ("--spool-numbers", "--spool-numbers-only"),
        ("--pipe", "--spool-numbers-only"),
    ]:
        with pytest.raises(SystemExit):
            build_parser().parse_args([
                "--input", "input", "--output", "output", first, second,
            ])


def test_fmr_cli_routes_spool_numbers_to_pipeline(tmp_path, capsys):
    source = tmp_path / "input"
    output = tmp_path / "output"
    source.mkdir()
    _save_package(source / "package.pdf", [("ISO-CLI", "0", ["0012"])])

    assert fmr_main([
        "--input", str(source),
        "--output", str(output),
        "--template", str(FMR_TEMPLATE),
        "--spool-numbers",
    ]) == 0

    summary = json.loads(capsys.readouterr().out)
    assert summary["spool_numbers_included"] is True
    workbook = load_workbook(summary["workbook"])
    assert workbook["IWP-SPOOL-100(00)"]["B8"].value == "SPOOL 0012"
    workbook.close()


def test_spool_rows_count_toward_dynamic_fmr_capacity(tmp_path):
    materials = [
        BomRow(
            point_number=str(index),
            description="ELL 90 LR",
            nominal_size="2",
            commodity_code=f"CODE-{index}",
            quantity="1",
            raw_text="ELL 90 LR",
            bbox=(0, 0, 1, 1),
        )
        for index in range(1, MATERIAL_CAPACITY + 1)
    ]
    iso = IsoPage(
        source_pdf="iso.pdf",
        source_path="iso.pdf",
        page=1,
        drawing_number="ISO-CAPACITY",
        revision="0",
        materials=materials,
        raw_text="source",
        content_hash="hash",
        spool_numbers=["0001", "0002"],
    )
    output = tmp_path / "capacity.xlsx"

    build_fmr_workbook(FMR_TEMPLATE, output, "IWP-CAPACITY", [iso])

    workbook = load_workbook(output)
    sheet = workbook["IWP-CAPACITY(00)"]
    assert sheet["B8"].value == "SPOOL 0001"
    assert sheet["B9"].value == "SPOOL 0002"
    assert sheet[f"B{MATERIAL_LAST_ROW + 2}"].value == f"CODE-{MATERIAL_CAPACITY}"
    assert sheet[f"B{MATERIAL_LAST_ROW + 3}"].value == "REASON REQUIRED"
    workbook.close()
