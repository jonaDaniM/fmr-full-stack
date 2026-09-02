import csv
import json
from pathlib import Path

import fitz
import pytest
from openpyxl import load_workbook

from iso_bom.mto_cli import main as mto_cli_main
from iso_bom.mto_parser import (
    classify_bolt_gasket,
    clean_material_description,
    extract_cwa,
    extract_pipe_schedule,
    material_for_scope,
    is_mto_material,
    parse_mto_iso_page,
    sheet_number,
)
from iso_bom.mto_pipeline import run_mto
from iso_bom.mto_model import (
    MTO_SCOPE_ALL_MATERIALS,
    MTO_SCOPE_BOLTS_GASKETS,
    MtoIsoPage,
    MtoMaterialRow,
)
from iso_bom.mto_workbook import build_mto_workbook, material_category_sheet
from iso_bom.model import BomRow


PROJECT_ROOT = Path(__file__).resolve().parents[1]
MTO_TEMPLATE = PROJECT_ROOT / "templates" / "MTO" / "Takeoff Spreadsheet Template.xlsx"
SUPPLIED_MTO_ISO = PROJECT_ROOT / "input" / "newFmr12" / "LP131-PV-941006-03_R0 FAB.PDF"
SUPPLIED_MTO_EXAMPLE = (
    PROJECT_ROOT / "templates" / "createMTO" / "MTO IP-SMM10D0012FPP-K447-135.xlsx"
)
MTO4_PACKAGE = PROJECT_ROOT / "input" / "newMTO4" / "8. IP-SMM10B0013FPP-K447-110 Combined.pdf"


def _insert_cover(document: fitz.Document, iwp_number: str, cwa: str = "10D") -> None:
    page = document.new_page(width=612, height=792)
    page.insert_text((170, 35), "Installation Work Package - TEST", fontsize=12)
    page.insert_text((50, 75), "IWP Number:", fontsize=10)
    page.insert_text((130, 75), iwp_number, fontsize=10)
    page.insert_text((50, 105), "Status: Issued to Construction", fontsize=10)
    if cwa:
        page.insert_text((260, 105), "CWA:", fontsize=10)
        page.insert_text((305, 105), cwa, fontsize=10)


def _insert_iso(
    document: fitz.Document,
    drawing_number: str,
    revision: str,
    rows,
    *,
    pipe_schedule: str = "315",
) -> None:
    page = document.new_page(width=900, height=700)
    page.insert_text((420, 35), "BILL OF MATERIALS", fontsize=9)
    for x, text in [
        (420, "PT"), (455, "DESCRIPTION"), (650, "NPD"),
        (710, "CMDTY"), (830, "QTY"),
    ]:
        page.insert_text((x, 60), text, fontsize=8)
    y = 90
    for row in rows:
        for x, text in zip((420, 455, 650, 710, 830), row):
            if text:
                page.insert_text((x, y), text, fontsize=8)
        y += 28
    if pipe_schedule is not None:
        page.insert_text((420, 560), "PIPE SCHEDULE:", fontsize=8)
        page.insert_text((525, 560), pipe_schedule, fontsize=8)
    page.insert_text((500, 610), "ISOMETRIC DRAWING NUMBER", fontsize=8)
    page.insert_text((790, 610), "REV", fontsize=8)
    page.insert_text((500, 635), drawing_number, fontsize=8)
    page.insert_text((795, 635), revision, fontsize=8)


def _create_mto_package(path: Path) -> None:
    document = fitz.open()
    _insert_cover(document, "IWP-MTO-100", "Andrew Longcake")
    _insert_iso(document, "LP131-PV-941006-03", "0", [
        ("1", "PIPE SCH 10S", "4", "P-OLD", "12.5'"),
        ("2", "ELL 90 LR SCH 10S", "4", "E-OLD", "2"),
        ("3", "BIRDSCREEN 316 SS 45 DEG", "4", "OLD-CODE", "1"),
    ])
    _insert_iso(document, "LP131-PV-941006-03", "2", [
        ("1", "PIPE SCH 10S", "4", "P-001", "12.5'"),
        ("2", "ELL 90 LR SCH 10S", "4", "E-001", "2"),
        ("3", "BIRDSCREEN 316 SS 45 DEG", "4", "LP131 BS-9411V1", "1"),
        ("4", "5UG, U-BOLT GUIDE", "4", "5UG-04", "4"),
    ])
    _insert_iso(document, "LP131-PV-941006-04", "0", [
        ("1", "TEE SCH 10S", "4", "T-001", "1"),
        ("2", '5CI, ISOLATION CRADLE, 4" PIPE, SS', "4", "5CI-04", "6"),
    ], pipe_schedule="316")
    document.save(path)
    document.close()


def _row(description: str) -> BomRow:
    return BomRow("1", description, "4", "CODE-1", "1", description, (0, 0, 1, 1))


def test_cwa_extraction_handles_codes_names_missing_and_multiple():
    document = fitz.open()
    _insert_cover(document, "IWP-CWA-1", "10D")
    _insert_cover(document, "IWP-CWA-2", "Andrew Longcake")
    _insert_cover(document, "IWP-CWA-3", "")
    _insert_cover(document, "IWP-CWA-4", "10D")
    document[-1].insert_text((260, 125), "CWA:", fontsize=10)
    document[-1].insert_text((305, 125), "11E", fontsize=10)

    assert extract_cwa(document[0]) == ("10D", [])
    assert extract_cwa(document[1]) == ("Andrew Longcake", [])
    assert extract_cwa(document[2]) == ("", ["cwa_not_detected"])
    assert extract_cwa(document[3]) == ("", ["multiple_cwa_values_on_cover"])
    document.close()


def test_mto_iso_parser_reads_schedule_sheet_and_multi_token_commodity():
    document = fitz.open()
    _insert_iso(document, "LP131-PV-941006-03", "0", [
        ("1", "PIPE SCH 10S", "4", "P-001", "12.5'"),
        ("2", "ELL 90 LR SCH 10S", "4", "E-001", "2"),
        ("3", "BIRDSCREEN 316 SS 45 DEG", "4", "LP131 BS-9411V1", "1"),
    ])
    page = document[0]
    assert extract_pipe_schedule(page) == ("315", [])
    assert sheet_number("LP131-PV-941006-03") == ("3", [])
    parsed = parse_mto_iso_page(page, 1, "package.pdf", "package.pdf")
    document.close()

    assert parsed.drawing_number == "LP131-PV-941006-03"
    assert parsed.pipe_schedule == "315"
    assert parsed.sheet_number == "3"
    assert parsed.review_reasons == []
    assert [row.description for row in parsed.materials] == ["BIRDSCREEN 316 SS 45 DEG"]
    assert parsed.materials[0].commodity_code == "LP131 BS-9411V1"
    assert not is_mto_material(_row("PIPE SCH 10S"))
    assert not is_mto_material(_row("ELL 90 LR SCH 10S"))
    assert is_mto_material(_row('5CI, ISOLATION CRADLE, 4" PIPE, SS'))


@pytest.mark.parametrize(("description", "code", "item_type"), [
    ("STUD-BOLT A193 GR B7 W/A194 GR 2H NUTS", "5675648L0", "BOLT"),
    ("MACHINE BOLT A307", "B-1", "BOLT"),
    ("ANCHOR-BOLT GALV", "B-2", "BOLT"),
    ("HEX BOLT A325", "B-3", "BOLT"),
    ("CAP SCREW HEX HEAD CR-MO", "5676576L1", "BOLT"),
    ("EARTHING WASHER 316 SS", "5676235L4", "WASHER"),
    ("GASKET 150# PTFE RING", "5669399L1", "GASKET"),
])
def test_bolt_gasket_classifier_includes_observed_standalone_families(
    description, code, item_type
):
    assert classify_bolt_gasket(description, code) == item_type
    assert not is_mto_material(BomRow(
        "1", description, "4", code, "1", description, (0, 0, 1, 1)
    ))


@pytest.mark.parametrize(("description", "code"), [
    ("5UG, U-BOLT GUIDE", "5UG-04"),
    ("5UGSP, U-BOLT GUIDE SPECIAL", "5UGSP-04"),
    ("5US, U-BOLT STOP", "5US-04"),
    ("5MUG12, MODIFIED U-BOLT GUIDE", "5MUG12-04"),
    ("COUPLING WITH GASKET AND RETAINER", "5714906L0"),
    ("FLG ADAPTER WITH GASKET", "5484771L0"),
])
def test_bolt_gasket_classifier_excludes_supports_and_gasket_assemblies(
    description, code
):
    assert classify_bolt_gasket(description, code) is None


def test_bolt_gasket_cleanup_preserves_code_suffixes_and_removes_revision_digits():
    assert clean_material_description(
        "STUD-BOLT A193 GR B7 - 3.75 in. Length 1"
    ) == "STUD-BOLT A193 GR B7 - 3.75 in. Length"
    assert clean_material_description(
        "STUD-BOLT A193 GR B7 - 1 3.75 in. Length"
    ) == "STUD-BOLT A193 GR B7 - 3.75 in. Length"
    assert clean_material_description(
        "1 PIPE SCH 10S ERW 316/316L SS A312"
    ) == "PIPE SCH 10S ERW 316/316L SS A312"
    for code in ("5676576L0", "5676576L1", "5676576L4"):
        material = material_for_scope(
            BomRow("1", "", "4", code, "1", code, (0, 0, 1, 1)),
            MTO_SCOPE_BOLTS_GASKETS,
        )
        assert material is not None
        assert material.commodity_code == code
        assert material.item_type == "BOLT"


@pytest.mark.parametrize(("field", "row", "expected_partial"), [
    (
        "description",
        BomRow("1", "", "4", "5676576L1", "1", "raw", (0, 0, 1, 1), ["missing_description"]),
        ["description"],
    ),
    (
        "size",
        BomRow("2", "GASKET 150# PTFE RING", "", "5669399L1", "1", "raw", (0, 0, 1, 1)),
        ["size"],
    ),
    (
        "commodity_code",
        BomRow("3", "CAP SCREW HEX HEAD", "4", "B-1 B-2", "1", "raw", (0, 0, 1, 1), ["multiple_commodity_code_candidates"]),
        ["commodity_code"],
    ),
    (
        "quantity",
        BomRow("4", "GASKET 300# SPIRAL WOUND", "4", "G-1", "1", "raw", (0, 0, 1, 1), ["multiple_quantity_candidates"]),
        ["quantity"],
    ),
])
def test_partial_bolt_gasket_material_blanks_only_uncertain_cell(
    field, row, expected_partial
):
    material = material_for_scope(row, MTO_SCOPE_BOLTS_GASKETS)
    assert material is not None
    assert getattr(material, "nominal_size" if field == "size" else field) == ""
    assert material.partial_fields == expected_partial
    confident = {
        "description": material.description,
        "size": material.nominal_size,
        "commodity_code": material.commodity_code,
        "quantity": material.quantity,
    }
    assert all(value for name, value in confident.items() if name != field)


def test_unknown_partial_row_is_not_classified_without_keyword_or_known_code():
    row = BomRow("1", "UNREADABLE MATERIAL", "4", "UNKNOWN-1", "1", "raw", (0, 0, 1, 1))
    assert material_for_scope(row, MTO_SCOPE_BOLTS_GASKETS) is None


@pytest.mark.parametrize(("description", "code", "quantity", "item_type", "uom"), [
    ("PIPE SCH 10S ERW 316 SS", "P-1", "18.9'", "PIPE", "LF"),
    ("ELL 90 LR SCH 10S 316 SS", "E-1", "2", "FITTING", "EA"),
    ("PIPET SCH 10S 316 SS", "PIPET-1", "1", "FITTING", "EA"),
    ("STUD-BOLT A193 GR B7", "5675648L0", "4", "BOLT", "EA"),
    ("GASKET 150# PTFE RING", "5669399L0", "1", "GASKET", "EA"),
    ("5CI, ISOLATION CRADLE", "5CI-04", "1", "SUPPORT", "EA"),
])
def test_all_materials_scope_keeps_every_material_family(
    description, code, quantity, item_type, uom
):
    material = material_for_scope(
        BomRow("1", description, "4", code, quantity, description, (0, 0, 1, 1)),
        MTO_SCOPE_ALL_MATERIALS,
    )
    assert material is not None
    assert material.item_type == item_type
    assert material.uom == uom
    assert material.quantity == ("18.9" if uom == "LF" else quantity)


@pytest.mark.parametrize(("item_type", "description", "sheet_name"), [
    ("FITTING", "FLG BLIND 150# RF 316 SS", "BLINDS"),
    ("GASKET", "GASKET 150# PTFE RING", "BOLTS & GASKETS"),
    ("SUPPORT", '5CI, ISOLATION CRADLE, 4" PIPE, SS', "SUPPORTS"),
    ("SPECIALTY", "BIRDSCREEN 316 SS 45 DEG", "BIRDSCREENS"),
    ("SPECIALTY", "BALL 150# RF 316SS TFE FP HNDL OP", "VALVES"),
    ("SPECIALTY", "FV - ON/OFF VALVE, BALL", "VALVES"),
    ("PIPE", "PIPE SCH 10S ERW 316 SS", "PIPE & FITTINGS"),
    ("SPECIALTY", "STUB END SCH10S 316 SS", "PIPE & FITTINGS"),
    ("SPECIALTY", "THERMOWELL ASSEMBLY", "OTHER MATERIALS"),
])
def test_consolidated_material_category_routing(
    item_type, description, sheet_name
):
    assert material_category_sheet(item_type, description) == sheet_name


def test_mto_pipeline_creates_single_combined_workbook(tmp_path):
    source = tmp_path / "input"
    output = tmp_path / "output"
    source.mkdir()
    _create_mto_package(source / "package.pdf")

    summary = run_mto(source, output, MTO_TEMPLATE)
    assert summary["status"] == "complete"
    assert summary["iwp_number"] == "IWP-MTO-100"
    assert summary["cwa"] == "Andrew Longcake"
    assert summary["iso_pages_detected"] == 3
    assert summary["unique_drawings"] == 2
    assert summary["selected_iso_drawings"] == 2
    assert summary["superseded_pages_ignored"] == 1
    assert summary["mto_rows_created"] == 3
    assert summary["quarantine_entries"] == 0
    assert summary["time_estimate_available"] is False
    assert summary["time_model_version"] == ""

    workbook_path = Path(summary["workbook"])
    assert workbook_path.name == "MTO IWP-MTO-100.xlsx"
    workbook = load_workbook(workbook_path, data_only=True)
    assert workbook.sheetnames == [
        "COMBINED", "BLINDS", "PIPE & FITTINGS", "BOLTS & GASKETS", "TOTALS"
    ]
    combined = workbook["COMBINED"]
    assert combined["B1"].value == "IWP"
    assert combined["H1"].value is None
    assert combined.column_dimensions["B"].width >= 30
    assert combined.max_row == 4
    assert sorted(table.ref for table in combined.tables.values()) == ["A1:G4", "I1:K4"]
    assert combined.sheet_view.showGridLines is False
    assert "A1:K4" in str(combined.print_area).replace("$", "")
    values = [
        [combined[f"{column}{row}"].value for column in ("A", "B", "C", "D", "E", "F", "G", "I", "J", "K")]
        for row in range(2, 5)
    ]
    assert values == [
        [
            "Andrew Longcake", "IWP-MTO-100", "LP131-PV-941006-03", 3, 315,
            "BIRDSCREEN 316 SS 45 DEG", 4, "LP131 BS-9411V1", 1, "EA",
        ],
        [
            "Andrew Longcake", "IWP-MTO-100", "LP131-PV-941006-03", 3, 315,
            "5UG, U-BOLT GUIDE", 4, "5UG-04", 4, "EA",
        ],
        [
            "Andrew Longcake", "IWP-MTO-100", "LP131-PV-941006-04", 4, 316,
            '5CI, ISOLATION CRADLE, 4" PIPE, SS', 4, "5CI-04", 6, "EA",
        ],
    ]
    assert all(
        workbook[sheet].max_row == 1
        for sheet in ("BLINDS", "PIPE & FITTINGS", "BOLTS & GASKETS")
    )
    assert all(
        workbook[sheet]["A2"].value is None
        for sheet in ("BLINDS", "PIPE & FITTINGS", "BOLTS & GASKETS")
    )
    totals = workbook["TOTALS"]
    assert totals.tables["MtoTotalsTable"].ref == "A1:I4"
    assert totals["G2"].value is None
    assert totals["G2"].number_format == "General"
    assert totals["H2"].number_format == "General"
    workbook.close()

    manifest_path = next((output / "analytics_runs").glob("*.json"))
    manifest = json.loads(manifest_path.read_text())
    assert manifest["workflow"] == "mto"
    assert manifest["package_number"] == "IWP-MTO-100"
    assert manifest["cwa"] == "Andrew Longcake"
    assert manifest["mto_scope"] == "combined"
    assert [item["material_rows"] for item in manifest["isos"]] == [2, 1]
    with (output / "mto_quarantine.csv").open() as handle:
        assert list(csv.DictReader(handle)) == []

    with pytest.raises(FileExistsError, match="already exists"):
        run_mto(source, output, MTO_TEMPLATE)
    overwritten = run_mto(source, output, MTO_TEMPLATE, overwrite=True)
    assert overwritten["status"] == "complete"


def test_bolts_gaskets_scope_writes_distinct_tab_filename_totals_and_partials(tmp_path):
    source = tmp_path / "input"
    output = tmp_path / "output"
    source.mkdir()
    document = fitz.open()
    _insert_cover(document, "IWP-MTO-BG-1", "10D")
    _insert_iso(document, "LP131-WPC-862026-01", "0", [
        ("1", "STUD-BOLT A193 GR B7 W/A194 GR 2H NUTS", "5/8", "5675648L0", "8"),
        ("2", "5UG, U-BOLT GUIDE", "4", "5UG-04", "2"),
        ("3", "GASKET 150# PTFE RING", "4", "5669399L1", "1.5"),
        ("4", "COUPLING WITH GASKET AND RETAINER", "4", "5714906L0", "1"),
    ])
    _insert_iso(document, "LP131-SCR-387001-09", "1", [
        ("1", "CAP SCREW HEX HEAD CR-MO", "", "5676576L4", "4"),
        ("2", "EARTHING WASHER 316 SS", "4", "5676235L", "2"),
        ("3", "FLG ADAPTER WITH GASKET", "4", "5484771L0", "1"),
    ], pipe_schedule="317")
    document.save(source / "package.pdf")
    document.close()

    combined_summary = run_mto(source, output, MTO_TEMPLATE)
    bolts_summary = run_mto(
        source,
        output,
        MTO_TEMPLATE,
        scope=MTO_SCOPE_BOLTS_GASKETS,
    )
    assert Path(combined_summary["workbook"]).name == "MTO IWP-MTO-BG-1.xlsx"
    assert Path(bolts_summary["workbook"]).name == "MTO IWP-MTO-BG-1 - BOLTS & GASKETS.xlsx"
    assert Path(combined_summary["workbook"]).is_file()
    assert bolts_summary["mto_scope"] == "bolts-gaskets"
    assert bolts_summary["mto_rows"] == 4
    assert bolts_summary["partial_rows"] == 1

    workbook = load_workbook(bolts_summary["workbook"], data_only=False)
    worksheet = workbook["BOLTS & GASKETS"]
    assert worksheet["B1"].value == "IWP"
    assert worksheet.tables["BoltsGasketsMtoTable"].ref == "A1:M5"
    values = [
        [worksheet.cell(row, column).value for column in range(1, 14)]
        for row in range(2, 6)
    ]
    assert [row[5] for row in values] == [
        "STUD-BOLT A193 GR B7 W/A194 GR 2H NUTS",
        "GASKET 150# PTFE RING",
        "CAP SCREW HEX HEAD CR-MO",
        "EARTHING WASHER 316 SS",
    ]
    assert all(row[0:3] == ["10D", "IWP-MTO-BG-1", row[2]] for row in values)
    assert values[2][6] is None
    assert all(row[10:13] == [None, None, None] for row in values)
    assert workbook["COMBINED"].max_row == 1
    assert workbook["COMBINED"]["A2"].value is None

    totals = workbook["TOTALS"]
    assert totals.tables["MtoTotalsTable"].ref == "A1:I4"
    assert totals["A2"].value == "BOLTS & GASKETS"
    assert totals["G2"].value.startswith("=SUMIFS('BOLTS & GASKETS'!")
    assert totals["H2"].value == '=IF(F2="LF",CEILING(G2,$K$2),G2)'
    assert totals["I2"].value == '=IF(F2="LF",TEXT($K$2,"0")&" FT STOCK","EXACT QTY")'
    assert totals["G2"].number_format == "General"
    assert totals["H2"].number_format == "General"
    assert totals["K2"].value == 20
    workbook.close()

    with (output / "mto_quarantine.csv").open() as handle:
        assert list(csv.DictReader(handle)) == []
    manifest = json.loads(sorted((output / "analytics_runs").glob("*.json"))[-1].read_text())
    assert manifest["mto_scope"] == "bolts-gaskets"
    assert manifest["metrics"]["partial_rows"] == 1
    assert sum(iso["partial_rows"] for iso in manifest["isos"]) == 1


def test_all_materials_scope_writes_every_bom_row_to_distinct_workbook(tmp_path):
    source = tmp_path / "input"
    output = tmp_path / "output"
    source.mkdir()
    document = fitz.open()
    _insert_cover(document, "IWP-MTO-ALL", "10D")
    _insert_iso(document, "LP131-PV-941006-03", "0", [
        ("1", "PIPE SCH 10S ERW 316 SS", "4", "P-1", "18.9'"),
        ("2", "ELL 90 LR SCH 10S 316 SS", "4", "E-1", "2"),
        ("3", "PIPET SCH 10S 316 SS", "4X1", "PIPET-1", "1"),
        ("4", "GASKET 150# PTFE RING", "4", "5669399L0", "1"),
        ("5", "5CI, ISOLATION CRADLE", "4", "5CI-04", "1"),
    ])
    document.save(source / "package.pdf")
    document.close()

    summary = run_mto(
        source,
        output,
        MTO_TEMPLATE,
        scope=MTO_SCOPE_ALL_MATERIALS,
    )

    assert summary["status"] == "complete"
    assert summary["mto_scope"] == "all-materials"
    assert summary["mto_rows"] == 5
    assert Path(summary["workbook"]).name == "MTO IWP-MTO-ALL - ALL MATERIALS.xlsx"
    workbook = load_workbook(summary["workbook"], data_only=False)
    combined = workbook["COMBINED"]
    assert combined.max_row == 6
    assert [combined[f"F{row}"].value for row in range(2, 7)] == [
        "PIPE SCH 10S ERW 316 SS",
        "ELL 90 LR SCH 10S 316 SS",
        "PIPET SCH 10S 316 SS",
        "GASKET 150# PTFE RING",
        "5CI, ISOLATION CRADLE",
    ]
    assert (combined["J2"].value, combined["K2"].value) == (18.9, "LF")
    assert combined.tables["CombinedMtoDetailsTable"].ref == "A1:G6"
    assert workbook["TOTALS"]["I2"].value == (
        '=IF(F2="LF",TEXT($K$2,"0")&" FT STOCK","EXACT QTY")'
    )
    workbook.close()


def test_totals_formulas_support_lf_stock_rounding(tmp_path):
    iso = MtoIsoPage(
        source_pdf="iso.pdf",
        source_path="iso.pdf",
        page=1,
        drawing_number="LP131-PV-941006-03",
        revision="0",
        pipe_schedule="315",
        sheet_number="3",
        materials=[
            MtoMaterialRow("1", "PIPE TYPE A", "4", "P-A", "11", "", "PIPE", uom="LF"),
            MtoMaterialRow("2", "PIPE TYPE B", "6", "P-B", "24", "", "PIPE", uom="LF"),
        ],
        raw_text="",
        content_hash="hash",
    )
    output = tmp_path / "totals.xlsx"
    build_mto_workbook(MTO_TEMPLATE, output, "IWP-LF", "10D", [iso])
    totals = load_workbook(output, data_only=False)["TOTALS"]
    assert totals["G2"].value.startswith("=SUMIFS(")
    assert totals["H2"].value == '=IF(F2="LF",CEILING(G2,$K$2),G2)'
    assert totals["H3"].value == '=IF(F3="LF",CEILING(G3,$K$2),G3)'
    assert (11 + 20 - 1) // 20 * 20 == 20
    assert (24 + 20 - 1) // 20 * 20 == 40


def test_bolts_gaskets_accepts_legacy_wer_header(tmp_path):
    legacy_template = tmp_path / "legacy-template.xlsx"
    workbook = load_workbook(MTO_TEMPLATE)
    workbook["BOLTS & GASKETS"]["B1"] = "WER"
    workbook.save(legacy_template)
    workbook.close()

    source = tmp_path / "input"
    source.mkdir()
    document = fitz.open()
    _insert_cover(document, "IWP-WER", "10D")
    _insert_iso(document, "LP131-PV-941006-03", "0", [
        ("1", "GASKET 150# PTFE RING", "4", "5669399L0", "1"),
    ])
    document.save(source / "package.pdf")
    document.close()
    summary = run_mto(
        source,
        tmp_path / "output",
        legacy_template,
        scope=MTO_SCOPE_BOLTS_GASKETS,
    )
    generated = load_workbook(summary["workbook"], data_only=True)
    assert generated["BOLTS & GASKETS"]["B1"].value == "IWP"
    generated.close()


def test_generated_sheet_removes_legacy_far_right_totals(tmp_path):
    legacy_template = tmp_path / "legacy-totals.xlsx"
    workbook = load_workbook(MTO_TEMPLATE)
    combined = workbook["COMBINED"]
    combined["L1"] = "TOTAL PIPE A"
    combined["R2"] = 40
    workbook.save(legacy_template)
    workbook.close()

    source = tmp_path / "input"
    source.mkdir()
    document = fitz.open()
    _insert_cover(document, "IWP-LEGACY-TOTALS", "10D")
    _insert_iso(document, "LP131-PV-941006-03", "0", [
        ("1", "BIRDSCREEN 316 SS 45 DEG", "4", "BS-1", "1"),
    ])
    document.save(source / "package.pdf")
    document.close()
    summary = run_mto(source, tmp_path / "output", legacy_template)
    generated = load_workbook(summary["workbook"], data_only=False)
    assert generated["COMBINED"].max_column == 11
    assert generated["COMBINED"]["L1"].value is None
    assert "TOTALS" in generated.sheetnames
    generated.close()


def test_invalid_scope_template_is_blocking(tmp_path):
    invalid_template = tmp_path / "invalid-template.xlsx"
    workbook = load_workbook(MTO_TEMPLATE)
    del workbook["BOLTS & GASKETS"]
    workbook.save(invalid_template)
    workbook.close()

    source = tmp_path / "input"
    source.mkdir()
    document = fitz.open()
    _insert_cover(document, "IWP-INVALID-TEMPLATE", "10D")
    _insert_iso(document, "LP131-PV-941006-03", "0", [
        ("1", "GASKET 150# PTFE RING", "4", "5669399L0", "1"),
    ])
    document.save(source / "package.pdf")
    document.close()

    with pytest.raises(ValueError, match="missing worksheet 'BOLTS & GASKETS'"):
        run_mto(
            source,
            tmp_path / "output",
            invalid_template,
            scope=MTO_SCOPE_BOLTS_GASKETS,
        )


@pytest.mark.skipif(
    not SUPPLIED_MTO_ISO.exists() or not SUPPLIED_MTO_EXAMPLE.exists(),
    reason="Local supplied MTO example files are unavailable",
)
def test_supplied_mto_example_regression(tmp_path):
    summary = run_mto(
        SUPPLIED_MTO_ISO.parent,
        tmp_path / "output",
        MTO_TEMPLATE,
        iwp_number_override="IP-SMM10D0012FPP-K447-135",
        cwa_override="10D",
    )
    assert summary["status"] == "complete"
    assert summary["mto_rows_created"] == 6

    generated = load_workbook(summary["workbook"], data_only=True)["COMBINED"]
    expected = load_workbook(SUPPLIED_MTO_EXAMPLE, data_only=True)["COMBINED"]
    for row in range(2, 8):
        assert [generated[f"{column}{row}"].value for column in "ABCDEFGHIJK"] == [
            expected[f"{column}{row}"].value for column in "ABCDEFGHIJK"
        ]
    assert generated["A8"].value is None


def test_mto_blocks_missing_and_conflicting_package_fields(tmp_path):
    source = tmp_path / "input"
    source.mkdir()
    document = fitz.open()
    _insert_cover(document, "IWP-MISSING-CWA", "")
    _insert_iso(document, "LP131-PV-941006-03", "0", [
        ("1", "BIRDSCREEN 316 SS 45 DEG", "4", "BS-1", "1"),
    ])
    document.save(source / "package.pdf")
    document.close()

    missing = run_mto(source, tmp_path / "missing-output", MTO_TEMPLATE)
    assert missing["status"] == "blocked"
    assert missing["workbook"] == ""
    with (tmp_path / "missing-output" / "mto_quarantine.csv").open() as handle:
        assert [row["reason_code"] for row in csv.DictReader(handle)] == ["cwa_missing"]

    conflict = run_mto(
        source,
        tmp_path / "conflict-output",
        MTO_TEMPLATE,
        cwa_override="10D",
    )
    assert conflict["status"] == "complete"

    conflict_cover = tmp_path / "conflict-cover"
    conflict_cover.mkdir()
    document = fitz.open()
    _insert_cover(document, "IWP-CWA-CONFLICT", "Andrew Longcake")
    _insert_iso(document, "LP131-PV-941006-03", "0", [
        ("1", "BIRDSCREEN 316 SS 45 DEG", "4", "BS-1", "1"),
    ])
    document.save(conflict_cover / "package.pdf")
    document.close()
    summary = run_mto(
        conflict_cover,
        tmp_path / "manual-conflict-output",
        MTO_TEMPLATE,
        cwa_override="10D",
    )
    assert summary["status"] == "blocked"
    assert summary["workbook"] == ""


def test_mto_keeps_blank_iso_metadata_but_blocks_duplicate_conflicts(tmp_path):
    invalid_schedule = tmp_path / "invalid-schedule"
    invalid_schedule.mkdir()
    document = fitz.open()
    _insert_cover(document, "IWP-MTO-BAD", "10D")
    _insert_iso(document, "LP131-PV-941006-03", "0", [
        ("1", "BIRDSCREEN 316 SS 45 DEG", "4", "BS-1", "1"),
    ], pipe_schedule=None)
    document.save(invalid_schedule / "package.pdf")
    document.close()
    summary = run_mto(invalid_schedule, tmp_path / "invalid-schedule-output", MTO_TEMPLATE)
    assert summary["status"] == "complete"
    assert summary["partial_rows"] == 1
    workbook = load_workbook(summary["workbook"], data_only=True)
    assert workbook["COMBINED"]["E2"].value is None
    workbook.close()
    with (tmp_path / "invalid-schedule-output" / "mto_quarantine.csv").open() as handle:
        assert list(csv.DictReader(handle)) == []

    invalid_sheet = tmp_path / "invalid-sheet"
    invalid_sheet.mkdir()
    document = fitz.open()
    _insert_cover(document, "IWP-MTO-BAD-SHEET", "10D")
    _insert_iso(document, "LP131-PV-941006", "0", [
        ("1", "BIRDSCREEN 316 SS 45 DEG", "4", "BS-1", "1"),
    ])
    document.save(invalid_sheet / "package.pdf")
    document.close()
    summary = run_mto(invalid_sheet, tmp_path / "invalid-sheet-output", MTO_TEMPLATE)
    assert summary["status"] == "complete"
    assert summary["partial_rows"] == 1
    workbook = load_workbook(summary["workbook"], data_only=True)
    assert workbook["COMBINED"]["D2"].value is None
    workbook.close()

    duplicate = tmp_path / "duplicate"
    duplicate.mkdir()
    document = fitz.open()
    _insert_cover(document, "IWP-MTO-DUP", "10D")
    _insert_iso(document, "LP131-PV-941006-03", "0", [
        ("1", "BIRDSCREEN 316 SS 45 DEG", "4", "BS-1", "1"),
    ])
    _insert_iso(document, "LP131-PV-941006-03", "0", [
        ("1", "BIRDSCREEN 316 SS 45 DEG", "4", "BS-2", "1"),
    ])
    document.save(duplicate / "package.pdf")
    document.close()
    summary = run_mto(duplicate, tmp_path / "duplicate-output", MTO_TEMPLATE)
    assert summary["status"] == "blocked"
    with (tmp_path / "duplicate-output" / "mto_quarantine.csv").open() as handle:
        assert "duplicate_revision_conflict" in {row["reason_code"] for row in csv.DictReader(handle)}


def test_unreadable_page_does_not_stop_valid_iso(tmp_path):
    source = tmp_path / "input"
    source.mkdir()
    document = fitz.open()
    sparse = document.new_page(width=612, height=792)
    sparse.insert_text((50, 50), "scan", fontsize=8)
    _insert_cover(document, "IWP-MTO-PARTIAL-PDF", "10D")
    _insert_iso(document, "LP131-PV-941006-03", "0", [
        ("1", "BIRDSCREEN 316 SS 45 DEG", "4", "BS-1", "1"),
    ])
    document.save(source / "package.pdf")
    document.close()

    summary = run_mto(source, tmp_path / "output", MTO_TEMPLATE)
    assert summary["status"] == "complete"
    assert summary["selected_iso_drawings"] == 1
    with (tmp_path / "output" / "mto_quarantine.csv").open() as handle:
        assert [row["reason_code"] for row in csv.DictReader(handle)] == ["ocr_required"]


def test_mto_cli_manual_iso_only_package(tmp_path, capsys):
    source = tmp_path / "input"
    output = tmp_path / "output"
    source.mkdir()
    document = fitz.open()
    _insert_iso(document, "LP131-PV-941006-03", "0", [
        ("1", "BIRDSCREEN 316 SS 45 DEG", "4", "BS-1", "1"),
    ])
    document.save(source / "iso.pdf")
    document.close()

    assert mto_cli_main([
        "--input", str(source),
        "--output", str(output),
        "--iwp-number", "IWP-MANUAL-MTO",
        "--cwa", "10D",
    ]) == 0
    captured = capsys.readouterr()
    assert '"status": "complete"' in captured.out
    assert (output / "MTO IWP-MANUAL-MTO.xlsx").is_file()


def test_cwa_batch_groups_pdfs_and_creates_one_workbook_per_iwp(tmp_path):
    source = tmp_path / "input"
    output = tmp_path / "output"
    source.mkdir()

    first = fitz.open()
    _insert_cover(first, "IWP-BATCH-100", "10D")
    _insert_iso(first, "LP131-PV-941006-03", "0", [
        ("1", "BIRDSCREEN 316 SS 45 DEG", "4", "BS-1", "1"),
    ])
    first.save(source / "package-one.pdf")
    first.close()

    second = fitz.open()
    _insert_iso(second, "LP131-PV-941006-04", "0", [
        ("1", '5CI, ISOLATION CRADLE, 4" PIPE, SS', "4", "5CI-04", "2"),
    ])
    second.save(source / "IP-BATCH-K447-002.pdf")
    second.close()

    summary = run_mto(
        source,
        output,
        MTO_TEMPLATE,
        cwa_override="10D",
        cwa_batch=True,
    )

    assert summary["status"] == "complete"
    assert summary["batch_mode"] == "cwa"
    assert summary["iwp_packages_detected"] == 2
    assert summary["packages_completed"] == 2
    assert summary["packages_blocked"] == 0
    assert summary["selected_iso_drawings"] == 2
    assert summary["mto_rows"] == 2
    assert {Path(path).name for path in summary["workbooks"]} == {
        "MTO IP-BATCH-K447-002.xlsx",
        "MTO IWP-BATCH-100.xlsx",
    }
    for iwp_number in ("IP-BATCH-K447-002", "IWP-BATCH-100"):
        workbook_path = output / f"MTO {iwp_number}.xlsx"
        workbook = load_workbook(workbook_path, data_only=True)
        assert workbook["COMBINED"]["B2"].value == iwp_number
        workbook.close()
        assert list((output / "_batch_runs" / iwp_number / "combined" / "analytics_runs").glob("*.json"))
    assert (output / "mto_batch_summary.json").is_file()
    with (output / "mto_batch_quarantine.csv").open() as handle:
        assert list(csv.DictReader(handle)) == []


def test_cwa_batch_continues_when_one_iwp_has_revision_conflict(tmp_path):
    source = tmp_path / "input"
    output = tmp_path / "output"
    source.mkdir()

    good = fitz.open()
    _insert_cover(good, "IWP-BATCH-GOOD", "10D")
    _insert_iso(good, "LP131-PV-941006-03", "0", [
        ("1", "BIRDSCREEN 316 SS 45 DEG", "4", "BS-1", "1"),
    ])
    good.save(source / "good.pdf")
    good.close()

    blocked = fitz.open()
    _insert_cover(blocked, "IWP-BATCH-BLOCKED", "10D")
    _insert_iso(blocked, "LP131-PV-941006-04", "0", [
        ("1", "BIRDSCREEN 316 SS 45 DEG", "4", "BS-1", "1"),
    ])
    _insert_iso(blocked, "LP131-PV-941006-04", "0", [
        ("1", "BIRDSCREEN 316 SS 45 DEG", "4", "BS-2", "1"),
    ])
    blocked.save(source / "blocked.pdf")
    blocked.close()

    summary = run_mto(
        source,
        output,
        MTO_TEMPLATE,
        cwa_override="10D",
        cwa_batch=True,
    )

    assert summary["status"] == "partial"
    assert summary["packages_completed"] == 1
    assert summary["packages_blocked"] == 1
    assert (output / "MTO IWP-BATCH-GOOD.xlsx").is_file()
    assert not (output / "MTO IWP-BATCH-BLOCKED.xlsx").exists()
    with (output / "mto_batch_quarantine.csv").open() as handle:
        reasons = {row["reason_code"] for row in csv.DictReader(handle)}
    assert "duplicate_revision_conflict" in reasons


def test_cwa_consolidation_creates_filtered_and_all_materials_workbooks(tmp_path):
    source = tmp_path / "input"
    output = tmp_path / "output"
    source.mkdir()

    first = fitz.open()
    _insert_cover(first, "IWP-CONSOLIDATE-100", "10D")
    _insert_iso(first, "LP131-PV-941006-03", "0", [
        ("1", "PIPE SCH 10S ERW 316 SS", "4", "P-1", "18.9'"),
        ("2", "BIRDSCREEN 316 SS 45 DEG", "4", "BS-1", "1"),
        ("3", "FLG BLIND 150# RF 316 SS", "4", "BL-1", "1"),
        ("4", "GASKET 150# PTFE RING", "4", "5669399L0", "1"),
    ])
    first.save(source / "package-one.pdf")
    first.close()

    second = fitz.open()
    _insert_cover(second, "IWP-CONSOLIDATE-200", "10D")
    _insert_iso(second, "LP131-PV-941006-04", "0", [
        ("1", "ELL 90 LR SCH 10S 316 SS", "4", "E-1", "2"),
        ("2", '5CI, ISOLATION CRADLE, 4" PIPE, SS', "4", "5CI-04", "3"),
        ("3", "BALL 150# RF 316SS TFE FP HNDL OP", "4", "V-1", "1"),
        ("4", "THERMOWELL ASSEMBLY", "4", "TW-1", "1"),
    ])
    second.save(source / "package-two.pdf")
    second.close()

    filtered = run_mto(
        source,
        output,
        MTO_TEMPLATE,
        cwa_override="10D",
        consolidate=True,
    )
    unrestricted = run_mto(
        source,
        output,
        MTO_TEMPLATE,
        cwa_override="10D",
        scope=MTO_SCOPE_ALL_MATERIALS,
        consolidate=True,
    )

    assert filtered["status"] == "complete"
    assert filtered["consolidation_mode"] == "cwa"
    assert filtered["packages_completed"] == 2
    assert filtered["packages_blocked"] == 0
    assert filtered["mto_rows"] == 4
    assert Path(filtered["workbook"]).name == "MTO CWA 10D - CONSOLIDATED.xlsx"
    assert unrestricted["status"] == "complete"
    assert unrestricted["mto_scope"] == "all-materials"
    assert unrestricted["mto_rows"] == 8
    assert Path(unrestricted["workbook"]).name == (
        "MTO CWA 10D - CONSOLIDATED - ALL MATERIALS.xlsx"
    )

    filtered_book = load_workbook(filtered["workbook"], data_only=False)
    filtered_sheet = filtered_book["COMBINED"]
    assert filtered_book.sheetnames == [
        "COMBINED",
        "BLINDS",
        "PIPE & FITTINGS",
        "BOLTS & GASKETS",
        "SUPPORTS",
        "VALVES",
        "BIRDSCREENS",
        "OTHER MATERIALS",
        "TOTALS",
    ]
    assert [filtered_sheet[f"B{row}"].value for row in range(2, 6)] == [
        "IWP-CONSOLIDATE-100",
        "IWP-CONSOLIDATE-200",
        "IWP-CONSOLIDATE-200",
        "IWP-CONSOLIDATE-200",
    ]
    assert [filtered_sheet[f"F{row}"].value for row in range(2, 6)] == [
        "BIRDSCREEN 316 SS 45 DEG",
        '5CI, ISOLATION CRADLE, 4" PIPE, SS',
        "BALL 150# RF 316SS TFE FP HNDL OP",
        "THERMOWELL ASSEMBLY",
    ]
    assert filtered_book["BLINDS"].max_row == 1
    assert filtered_book["PIPE & FITTINGS"].max_row == 1
    assert filtered_book["BOLTS & GASKETS"].max_row == 1
    assert filtered_book["SUPPORTS"]["F2"].value == (
        '5CI, ISOLATION CRADLE, 4" PIPE, SS'
    )
    assert filtered_book["VALVES"]["F2"].value == (
        "BALL 150# RF 316SS TFE FP HNDL OP"
    )
    assert filtered_book["BIRDSCREENS"]["F2"].value == (
        "BIRDSCREEN 316 SS 45 DEG"
    )
    assert filtered_book["OTHER MATERIALS"]["F2"].value == "THERMOWELL ASSEMBLY"
    assert filtered_book["TOTALS"]["A2"].value == "SUPPORTS"
    assert filtered_book["TOTALS"]["G2"].value.startswith("=SUMIFS('SUPPORTS'!")
    filtered_book.close()

    all_book = load_workbook(unrestricted["workbook"], data_only=True)
    all_sheet = all_book["COMBINED"]
    assert {all_sheet[f"F{row}"].value for row in range(2, 10)} == {
        "PIPE SCH 10S ERW 316 SS",
        "BIRDSCREEN 316 SS 45 DEG",
        "FLG BLIND 150# RF 316 SS",
        "GASKET 150# PTFE RING",
        "ELL 90 LR SCH 10S 316 SS",
        '5CI, ISOLATION CRADLE, 4" PIPE, SS',
        "BALL 150# RF 316SS TFE FP HNDL OP",
        "THERMOWELL ASSEMBLY",
    }
    expected_tab_descriptions = {
        "BLINDS": ["FLG BLIND 150# RF 316 SS"],
        "PIPE & FITTINGS": [
            "PIPE SCH 10S ERW 316 SS",
            "ELL 90 LR SCH 10S 316 SS",
        ],
        "BOLTS & GASKETS": ["GASKET 150# PTFE RING"],
        "SUPPORTS": ['5CI, ISOLATION CRADLE, 4" PIPE, SS'],
        "VALVES": ["BALL 150# RF 316SS TFE FP HNDL OP"],
        "BIRDSCREENS": ["BIRDSCREEN 316 SS 45 DEG"],
        "OTHER MATERIALS": ["THERMOWELL ASSEMBLY"],
    }
    for sheet_name, expected in expected_tab_descriptions.items():
        worksheet = all_book[sheet_name]
        assert worksheet["B1"].value == "IWP"
        assert [worksheet[f"F{row}"].value for row in range(2, worksheet.max_row + 1)] == expected
        assert next(iter(worksheet.tables.values())).ref == f"A1:M{worksheet.max_row}"
    all_book.close()

    assert filtered["material_rows_by_tab"] == {
        "COMBINED": 4,
        "BLINDS": 0,
        "PIPE & FITTINGS": 0,
        "BOLTS & GASKETS": 0,
        "SUPPORTS": 1,
        "VALVES": 1,
        "BIRDSCREENS": 1,
        "OTHER MATERIALS": 1,
    }
    assert unrestricted["material_rows_by_tab"] == {
        "COMBINED": 8,
        "BLINDS": 1,
        "PIPE & FITTINGS": 2,
        "BOLTS & GASKETS": 1,
        "SUPPORTS": 1,
        "VALVES": 1,
        "BIRDSCREENS": 1,
        "OTHER MATERIALS": 1,
    }

    assert (output / "mto_consolidated_summary.json").is_file()
    assert (output / "mto_consolidated_quarantine.csv").is_file()
    assert (output / "mto_consolidated_all_materials_summary.json").is_file()
    assert (output / "mto_consolidated_all_materials_quarantine.csv").is_file()


def test_cwa_consolidation_excludes_conflicted_iwp(tmp_path):
    source = tmp_path / "input"
    output = tmp_path / "output"
    source.mkdir()

    good = fitz.open()
    _insert_cover(good, "IWP-CONSOLIDATE-GOOD", "10D")
    _insert_iso(good, "LP131-PV-941006-03", "0", [
        ("1", "BIRDSCREEN 316 SS 45 DEG", "4", "BS-1", "1"),
    ])
    good.save(source / "good.pdf")
    good.close()

    blocked = fitz.open()
    _insert_cover(blocked, "IWP-CONSOLIDATE-BLOCKED", "10D")
    _insert_iso(blocked, "LP131-PV-941006-04", "0", [
        ("1", "BIRDSCREEN 316 SS 45 DEG", "4", "BS-1", "1"),
    ])
    _insert_iso(blocked, "LP131-PV-941006-04", "0", [
        ("1", "BIRDSCREEN 316 SS 45 DEG", "4", "BS-2", "1"),
    ])
    blocked.save(source / "blocked.pdf")
    blocked.close()

    summary = run_mto(
        source,
        output,
        MTO_TEMPLATE,
        cwa_override="10D",
        consolidate=True,
    )

    assert summary["status"] == "partial"
    assert summary["packages_completed"] == 1
    assert summary["packages_blocked"] == 1
    assert summary["iwp_numbers_included"] == ["IWP-CONSOLIDATE-GOOD"]
    assert summary["iwp_numbers_blocked"] == ["IWP-CONSOLIDATE-BLOCKED"]
    workbook = load_workbook(summary["workbook"], data_only=True)
    assert workbook["COMBINED"]["B2"].value == "IWP-CONSOLIDATE-GOOD"
    assert workbook["COMBINED"].max_row == 2
    workbook.close()
    with (output / "mto_consolidated_quarantine.csv").open() as handle:
        reasons = {row["reason_code"] for row in csv.DictReader(handle)}
    assert "duplicate_revision_conflict" in reasons


def test_mto_cli_cwa_batch_mode(tmp_path, capsys):
    source = tmp_path / "input"
    output = tmp_path / "output"
    source.mkdir()
    document = fitz.open()
    _insert_cover(document, "IWP-CLI-BATCH", "10D")
    _insert_iso(document, "LP131-PV-941006-03", "0", [
        ("1", "BIRDSCREEN 316 SS 45 DEG", "4", "BS-1", "1"),
    ])
    document.save(source / "package.pdf")
    document.close()

    assert mto_cli_main([
        "--input", str(source),
        "--output", str(output),
        "--cwa", "10D",
        "--cwa-batch",
    ]) == 0
    summary = json.loads(capsys.readouterr().out)
    assert summary["batch_mode"] == "cwa"
    assert summary["packages_completed"] == 1
    assert (output / "MTO IWP-CLI-BATCH.xlsx").is_file()


def test_mto_cli_consolidation_mode(tmp_path, capsys):
    source = tmp_path / "input"
    output = tmp_path / "output"
    source.mkdir()
    document = fitz.open()
    _insert_cover(document, "IWP-CLI-CONSOLIDATE", "10D")
    _insert_iso(document, "LP131-PV-941006-03", "0", [
        ("1", "BIRDSCREEN 316 SS 45 DEG", "4", "BS-1", "1"),
    ])
    document.save(source / "package.pdf")
    document.close()

    assert mto_cli_main([
        "--input", str(source),
        "--output", str(output),
        "--cwa", "10D",
        "--consolidate",
    ]) == 0
    summary = json.loads(capsys.readouterr().out)
    assert summary["consolidation_mode"] == "cwa"
    assert summary["packages_completed"] == 1
    assert (output / "MTO CWA 10D - CONSOLIDATED.xlsx").is_file()


def test_mto_cli_all_materials_bypass(tmp_path, capsys):
    source = tmp_path / "input"
    output = tmp_path / "output"
    source.mkdir()
    document = fitz.open()
    _insert_cover(document, "IWP-CLI-ALL", "10D")
    _insert_iso(document, "LP131-PV-941006-03", "0", [
        ("1", "ELL 90 LR SCH 10S 316 SS", "4", "E-1", "2"),
    ])
    document.save(source / "package.pdf")
    document.close()

    assert mto_cli_main([
        "--input", str(source),
        "--output", str(output),
        "--all-materials",
    ]) == 0
    summary = json.loads(capsys.readouterr().out)
    assert summary["mto_scope"] == "all-materials"
    assert summary["mto_rows"] == 1
    assert (output / "MTO IWP-CLI-ALL - ALL MATERIALS.xlsx").is_file()


@pytest.mark.skipif(not MTO4_PACKAGE.exists(), reason="Local mto4 package is unavailable")
def test_mto4_all_materials_regression_keeps_all_selected_bom_rows(tmp_path):
    summary = run_mto(
        MTO4_PACKAGE.parent,
        tmp_path / "output",
        MTO_TEMPLATE,
        cwa_override="10B",
        scope=MTO_SCOPE_ALL_MATERIALS,
    )

    assert summary["status"] == "complete"
    assert summary["selected_iso_drawings"] == 3
    assert summary["mto_rows"] == 27
    workbook = load_workbook(summary["workbook"], data_only=True)
    combined = workbook["COMBINED"]
    descriptions = [combined[f"F{row}"].value for row in range(2, 29)]
    assert sum(value.startswith("ELL 90") for value in descriptions) == 5
    assert sum(value.startswith("NIPPLE") for value in descriptions) == 3
    assert combined.max_row == 28
    workbook.close()
