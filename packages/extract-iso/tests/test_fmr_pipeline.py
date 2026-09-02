import base64
import csv
import json
from pathlib import Path

import fitz
import pytest
from openpyxl.cell.rich_text import CellRichText
from openpyxl import load_workbook

from iso_bom.fmr_cli import build_parser
from iso_bom.fmr_model import IsoPage
from iso_bom.fmr_parser import (
    extract_iwp_number,
    is_fmr_attachment_page,
    is_iso_workflow_status_page,
    is_iso_page,
    is_iwp_cover_page,
    is_logistics_page,
    is_pipe_support_page,
    is_pipe_material,
    is_weld_log_page,
    parse_iso_page,
)
from iso_bom.fmr_pipeline import _select_revisions, run_fmr
from iso_bom.fmr_workbook import (
    MATERIAL_CAPACITY,
    MATERIAL_LAST_ROW,
    build_fmr_workbook,
    sheet_names,
)
from iso_bom.model import BomRow
from iso_bom.pdf_parser import parse_page


PROJECT_ROOT = Path(__file__).resolve().parents[1]
FMR_TEMPLATE = PROJECT_ROOT / "templates" / "FMR" / "blankFMR.xlsx"
SUPPLIED_PACKAGE = PROJECT_ROOT / "templates" / "CreateFMR" / "IP-SMM30R127MMPP-K447.pdf"
NEW_FMR116_PACKAGE = PROJECT_ROOT / "input" / "newFmr116" / "IP-SMM30B0012FPP-K447-116.pdf"
ONE_PIXEL_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def _insert_cover(document: fitz.Document, iwp_number: str) -> None:
    page = document.new_page(width=612, height=792)
    page.insert_text((170, 35), "Installation Work Package - TEST", fontsize=12)
    page.insert_text((50, 75), "IWP Number:", fontsize=10)
    page.insert_text((130, 75), iwp_number, fontsize=10)


def _insert_iso(
    document: fitz.Document,
    drawing_number: str,
    revision: str,
    rows,
) -> None:
    page = document.new_page(width=900, height=700)
    page.insert_text((420, 35), "BILL OF MATERIALS", fontsize=9)
    for x, text in [
        (420, "PT"), (455, "DESCRIPTION"), (650, "NPD"),
        (710, "CMDTY"), (830, "QTY"),
    ]:
        page.insert_text((x, 60), text, fontsize=8)
    y = 90
    for row in rows:
        for x, text in zip((420, 455, 650, 710, 830), row):
            if text:
                page.insert_text((x, y), text, fontsize=8)
        y += 28
    page.insert_text((500, 610), "ISOMETRIC DRAWING NUMBER", fontsize=8)
    page.insert_text((790, 610), "REV", fontsize=8)
    page.insert_text((500, 635), drawing_number, fontsize=8)
    page.insert_text((795, 635), revision, fontsize=8)


def _insert_weld_log(document: fitz.Document) -> None:
    page = document.new_page(width=612, height=792)
    page.insert_text((50, 50), "SECTION 15117", fontsize=10)
    page.insert_text((50, 70), "FABRICATION OF METALLIC PIPE AND TUBING", fontsize=10)
    page.insert_text((50, 90), "ATTACHMENT A - WELD/BRAZE LOG SHEET", fontsize=10)


def _insert_view_attachment(document: fitz.Document, label: str = "LOOKING EAST") -> None:
    page = document.new_page(width=500, height=300)
    page.insert_image(page.rect, stream=ONE_PIXEL_PNG)
    page.insert_text((340, 120), label, fontsize=8)


def _create_test_package(path: Path) -> None:
    document = fitz.open()
    _insert_cover(document, "IWP-TEST-100")
    _insert_iso(document, "ISO-001", "0", [
        ("1", "PIPE SCH 10S", "2", "1111111", "10"),
        ("2", "ELL 90 LR", "2", "OLD-001", "1"),
    ])
    _insert_iso(document, "ISO-001", "2", [
        ("1", "PIPE SCH 10S", "2", "1111111", "10"),
        ("2", "ELL 90 LR", "2", "NEW-001", "2"),
    ])
    _insert_iso(document, "ISO-002", "0", [
        ("1", "PIPE SCH 10S", "6", "2222222", "12"),
        ("2", "PIPET SCH 10S 316 SS", "6X2", "FIT-002", "1"),
    ])
    _insert_weld_log(document)
    document.save(path)
    document.close()


def _row(point="1", description="ELL 90", code="CODE-1", quantity="1") -> BomRow:
    return BomRow(point, description, "2", code, quantity, description, (0, 0, 1, 1))


def _iso(page: int, revision: str, content_hash: str, materials=None) -> IsoPage:
    return IsoPage(
        source_pdf="package.pdf",
        source_path="package.pdf",
        page=page,
        drawing_number="ISO-001",
        revision=revision,
        materials=materials or [_row()],
        raw_text="source text",
        content_hash=content_hash,
    )


def test_pipe_filter_excludes_pipe_word_but_keeps_pipet():
    assert is_pipe_material(_row(description="PIPE SCH 10S ERW"))
    assert is_pipe_material(_row(description=" pipe SCH 40"))
    assert not is_pipe_material(_row(description="PIPET 10S X 3000#"))


def test_standalone_bom_point_annotation_is_ignored():
    document = fitz.open()
    _insert_iso(document, "ISO-ANNOTATED", "0", [
        ("1", "PIPE SCH 10S", "4", "P-001", "10.0'"),
        ("2", "ELL 90 LR", "4", "E-001", "1"),
        ("1", "", "", "", ""),
        ("3", "VALVE BALL", "4", "V-001", "1"),
    ])
    iso = parse_iso_page(document[0], 1, "package.pdf", "package.pdf")
    assert iso.review_reasons == []
    assert [(row.point_number, row.description) for row in iso.materials] == [
        ("1", "PIPE SCH 10S"),
        ("2", "ELL 90 LR"),
        ("3", "VALVE BALL"),
    ]
    document.close()


def test_bottom_bom_notes_do_not_pollute_material_cells():
    document = fitz.open()
    page = document.new_page(width=900, height=700)
    page.insert_text((420, 35), "BILL OF MATERIALS", fontsize=9)
    for x, text in [
        (420, "PT"), (455, "DESCRIPTION"), (650, "NPD"),
        (710, "CMDTY"), (830, "QTY"),
    ]:
        page.insert_text((x, 60), text, fontsize=8)
    for x, text in [
        (420, "25"),
        (455, "5UG, U-BOLT GUIDE FOR UNINSULATED LINES"),
        (650, "2"),
        (710, "5UG-02"),
        (830, "1"),
    ]:
        page.insert_text((x, 90), text, fontsize=8)
    page.insert_text((455, 105), '2" PIPE', fontsize=8)
    page.insert_text((650, 128), "FIELD WELDS TO BE", fontsize=8)
    page.insert_text((650, 151), "CUT AND BEVELED", fontsize=8)
    page.insert_text((710, 170), "112002-02-1503", fontsize=8)

    parsed = parse_page(page, 1)
    assert parsed.audit["status"] == "ok"
    assert len(parsed.rows) == 1
    row = parsed.rows[0]
    assert row.description == '5UG, U-BOLT GUIDE FOR UNINSULATED LINES 2" PIPE'
    assert row.nominal_size == "2"
    assert row.commodity_code == "5UG-02"
    assert row.quantity == "1"
    assert row.structural_notes == []
    document.close()


def test_close_first_bom_row_and_revision_history_are_handled():
    document = fitz.open()
    page = document.new_page(width=900, height=700)
    for x, text in [
        (420, "PT"), (455, "DESCRIPTION"), (650, "NPD"),
        (710, "CMDTY"), (830, "QTY"),
    ]:
        page.insert_text((x, 60), text, fontsize=8)
    for x, text in zip(
        (420, 430, 650, 710, 830),
        ("1", "PIPE SCH 80 ERW STL A53-B", "1 1/2", "5355898L", "2.1'"),
    ):
        page.insert_text((x, 73), text, fontsize=8)
    for x, text in zip(
        (420, 430, 650, 710, 830),
        ("2", "TEE 3000# SW STL A105", "1 1/2X1 1/2", "5478107", "1"),
    ):
        page.insert_text((x, 90), text, fontsize=8)
    page.insert_text((400, 90), "NORTH", fontsize=8)
    for x, text in zip(
        (420, 455, 710, 770),
        ("1", "ISSUED FOR CONSTRUCTION", "07-APR-2026", "RNR"),
    ):
        page.insert_text((x, 200), text, fontsize=8)
    for x, text in zip(
        (420, 455, 710, 770),
        ("0", "ISSUED FOR CONSTRUCTION", "11-SEP-2025", "RJC"),
    ):
        page.insert_text((x, 215), text, fontsize=8)

    parsed = parse_page(page, 1)
    assert parsed.audit["status"] == "ok"
    assert [(row.point_number, row.description) for row in parsed.rows] == [
        ("1", "PIPE SCH 80 ERW STL A53-B"),
        ("2", "TEE 3000# SW STL A105"),
    ]
    assert parsed.rows[0].commodity_code == "5355898L"
    assert parsed.rows[1].nominal_size == "1 1/2X1 1/2"
    assert parsed.rows[1].commodity_code == "5478107"
    document.close()


def test_numeric_wrap_artifact_does_not_create_false_bom_point():
    document = fitz.open()
    page = document.new_page(width=900, height=700)
    for x, text in [
        (420, "PT"), (455, "DESCRIPTION"), (650, "NPD"),
        (710, "CMDTY"), (830, "QTY"),
    ]:
        page.insert_text((x, 60), text, fontsize=8)
    rows = [
        ("1", "PIPE SCH 40 ERW STL A53-B", "2", "5356648", "6.6'"),
        ("2", "ELL 90 DEG LR SCH 40 STL A234 WPB", "2", "5374060", "5"),
        ("3", "FLG WN 300# RF STL 40 BORE A105", "2", "5554634", "2"),
        ("4", "GASKET 300# FILLED TFE RING 1/16\" THK", "2", "5669391", "2"),
        ("5", "STUD-BOLT A193 GR B7 W/A194 GR 2H NUTS -", "5/8", "5675648", "16"),
    ]
    y = 74
    for point, description, size, code, quantity in rows:
        for x, text in zip((420, 455, 650, 710, 830), (point, description, size, code, quantity)):
            page.insert_text((x, y), text, fontsize=8)
        y += 18
    page.insert_text((420, y - 6), "1", fontsize=8)
    page.insert_text((650, y - 8), "1", fontsize=8)
    page.insert_text((455, y - 2), "3.75 in. Length", fontsize=8)
    for x, text in zip(
        (420, 455, 650, 710, 830),
        ("6", "BALL 300# RF STL, 316 TRIM, RP, EXT BON,", "2", "5301238L8", "1"),
    ):
        page.insert_text((x, y + 16), text, fontsize=8)
    page.insert_text((455, y + 31), "HNDL OP, FS", fontsize=8)

    parsed = parse_page(page, 1)
    assert parsed.audit["status"] == "ok"
    assert [row.point_number for row in parsed.rows] == ["1", "2", "3", "4", "5", "6"]
    assert parsed.rows[4].description == (
        "STUD-BOLT A193 GR B7 W/A194 GR 2H NUTS - 3.75 in. Length"
    )
    assert parsed.rows[4].nominal_size == "5/8"
    assert parsed.rows[4].commodity_code == "5675648"
    assert parsed.rows[4].quantity == "16"
    assert parsed.rows[5].description == (
        "BALL 300# RF STL, 316 TRIM, RP, EXT BON, HNDL OP, FS"
    )
    document.close()


def test_inline_numeric_wrap_artifact_is_removed_from_wrapped_description():
    document = fitz.open()
    page = document.new_page(width=900, height=700)
    for x, text in [
        (420, "PT"), (455, "DESCRIPTION"), (650, "NPD"),
        (710, "CMDTY"), (830, "QTY"),
    ]:
        page.insert_text((x, 60), text, fontsize=8)
    for x, text in zip(
        (420, 455, 650, 710, 830),
        ("5", "STUD-BOLT A193 GR B7 W/A194 GR 2H NUTS -", "5/8", "5675648", "16"),
    ):
        page.insert_text((x, 90), text, fontsize=8)
    page.insert_text((438, 103), "1", fontsize=8)
    page.insert_text((455, 103), "3.75 in. Length", fontsize=8)
    page.insert_text((650, 101), "1", fontsize=8)
    for x, text in zip(
        (420, 455, 650, 710, 830),
        ("6", "BALL 300# RF STL, 316 TRIM, RP, EXT BON,", "2", "5301238L8", "1"),
    ):
        page.insert_text((x, 120), text, fontsize=8)

    parsed = parse_page(page, 1)
    assert [row.point_number for row in parsed.rows] == ["5", "6"]
    assert parsed.rows[0].description == (
        "STUD-BOLT A193 GR B7 W/A194 GR 2H NUTS - 3.75 in. Length"
    )
    assert parsed.rows[0].nominal_size == "5/8"
    assert parsed.rows[0].commodity_code == "5675648"
    assert parsed.rows[0].quantity == "16"
    document.close()


def test_title_block_metadata_after_final_support_row_is_not_commodity():
    document = fitz.open()
    page = document.new_page(width=900, height=700)
    for x, text in [
        (420, "PT"), (455, "DESCRIPTION"), (650, "NPD"),
        (710, "CMDTY"), (830, "QTY"),
    ]:
        page.insert_text((x, 60), text, fontsize=8)
    for x, text in zip(
        (420, 455, 650, 710, 830),
        ("27", '5CI, ISOLATION CRADLE, 2" PIPE, SS', "2", "5CI-02", "1"),
    ):
        page.insert_text((x, 90), text, fontsize=8)
    page.insert_text((660, 120), "112002-30287", fontsize=8)
    page.insert_text((660, 142), "MOD-SMM30R109MM", fontsize=8)
    page.insert_text((670, 164), "SPEC 315", fontsize=8)

    parsed = parse_page(page, 1)
    assert parsed.audit["status"] == "ok"
    assert len(parsed.rows) == 1
    row = parsed.rows[0]
    assert row.point_number == "27"
    assert row.description == '5CI, ISOLATION CRADLE, 2" PIPE, SS'
    assert row.nominal_size == "2"
    assert row.commodity_code == "5CI-02"
    assert row.quantity == "1"
    assert row.structural_notes == []
    document.close()


def test_standard_threaded_joint_notes_do_not_pollute_final_material_row():
    document = fitz.open()
    page = document.new_page(width=900, height=700)
    for x, text in [
        (420, "PT"), (455, "DESCRIPTION"), (650, "NPD"),
        (710, "CMDTY"), (830, "QTY"),
    ]:
        page.insert_text((x, 60), text, fontsize=8)
    for x, text in zip(
        (420, 455, 650, 710, 830),
        (
            "7",
            "5UGSP, U-BOLT GUIDE FOR INSULATED LINES",
            "3/4",
            "5UGSP-75-01",
            "1",
        ),
    ):
        page.insert_text((x, 90), text, fontsize=8)
    page.insert_text((455, 108), '3/4" PIPE W/ 1" INSUL', fontsize=8)
    page.insert_text((455, 126), "THREADED JOINTS ON THIS", fontsize=8)
    page.insert_text((650, 126), "ISO ARE", fontsize=8)
    page.insert_text((710, 126), "NOW SW.", fontsize=8)
    page.insert_text((455, 144), "VALVES AND INSTRUMENTS", fontsize=8)
    page.insert_text((650, 144), "TO REMAIN", fontsize=8)
    page.insert_text((710, 144), "THREADED.", fontsize=8)

    parsed = parse_page(page, 1)
    assert parsed.audit["status"] == "ok"
    assert len(parsed.rows) == 1
    row = parsed.rows[0]
    assert row.description == (
        '5UGSP, U-BOLT GUIDE FOR INSULATED LINES 3/4" PIPE W/ 1" INSUL'
    )
    assert row.nominal_size == "3/4"
    assert row.commodity_code == "5UGSP-75-01"
    assert row.quantity == "1"
    assert row.structural_notes == []
    document.close()


def test_ocr_header_aliases_are_accepted_for_standard_iso_bom():
    document = fitz.open()
    page = document.new_page(width=900, height=700)
    page.insert_text((420, 35), "BILL OF MATERIALS", fontsize=9)
    for x, text in [
        (420, "PT"), (455, "DESCRIPTION"), (650, "NPD"),
        (710, "CMD1Y"), (760, "CODE"), (830, "Q1Y"),
    ]:
        page.insert_text((x, 60), text, fontsize=8)
    for x, text in zip(
        (420, 455, 650, 710, 830),
        ("1", "ELL 90 LR SCH 40 STL A234 WPB", "2", "5450706", "9"),
    ):
        page.insert_text((x, 90), text, fontsize=8)

    parsed = parse_page(page, 1)
    assert parsed.audit["status"] == "ok"
    assert len(parsed.rows) == 1
    row = parsed.rows[0]
    assert row.point_number == "1"
    assert row.description == "ELL 90 LR SCH 40 STL A234 WPB"
    assert row.nominal_size == "2"
    assert row.commodity_code == "5450706"
    assert row.quantity == "9"
    document.close()


def test_support_ocr_s_prefix_is_normalized_to_standard_5_code():
    document = fitz.open()
    page = document.new_page(width=900, height=700)
    page.insert_text((420, 35), "BILL OF MATERIALS", fontsize=9)
    for x, text in [
        (420, "PT"), (455, "DESCRIPTION"), (650, "NPD"),
        (710, "CMDTY"), (830, "QTY"),
    ]:
        page.insert_text((x, 60), text, fontsize=8)
    rows = [
        ("12", "SUPPORT CRADLE COLD SERVICE", "2", "SCC-02-50", "5"),
        ("13", 'SUGSP, U-BOLT GUIDE FOR INSULATED LINES 2" PIPE', "2", "SUGSP-02-50", "4"),
        ("14", "SFSS4, FIELD SUPPORT STEEL, SIZE L, LC 3", "3/4", "SFSS4-L-3", "1"),
        ("15", "SMHR2-5-L-2", "3/4", "SMHR2-5-L", "1"),
        ("16", 'SC, INSULATION CRADLE 3/4" PIPE, 1/2" INS', "3/4", "SC-75-50", "1"),
    ]
    y = 90
    for row in rows:
        for x, text in zip((420, 455, 650, 710, 830), row):
            page.insert_text((x, y), text, fontsize=8)
        y += 24

    parsed = parse_page(page, 1)
    assert parsed.audit["status"] == "ok"
    assert [
        (row.description, row.commodity_code)
        for row in parsed.rows
    ] == [
        ("5CC, SUPPORT CRADLE COLD SERVICE", "5CC-02-50"),
        ('5UGSP, U-BOLT GUIDE FOR INSULATED LINES 2" PIPE', "5UGSP-02-50"),
        ("5FSS4, FIELD SUPPORT STEEL, SIZE L, LC 3", "5FSS4-L-3"),
        ("5MHR2-S-L-2", "5MHR2-S-L"),
        ('5C, INSULATION CRADLE 3/4" PIPE, 1/2" INS', "5C-75-50"),
    ]
    document.close()


def test_standard_stainless_description_ocr_is_normalized():
    document = fitz.open()
    page = document.new_page(width=900, height=700)
    page.insert_text((420, 35), "BILL OF MATERIALS", fontsize=9)
    for x, text in [
        (420, "PT"), (455, "DESCRIPTION"), (650, "NPD"),
        (710, "CMDTY"), (830, "QTY"),
    ]:
        page.insert_text((x, 60), text, fontsize=8)
    rows = [
        ("1", "PIPE SCH 105 ERW 316/316L 55 A312", "2", "5368751", "1.0'"),
        ("2", "PIPET 105 X 3000# SW 316/316L 55", "2X3/4", "5537081", "1"),
        ("3", "FLG SW 150# RF 316/316L 55 105 BORE", "3/4", "5608148", "1"),
        ("4", "BALL 150# RF 31655 TFE FP EXTENDED STEM", "3/4", "5352193L2", "1"),
        ("5", "FLG BLIND 150# RF 316/316L 55 sec,", "3/4", "5616520", "1"),
    ]
    y = 90
    for row in rows:
        for x, text in zip((420, 455, 650, 710, 830), row):
            page.insert_text((x, y), text, fontsize=8)
        y += 24

    parsed = parse_page(page, 1)
    assert parsed.audit["status"] == "ok"
    assert [row.description for row in parsed.rows] == [
        "PIPE SCH 10S ERW 316/316L SS A312",
        "PIPET 10S X 3000# SW 316/316L SS",
        "FLG SW 150# RF 316/316L SS 10S BORE",
        "BALL 150# RF 316SS TFE FP EXTENDED STEM",
        "FLG BLIND 150# RF 316/316L SS",
    ]
    document.close()


def test_description_fragments_are_not_left_in_size_column():
    document = fitz.open()
    page = document.new_page(width=900, height=700)
    for x, text in [
        (420, "PT"), (455, "DESCRIPTION"), (650, "NPD"),
        (710, "CMDTY"), (830, "QTY"),
    ]:
        page.insert_text((x, 60), text, fontsize=8)
    rows = [
        (
            "1", ["STRAINER Y-TYPE 600# STL FNPT 304 .045 PERF"],
            "SCR 1 1/2X3/4", "LP131 STR-6065U2", "1",
        ),
        (
            "2", ["STEAM TRAP FLOAT HZ 465 PSIG STL"],
            "FNPT 1 1/2", "LP131 STT-6065U1", "1",
        ),
        (
            "3", ['5ABS3, T-SHAPED BASE SUPPORT, FOR', 'SIZE 6" NPD AND BELOW'],
            "PIPE 1 1/2", "5ABS3", "1",
        ),
        (
            "4", ['5MS, CLAMP SHOE, TYPE 1, 5" HIGH,', "NPD"],
            '1-1/2" 1 1/2', "5MS15-15", "1",
        ),
        (
            "5", ["5FS1, FIELD SUPPORT STEEL ATTACHMENT,", "AND SMALLER PIPE"],
            '2" 1 1/2', "5FS1S1", "1",
        ),
    ]
    y = 74
    for point, description_lines, size, code, quantity in rows:
        page.insert_text((420, y), point, fontsize=8)
        for offset, text in enumerate(description_lines):
            page.insert_text((455, y + (offset * 5)), text, fontsize=8)
        for x, text in ((650, size), (710, code), (830, quantity)):
            page.insert_text((x, y), text, fontsize=8)
        y += 20

    parsed = parse_page(page, 1)
    assert [(row.commodity_code, row.nominal_size, row.description) for row in parsed.rows] == [
        (
            "LP131 STR-6065U2",
            "1 1/2X3/4",
            "STRAINER Y-TYPE 600# STL FNPT 304 .045 PERF",
        ),
        (
            "LP131 STT-6065U1",
            "1 1/2",
            "STEAM TRAP FLOAT HZ 465 PSIG STL FNPT",
        ),
        (
            "5ABS3",
            "1 1/2",
            '5ABS3, T-SHAPED BASE SUPPORT, FOR PIPE SIZE 6" NPD AND BELOW',
        ),
        (
            "5MS15-15",
            "1 1/2",
            '5MS, CLAMP SHOE, TYPE 1, 5" HIGH, 1-1/2" NPD',
        ),
        (
            "5FS1S1",
            "1 1/2",
            '5FS1, FIELD SUPPORT STEEL ATTACHMENT, 2" AND SMALLER PIPE',
        ),
    ]
    document.close()


def test_cover_extraction_and_weld_log_classification():
    document = fitz.open()
    _insert_cover(document, "IWP-TEST-200")
    bare_cover = document.new_page(width=612, height=792)
    bare_cover.insert_text((50, 50), "IP-SMM10R125MMPP-K447-101 Issued to Construction", fontsize=10)
    bare_cover.insert_text((50, 75), "IP-SMM10R125MMPP-K447-101 Issued to Construction", fontsize=10)
    for_construction_cover = document.new_page(width=612, height=792)
    for_construction_cover.insert_text(
        (50, 50),
        "IP-SMM00U0011FPP-K447-105 Install AI/APL piping Issued for Construction",
        fontsize=10,
    )
    for_construction_cover.insert_text((50, 75), "Please see below JSA's for this IWP:", fontsize=10)
    singular_jsa_cover = document.new_page(width=612, height=792)
    singular_jsa_cover.insert_text(
        (50, 50),
        "Please see below for the following JSA for this IWP",
        fontsize=10,
    )
    singular_jsa_cover.insert_text(
        (50, 75),
        "IP-SMM20E0011FPP-K447-102 IP-SMM20E011FPP Install First Fix Piping 20E",
        fontsize=10,
    )
    base_iwp_cover = document.new_page(width=612, height=792)
    base_iwp_cover.insert_text(
        (50, 50),
        "IP-SMM30R117MMPP-K447 Issued to Construction",
        fontsize=10,
    )
    _insert_weld_log(document)
    iwp_number, reasons = extract_iwp_number(document[0])
    assert (iwp_number, reasons) == ("IWP-TEST-200", [])
    iwp_number, reasons = extract_iwp_number(document[1])
    assert is_iwp_cover_page(document[1].get_text("text"))
    assert (iwp_number, reasons) == ("IP-SMM10R125MMPP-K447-101", [])
    iwp_number, reasons = extract_iwp_number(document[2])
    assert is_iwp_cover_page(document[2].get_text("text"))
    assert (iwp_number, reasons) == ("IP-SMM00U0011FPP-K447-105", [])
    iwp_number, reasons = extract_iwp_number(document[3])
    assert is_iwp_cover_page(document[3].get_text("text"))
    assert (iwp_number, reasons) == ("IP-SMM20E0011FPP-K447-102", [])
    iwp_number, reasons = extract_iwp_number(document[4])
    assert is_iwp_cover_page(document[4].get_text("text"))
    assert (iwp_number, reasons) == ("IP-SMM30R117MMPP-K447", [])
    assert is_weld_log_page(document[5].get_text("text"))
    document.close()


def test_iso_construction_status_is_not_mistaken_for_cover():
    document = fitz.open()
    _insert_iso(document, "ISO-STATUS", "0", [
        ("1", "PIPE SCH 10S", "2", "P-001", "5.0'"),
    ])
    document[0].insert_text(
        (40, 660),
        "IP-SMM30R117MMPP-K447 Issued for Construction",
        fontsize=8,
    )
    text = document[0].get_text("text")
    assert is_iso_page(text)
    assert not is_iwp_cover_page(text)
    document.close()


def test_known_non_iso_attachments_are_classified():
    assert is_pipe_support_page(
        "PIPE HANGERS AND SUPPORTS - ATTACHMENT C-1"
    )
    assert is_fmr_attachment_page("FIELD MATERIAL REQUEST / RETURN (FMR)")
    assert is_logistics_page("LP1 LOGISTICS PLAN")
    assert is_iso_workflow_status_page(
        "Isos In Processing Iso Rev WorkflowStatus LP131-AI(100)-852045-03 1 4-Fab Shop"
    )


def test_revision_selection_and_conflict_quarantine():
    accepted, quarantine, superseded, duplicate, unique = _select_revisions(
        [_iso(1, "0", "old"), _iso(2, "2", "new")], "IWP-1"
    )
    assert [item.revision for item in accepted] == ["2"]
    assert not quarantine
    assert (superseded, duplicate, unique) == (1, 0, 1)

    accepted, quarantine, _, _, _ = _select_revisions(
        [_iso(1, "A", "a"), _iso(2, "B", "b")], "IWP-1"
    )
    assert not accepted
    assert {entry.reason_code for entry in quarantine} == {"unorderable_multiple_revisions"}

    accepted, quarantine, _, duplicate, _ = _select_revisions(
        [_iso(1, "2", "same"), _iso(2, "2", "same")], "IWP-1"
    )
    assert len(accepted) == 1 and not quarantine and duplicate == 1

    accepted, quarantine, _, _, _ = _select_revisions(
        [_iso(1, "2", "one"), _iso(2, "2", "two")], "IWP-1"
    )
    assert not accepted
    assert {entry.reason_code for entry in quarantine} == {"duplicate_revision_conflict"}


def test_material_capacity_and_sheet_name_validation():
    materials = [_row(str(index), code=f"C-{index}") for index in range(MATERIAL_CAPACITY + 1)]
    accepted, quarantine, _, _, _ = _select_revisions(
        [_iso(1, "0", "overflow", materials)], "IWP-1"
    )
    assert len(accepted) == 1
    assert not quarantine
    with pytest.raises(ValueError, match="31 characters"):
        sheet_names("IWP-" + "X" * 30, 1)


def test_fmr_workbook_extends_material_rows_before_footer(tmp_path):
    material_count = MATERIAL_CAPACITY + 2
    materials = [
        _row(str(index), code=f"C-{index}")
        for index in range(1, material_count + 1)
    ]
    output = tmp_path / "extended.xlsx"

    build_fmr_workbook(
        FMR_TEMPLATE,
        output,
        "IWP-EXT",
        [_iso(1, "0", "overflow", materials)],
    )

    workbook = load_workbook(output, rich_text=True)
    worksheet = workbook["IWP-EXT(00)"]
    assert worksheet[f"B{MATERIAL_LAST_ROW + 1}"].value == "C-24"
    assert worksheet[f"B{MATERIAL_LAST_ROW + 2}"].value == "C-25"
    assert worksheet[f"B{MATERIAL_LAST_ROW + 3}"].value == "REASON REQUIRED"
    merged_ranges = {str(item) for item in worksheet.merged_cells.ranges}
    assert f"B{MATERIAL_LAST_ROW + 1}:E{MATERIAL_LAST_ROW + 1}" not in merged_ranges
    assert f"B{MATERIAL_LAST_ROW + 3}:E{MATERIAL_LAST_ROW + 3}" in merged_ranges
    assert f"E{MATERIAL_LAST_ROW + 2}:H{MATERIAL_LAST_ROW + 2}" in merged_ranges
    assert min(image.anchor._from.row + 1 for image in worksheet._images) == 33


def test_fmr_pipeline_creates_formatted_workbook(tmp_path):
    source = tmp_path / "input"
    output = tmp_path / "output"
    source.mkdir()
    _create_test_package(source / "package.pdf")

    summary = run_fmr(source, output, FMR_TEMPLATE)
    assert summary["status"] == "complete"
    assert summary["iso_pages_detected"] == 3
    assert summary["unique_drawings"] == 2
    assert summary["superseded_pages_ignored"] == 1
    assert summary["weld_log_pages_ignored"] == 1
    assert summary["fmr_sheets_created"] == 2
    assert summary["quarantine_entries"] == 0
    assert summary["pipe_rows_included"] is True
    assert summary["pipe_rows_only"] is False
    assert summary["estimated_time_saved_seconds"] == "164.02"
    assert summary["time_estimate_available"] is True
    assert float(summary["processing_seconds"]) > 0

    manifests = list((output / "analytics_runs").glob("*.json"))
    assert len(manifests) == 1
    manifest = json.loads(manifests[0].read_text())
    assert manifest["workflow"] == "fmr"
    assert manifest["status"] == "complete"
    assert manifest["package_number"] == "IWP-TEST-100"
    assert [item["drawing_number"] for item in manifest["isos"]] == [
        "ISO-001", "ISO-002",
    ]
    assert manifest["metrics"]["fmr_sheets_created"] == 2
    assert manifest["metrics"]["estimated_time_saved_seconds"] == "164.02"
    assert [item["overflow_material_rows"] for item in manifest["isos"]] == [0, 0]

    workbook_path = Path(summary["workbook"])
    workbook = load_workbook(workbook_path, rich_text=True)
    assert workbook.sheetnames == ["IWP-TEST-100(00)", "IWP-TEST-100(01)"]
    first = workbook[workbook.sheetnames[0]]
    second = workbook[workbook.sheetnames[1]]
    assert str(first["B4"].value) == "DESTINATION:  \nField"
    assert isinstance(first["B5"].value, CellRichText)
    assert str(first["B5"].value) == "REQUESTED BY:\n"
    assert first["B6"].value == "DELIVER TO:\n"
    assert first["H5"].value == "IWP: IWP-TEST-100"
    assert first["H6"].value == "LINE NO:\nISO-001"
    assert first["K6"].value == "REV:\n2"
    assert (first["B8"].value, first["C8"].value, first["D8"].value) == ("1111111", "2", 10)
    assert first["B9"].value == "NEW-001"
    assert second["B8"].value == "2222222"
    assert second["E9"].value.startswith("PIPET")
    assert first["E8"].alignment.shrink_to_fit
    assert all(
        f"E{row}:H{row}" in {str(item) for item in first.merged_cells.ranges}
        for row in range(8, 31)
    )
    assert len(first._images) == 7
    assert all(len(worksheet._images) == 7 for worksheet in workbook.worksheets)
    assert first.page_setup.orientation == "portrait"

    all_values = "\n".join(
        str(cell.value or "")
        for worksheet in workbook.worksheets
        for row in worksheet.iter_rows(min_row=1, max_row=40, min_col=1, max_col=11)
        for cell in row
    )
    assert "Manuel Garcia Jr" not in all_values
    assert "Cedric Labassiere" not in all_values
    with (output / "fmr_quarantine.csv").open() as handle:
        assert list(csv.DictReader(handle)) == []


def test_fmr_cli_assignment_fields_populate_every_sheet(tmp_path):
    source = tmp_path / "input"
    output = tmp_path / "output"
    source.mkdir()
    _create_test_package(source / "package.pdf")

    args = build_parser().parse_args([
        "--input", str(source),
        "--output", str(output),
        "--destination", "3rd Floor",
        "--requested-by", "Manuel Garcia",
        "--deliver-to", "Cedric",
    ])
    summary = run_fmr(
        args.input,
        args.output,
        FMR_TEMPLATE,
        destination=args.destination,
        requested_by=args.requested_by,
        deliver_to=args.deliver_to,
    )

    workbook = load_workbook(summary["workbook"], rich_text=True)
    assert len(workbook.worksheets) == 2
    for worksheet in workbook.worksheets:
        assert str(worksheet["B4"].value) == "DESTINATION:\n3rd Floor"
        assert str(worksheet["B5"].value) == "REQUESTED BY:\nManuel Garcia"
        assert worksheet["B6"].value == "DELIVER TO:\nCedric"


def test_fmr_cli_pipe_flag_includes_pipe_stock_rows(tmp_path):
    source = tmp_path / "input"
    output = tmp_path / "output"
    source.mkdir()
    _create_test_package(source / "package.pdf")

    args = build_parser().parse_args([
        "--input", str(source),
        "--output", str(output),
        "--pipe",
    ])
    summary = run_fmr(
        args.input,
        args.output,
        FMR_TEMPLATE,
        include_pipe=args.pipe,
    )

    assert summary["status"] == "complete"
    assert summary["pipe_rows_included"] is True
    assert summary["pipe_rows_only"] is False

    workbook = load_workbook(summary["workbook"])
    first = workbook["IWP-TEST-100(00)"]
    second = workbook["IWP-TEST-100(01)"]
    assert (first["B8"].value, first["E8"].value) == ("1111111", "PIPE SCH 10S")
    assert (first["B9"].value, first["E9"].value) == ("NEW-001", "ELL 90 LR")
    assert (second["B8"].value, second["E8"].value) == ("2222222", "PIPE SCH 10S")
    assert (second["B9"].value, second["E9"].value) == ("FIT-002", "PIPET SCH 10S 316 SS")

    manifests = list((output / "analytics_runs").glob("*.json"))
    assert len(manifests) == 1
    manifest = json.loads(manifests[0].read_text())
    assert [item["pipe_rows"] for item in manifest["isos"]] == [1, 1]


def test_fmr_cli_pipe_only_flag_keeps_only_pipe_stock_rows(tmp_path):
    source = tmp_path / "input"
    output = tmp_path / "output"
    source.mkdir()
    _create_test_package(source / "package.pdf")

    args = build_parser().parse_args([
        "--input", str(source),
        "--output", str(output),
        "--pipe-only",
    ])
    summary = run_fmr(
        args.input,
        args.output,
        FMR_TEMPLATE,
        pipe_only=args.pipe_only,
    )

    assert summary["status"] == "complete"
    assert summary["pipe_rows_included"] is True
    assert summary["pipe_rows_only"] is True

    workbook = load_workbook(summary["workbook"])
    first = workbook["IWP-TEST-100(00)"]
    second = workbook["IWP-TEST-100(01)"]
    assert (first["B8"].value, first["E8"].value) == ("1111111", "PIPE SCH 10S")
    assert not first["B9"].value
    assert not first["E9"].value
    assert (second["B8"].value, second["E8"].value) == ("2222222", "PIPE SCH 10S")
    assert not second["B9"].value
    assert not second["E9"].value

    manifests = list((output / "analytics_runs").glob("*.json"))
    assert len(manifests) == 1
    manifest = json.loads(manifests[0].read_text())
    assert [item["pipe_rows"] for item in manifest["isos"]] == [1, 1]
    assert [item["material_rows"] for item in manifest["isos"]] == [1, 1]


def test_fmr_cli_pipe_modes_are_mutually_exclusive():
    with pytest.raises(SystemExit):
        build_parser().parse_args([
            "--input", "input",
            "--output", "output",
            "--pipe",
            "--pipe-only",
        ])


def test_scanned_and_partial_iso_pages_are_quarantined(tmp_path):
    source = tmp_path / "input"
    output = tmp_path / "output"
    source.mkdir()
    _create_test_package(source / "package.pdf")

    extras = fitz.open()
    extras.new_page(width=612, height=792)
    partial = extras.new_page(width=612, height=792)
    partial.insert_text(
        (40, 50),
        "BILL OF MATERIALS PRESENT BUT THE REQUIRED ISO TITLE BLOCK IS MISSING",
        fontsize=10,
    )
    extras.save(source / "extras.pdf")
    extras.close()

    summary = run_fmr(source, output, FMR_TEMPLATE)
    assert summary["fmr_sheets_created"] == 2
    assert summary["quarantine_entries"] == 2
    with (output / "fmr_quarantine.csv").open() as handle:
        entries = list(csv.DictReader(handle))
    assert {entry["reason_code"] for entry in entries} == {
        "ocr_required", "partial_iso_structure",
    }
    assert all(entry["raw_text"] or entry["reason_code"] == "ocr_required" for entry in entries)


def test_post_iso_looking_view_attachments_are_ignored(tmp_path):
    source = tmp_path / "input"
    output = tmp_path / "output"
    source.mkdir()
    document = fitz.open()
    _insert_cover(document, "IWP-VIEW-RUN")
    _insert_iso(document, "ISO-001", "0", [
        ("1", "ELL 90 LR", "2", "C-1", "1"),
    ])
    _insert_view_attachment(document, "LOOKING EAST")
    _insert_view_attachment(document, "LOOKING NORTH")
    _insert_weld_log(document)
    document.save(source / "package.pdf")
    document.close()

    summary = run_fmr(source, output, FMR_TEMPLATE)
    assert summary["status"] == "complete"
    assert summary["fmr_sheets_created"] == 1
    assert summary["attachment_image_pages_ignored"] == 2
    assert summary["weld_log_pages_ignored"] == 1
    assert summary["quarantine_entries"] == 0


def test_post_iso_side_view_attachments_are_ignored(tmp_path):
    source = tmp_path / "input"
    output = tmp_path / "output"
    source.mkdir()
    document = fitz.open()
    _insert_cover(document, "IWP-SIDE-VIEW")
    _insert_iso(document, "ISO-001", "0", [
        ("1", "ELL 90 LR", "2", "C-1", "1"),
    ])
    _insert_view_attachment(document, "Coming of Mod 30R128 on the East side")
    _insert_weld_log(document)
    document.save(source / "package.pdf")
    document.close()

    summary = run_fmr(source, output, FMR_TEMPLATE)
    assert summary["status"] == "complete"
    assert summary["fmr_sheets_created"] == 1
    assert summary["attachment_image_pages_ignored"] == 1
    assert summary["quarantine_entries"] == 0


def test_uncertain_material_row_is_written_and_marked_for_review(tmp_path):
    source = tmp_path / "input"
    output = tmp_path / "output"
    source.mkdir()
    document = fitz.open()
    _insert_cover(document, "IWP-ROW-REVIEW")
    _insert_iso(document, "ISO-REVIEW", "0", [
        ("1", "ELL 90 LR", "", "CODE-1", "1"),
    ])
    document.save(source / "package.pdf")
    document.close()

    summary = run_fmr(source, output, FMR_TEMPLATE)
    assert summary["status"] == "complete"
    assert summary["fmr_sheets_created"] == 1
    assert summary["material_review_entries"] == 1
    assert summary["quarantine_entries"] == 1

    workbook = load_workbook(summary["workbook"])
    worksheet = workbook["IWP-ROW-REVIEW(00)"]
    assert (worksheet["B8"].value, worksheet["C8"].value) == ("CODE-1", None)
    assert (worksheet["D8"].value, worksheet["E8"].value) == (1, "ELL 90 LR")
    assert worksheet["K8"].value == "REVIEW"
    assert worksheet["C8"].comment is not None
    assert "missing nominal size" in worksheet["C8"].comment.text
    assert worksheet["C8"].fill.fgColor.rgb == "00FFF2CC"
    workbook.close()

    with (output / "fmr_quarantine.csv").open() as handle:
        entries = list(csv.DictReader(handle))
    assert [entry["reason_code"] for entry in entries] == ["material_line_review"]
    assert "BOM point 1 was written to the FMR" in entries[0]["reason_detail"]

    manifest_path = next((output / "analytics_runs").glob("*.json"))
    manifest = json.loads(manifest_path.read_text())
    assert manifest["isos"][0]["status"] == "review"
    assert manifest["isos"][0]["material_review_rows"] == 1


@pytest.mark.skipif(
    not NEW_FMR116_PACKAGE.exists(),
    reason="Local newFmr116 package is unavailable",
)
def test_newfmr116_short_weld_annotations_are_not_spool_candidates(tmp_path):
    summary = run_fmr(
        NEW_FMR116_PACKAGE.parent,
        tmp_path / "output",
        FMR_TEMPLATE,
        include_spool_numbers=True,
    )

    assert summary["status"] == "complete"
    assert summary["fmr_sheets_created"] == 3
    assert summary["spool_rows_written"] == 6
    assert summary["spool_review_entries"] == 0
    assert summary["material_review_entries"] == 0
    assert summary["attachment_image_pages_ignored"] == 1
    assert summary["quarantine_entries"] == 0


def test_missing_or_conflicting_iwp_blocks_workbook_creation(tmp_path):
    missing = tmp_path / "missing"
    missing.mkdir()
    document = fitz.open()
    _insert_iso(document, "ISO-001", "0", [("1", "ELL 90 LR", "2", "C-1", "1")])
    document.save(missing / "iso.pdf")
    document.close()
    summary = run_fmr(missing, tmp_path / "missing-output", FMR_TEMPLATE)
    assert summary["status"] == "blocked"
    assert summary["time_estimate_available"] is False
    assert summary["estimated_time_saved_seconds"] is None
    assert summary["iwp_number"] == ""
    assert summary["workbook"] == ""

    summary = run_fmr(
        missing,
        tmp_path / "manual-output",
        FMR_TEMPLATE,
        iwp_number_override="SMM30G0012FPP-K477-103",
    )
    assert summary["status"] == "complete"
    assert summary["iwp_number"] == "SMM30G0012FPP-K477-103"
    assert summary["fmr_sheets_created"] == 1
    assert Path(summary["workbook"]).name == "SMM30G0012FPP-K477-103_FMR.xlsx"

    existing_output = tmp_path / "existing-output"
    run_fmr(
        missing,
        existing_output,
        FMR_TEMPLATE,
        iwp_number_override="SMM30G0012FPP-K477-103",
    )
    summary = run_fmr(
        missing,
        existing_output,
        FMR_TEMPLATE,
        overwrite=True,
        requested_by="Manuel Garcia",
    )
    assert summary["status"] == "complete"
    assert summary["iwp_number"] == "SMM30G0012FPP-K477-103"
    assert summary["fmr_sheets_created"] == 1
    workbook = load_workbook(summary["workbook"])
    assert workbook[workbook.sheetnames[0]]["B5"].value == "REQUESTED BY:\nManuel Garcia"
    workbook.close()

    conflicting = tmp_path / "conflicting"
    conflicting.mkdir()
    document = fitz.open()
    _insert_cover(document, "IWP-A")
    _insert_cover(document, "IWP-B")
    _insert_iso(document, "ISO-001", "0", [("1", "ELL 90 LR", "2", "C-1", "1")])
    document.save(conflicting / "package.pdf")
    document.close()
    summary = run_fmr(conflicting, tmp_path / "conflicting-output", FMR_TEMPLATE)
    assert summary["status"] == "blocked"
    assert summary["iwp_number"] == "IWP-A;IWP-B"
    assert summary["workbook"] == ""

    summary = run_fmr(
        conflicting,
        tmp_path / "manual-conflicting-output",
        FMR_TEMPLATE,
        iwp_number_override="IWP-A",
    )
    assert summary["status"] == "blocked"
    assert summary["iwp_number"] == "IWP-A;IWP-B"
    assert summary["workbook"] == ""


def test_existing_workbook_requires_overwrite(tmp_path):
    source = tmp_path / "input"
    output = tmp_path / "output"
    source.mkdir()
    _create_test_package(source / "package.pdf")
    run_fmr(source, output, FMR_TEMPLATE)
    with pytest.raises(FileExistsError, match="already exists"):
        run_fmr(source, output, FMR_TEMPLATE)
    summary = run_fmr(source, output, FMR_TEMPLATE, overwrite=True)
    assert summary["fmr_sheets_created"] == 2


@pytest.mark.skipif(
    not SUPPLIED_PACKAGE.exists() or not FMR_TEMPLATE.exists(),
    reason="Local supplied IWP package/template are unavailable",
)
def test_supplied_package_end_to_end(tmp_path):
    source = SUPPLIED_PACKAGE.parent
    summary = run_fmr(source, tmp_path / "output", FMR_TEMPLATE)
    assert summary["pages_scanned"] == 85
    assert summary["cover_pages_detected"] == 1
    assert summary["iso_pages_detected"] == 51
    assert summary["unique_drawings"] == 49
    assert summary["weld_log_pages_ignored"] == 33
    assert summary["superseded_pages_ignored"] == 2
    assert summary["fmr_sheets_created"] == 49
    assert summary["quarantine_entries"] == 0

    workbook = load_workbook(summary["workbook"])
    first = workbook["IP-SMM30R127MMPP-K408A-100(00)"]
    assert first["H5"].value == "IWP: IP-SMM30R127MMPP-K408A-100"
    assert first["H6"].value == "LINE NO:\nLP131-AI(100)-852045-04"
    assert first["K6"].value == "REV:\n0"
    assert [first[f"B{row}"].value for row in range(8, 16)] == [
        "5368751", "5368751", "5450714", "5450706",
        "5UG-02", "5CI-02", "5CI-01", "5UG-01",
    ]


@pytest.mark.skipif(
    not SUPPLIED_PACKAGE.exists(),
    reason="Local supplied IWP package is unavailable",
)
def test_supplied_rotated_wrapped_and_revision_annotation_pages():
    document = fitz.open(SUPPLIED_PACKAGE)
    try:
        rotated = document[1]
        assert rotated.rotation == 90
        first = parse_iso_page(rotated, 2, SUPPLIED_PACKAGE.name, SUPPLIED_PACKAGE.name)
        assert first.drawing_number == "LP131-AI(100)-852045-04"
        assert first.revision == "0"
        assert any(" | " in row.raw_text for row in first.materials)

        annotated = parse_iso_page(
            document[34], 35, SUPPLIED_PACKAGE.name, SUPPLIED_PACKAGE.name
        )
        notes = {note for row in annotated.materials for note in row.structural_notes}
        assert "ignored_numeric_annotation" in notes
        assert annotated.review_reasons == []
    finally:
        document.close()
