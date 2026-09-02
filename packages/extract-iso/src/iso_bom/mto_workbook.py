import os
import re
import tempfile
from copy import copy
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.table import Table, TableStyleInfo

from .mto_model import (
    MTO_SCOPE_ALL_MATERIALS,
    MTO_SCOPE_BOLTS_GASKETS,
    MTO_SCOPE_COMBINED,
    MTO_SCOPES,
    MtoIsoPage,
)


COMBINED_SHEET = "COMBINED"
BLINDS_SHEET = "BLINDS"
PIPE_FITTINGS_SHEET = "PIPE & FITTINGS"
BOLTS_GASKETS_SHEET = "BOLTS & GASKETS"
SUPPORTS_SHEET = "SUPPORTS"
VALVES_SHEET = "VALVES"
BIRDSCREENS_SHEET = "BIRDSCREENS"
OTHER_MATERIALS_SHEET = "OTHER MATERIALS"
TOTALS_SHEET = "TOTALS"
FIRST_DATA_ROW = 2
IWP_COLUMN_WIDTH = 30.29
INVALID_FILENAME_CHARS = re.compile(r"[<>:\"/\\|?*]")
BLIND_DESCRIPTION_RE = re.compile(r"\bBLINDS?\b", re.IGNORECASE)
BIRDSCREEN_DESCRIPTION_RE = re.compile(r"\bBIRD\s*SCREENS?\b", re.IGNORECASE)
VALVE_DESCRIPTION_RE = re.compile(
    r"(?:\bVALVES?\b|^\s*(?:BALL|CHECK|GATE|GLOBE|BUTTERFLY|"
    r"DIAPHRAGM|PLUG|NEEDLE|CONTROL|RELIEF|SAFETY)\b)",
    re.IGNORECASE,
)
PIPE_FITTING_DESCRIPTION_RE = re.compile(
    r"^\s*(?:PIPE|PIPET|STUB\s+END|ELL|ELBOW|TEE|RED|REDUCER|"
    r"CONC\s+RED|ECC\s+RED|CAP|COUPLING|COUP|CPLG|NIPPLE|NIP|"
    r"FLG|FLANGE|OLET|WELDOLET|SOCKOLET|THREDOLET|UNION)\b",
    re.IGNORECASE,
)
TOTAL_HEADERS = (
    "SOURCE SHEET",
    "ITEM TYPE",
    "DESCRIPTION",
    "SIZE",
    "COMMODITY CODE",
    "UOM",
    "RAW TOTAL",
    "ORDER TOTAL",
    "ORDER BASIS",
)
SCOPE_CONFIG: Dict[str, Dict[str, object]] = {
    MTO_SCOPE_COMBINED: {
        "sheet": COMBINED_SHEET,
        "columns": ("A", "B", "C", "D", "E", "F", "G", "I", "J", "K"),
        "clear_columns": tuple("ABCDEFGHIJK"),
        "last_column": 11,
        "description_column": "F",
        "size_column": "G",
        "code_column": "I",
        "quantity_column": "J",
        "uom_column": "K",
        "headers": {
            "A1": ("CWA",), "B1": ("IWP",), "C1": ("LINE NUMBER",),
            "D1": ("SHEET",), "E1": ("PIPE SPEC",), "F1": ("DESCRIPTION",),
            "G1": ("SIZE",), "I1": ("COMMODITY CODE",), "J1": ("QTY",),
            "K1": ("UOM",),
        },
    },
    MTO_SCOPE_BOLTS_GASKETS: {
        "sheet": BOLTS_GASKETS_SHEET,
        "columns": tuple("ABCDEFGHIJ"),
        "clear_columns": tuple("ABCDEFGHIJKLM"),
        "last_column": 13,
        "table_name": "BoltsGasketsMtoTable",
        "description_column": "F",
        "size_column": "G",
        "code_column": "H",
        "quantity_column": "I",
        "uom_column": "J",
        "headers": {
            "A1": ("CWA",), "B1": ("IWP", "WER"), "C1": ("LINE NUMBER",),
            "D1": ("SHEET",), "E1": ("PIPE SPEC",), "F1": ("DESCRIPTION",),
            "G1": ("SIZE",), "H1": ("COMMODITY CODE",), "I1": ("QTY",),
            "J1": ("UOM",), "K1": ("EPIC PAINT CODE",),
            "L1": ("CUST. PAINT CODE",), "M1": ("COLOR",),
        },
    },
}
SCOPE_CONFIG[MTO_SCOPE_ALL_MATERIALS] = dict(SCOPE_CONFIG[MTO_SCOPE_COMBINED])

CATEGORY_CONFIG: Dict[str, object] = {
    "columns": tuple("ABCDEFGHIJ"),
    "clear_columns": tuple("ABCDEFGHIJKLM"),
    "last_column": 13,
    "description_column": "F",
    "size_column": "G",
    "code_column": "H",
    "quantity_column": "I",
    "uom_column": "J",
    "headers": {
        "A1": ("CWA",), "B1": ("IWP", "WER"), "C1": ("LINE NUMBER",),
        "D1": ("SHEET",), "E1": ("PIPE SPEC",), "F1": ("DESCRIPTION",),
        "G1": ("SIZE",), "H1": ("COMMODITY CODE",), "I1": ("QTY",),
        "J1": ("UOM",), "K1": ("EPIC PAINT CODE",),
        "L1": ("CUST. PAINT CODE",), "M1": ("COLOR",),
    },
}
CATEGORY_SHEETS = (
    BLINDS_SHEET,
    PIPE_FITTINGS_SHEET,
    BOLTS_GASKETS_SHEET,
    SUPPORTS_SHEET,
    VALVES_SHEET,
    BIRDSCREENS_SHEET,
    OTHER_MATERIALS_SHEET,
)
CATEGORY_TABLE_NAMES = {
    BLINDS_SHEET: "ConsolidatedBlindsTable",
    PIPE_FITTINGS_SHEET: "ConsolidatedPipeFittingsTable",
    BOLTS_GASKETS_SHEET: "ConsolidatedBoltsGasketsTable",
    SUPPORTS_SHEET: "ConsolidatedSupportsTable",
    VALVES_SHEET: "ConsolidatedValvesTable",
    BIRDSCREENS_SHEET: "ConsolidatedBirdscreensTable",
    OTHER_MATERIALS_SHEET: "ConsolidatedOtherMaterialsTable",
}


def workbook_filename(iwp_number: str, scope: str = MTO_SCOPE_COMBINED) -> str:
    if not iwp_number or INVALID_FILENAME_CHARS.search(iwp_number):
        raise ValueError(f"IWP number cannot be used in an output filename: {iwp_number!r}")
    if scope not in MTO_SCOPES:
        raise ValueError(f"Unsupported MTO scope: {scope!r}")
    if scope == MTO_SCOPE_BOLTS_GASKETS:
        return f"MTO {iwp_number} - BOLTS & GASKETS.xlsx"
    if scope == MTO_SCOPE_ALL_MATERIALS:
        return f"MTO {iwp_number} - ALL MATERIALS.xlsx"
    return f"MTO {iwp_number}.xlsx"


def consolidated_workbook_filename(
    cwa: str,
    scope: str = MTO_SCOPE_COMBINED,
) -> str:
    if not cwa or INVALID_FILENAME_CHARS.search(cwa):
        raise ValueError(f"CWA cannot be used in an output filename: {cwa!r}")
    if scope not in MTO_SCOPES:
        raise ValueError(f"Unsupported MTO scope: {scope!r}")
    base = f"MTO CWA {cwa} - CONSOLIDATED"
    if scope == MTO_SCOPE_BOLTS_GASKETS:
        return f"{base} - BOLTS & GASKETS.xlsx"
    if scope == MTO_SCOPE_ALL_MATERIALS:
        return f"{base} - ALL MATERIALS.xlsx"
    return f"{base}.xlsx"


def _typed_value(value: str):
    value = str(value or "").strip()
    if re.fullmatch(r"\d+", value):
        return int(value)
    if re.fullmatch(r"\d+\.\d+", value):
        return float(value)
    return value


def _mto_rows(
    cwa: str,
    iwp_number: str,
    iso_pages: Sequence[MtoIsoPage],
) -> List[Tuple[List[object], str, bool]]:
    rows: List[Tuple[List[object], str, bool]] = []
    for iso in iso_pages:
        for material in iso.materials:
            values = [
                cwa,
                iwp_number,
                iso.drawing_number,
                _typed_value(iso.sheet_number),
                _typed_value(iso.pipe_schedule),
                material.description,
                _typed_value(material.nominal_size),
                material.commodity_code,
                _typed_value(material.quantity),
                material.uom or "EA",
            ]
            rows.append((values, material.item_type, material.is_partial or not iso.sheet_number or not iso.pipe_schedule))
    return rows


def _consolidated_mto_rows(
    cwa: str,
    packages: Sequence[Tuple[str, Sequence[MtoIsoPage]]],
) -> List[Tuple[List[object], str, bool]]:
    rows: List[Tuple[List[object], str, bool]] = []
    for iwp_number, iso_pages in packages:
        rows.extend(_mto_rows(cwa, iwp_number, iso_pages))
    return rows


def material_category_sheet(item_type: str, description: str) -> str:
    normalized_type = str(item_type or "").strip().upper()
    normalized_description = " ".join(str(description or "").split())
    if BLIND_DESCRIPTION_RE.search(normalized_description):
        return BLINDS_SHEET
    if normalized_type in {"BOLT", "GASKET", "WASHER"}:
        return BOLTS_GASKETS_SHEET
    if normalized_type == "SUPPORT":
        return SUPPORTS_SHEET
    if BIRDSCREEN_DESCRIPTION_RE.search(normalized_description):
        return BIRDSCREENS_SHEET
    if VALVE_DESCRIPTION_RE.search(normalized_description):
        return VALVES_SHEET
    if normalized_type in {"PIPE", "FITTING"} or PIPE_FITTING_DESCRIPTION_RE.search(
        normalized_description
    ):
        return PIPE_FITTINGS_SHEET
    return OTHER_MATERIALS_SHEET


def _categorized_mto_rows(
    rows: Sequence[Tuple[List[object], str, bool]],
) -> Dict[str, List[Tuple[List[object], str, bool]]]:
    categorized: Dict[str, List[Tuple[List[object], str, bool]]] = {
        sheet_name: [] for sheet_name in CATEGORY_SHEETS
    }
    for row in rows:
        values, item_type, _partial = row
        description = str(values[5] or "")
        categorized[material_category_sheet(item_type, description)].append(row)
    return categorized


def consolidated_tab_row_counts(
    cwa: str,
    packages: Sequence[Tuple[str, Sequence[MtoIsoPage]]],
) -> Dict[str, int]:
    rows = _consolidated_mto_rows(cwa, packages)
    categorized = _categorized_mto_rows(rows)
    return {
        COMBINED_SHEET: len(rows),
        **{sheet_name: len(categorized[sheet_name]) for sheet_name in CATEGORY_SHEETS},
    }


def _copy_row_style(worksheet, source_row: int, target_row: int) -> None:
    worksheet.row_dimensions[target_row].height = worksheet.row_dimensions[source_row].height
    for column in range(1, worksheet.max_column + 1):
        source = worksheet.cell(source_row, column)
        target = worksheet.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.alignment:
            target.alignment = copy(source.alignment)
        if source.font:
            target.font = copy(source.font)
        if source.fill:
            target.fill = copy(source.fill)
        if source.border:
            target.border = copy(source.border)
        if source.protection:
            target.protection = copy(source.protection)


def _validate_template(worksheet, config: Dict[str, object]) -> None:
    for coordinate, expected_values in config["headers"].items():
        actual = worksheet[coordinate].value
        normalized = str(actual or "").strip().upper()
        if normalized not in expected_values:
            raise ValueError(
                f"MTO template header {coordinate} must be one of "
                f"{expected_values!r}, found {actual!r}"
            )


def validate_mto_template(template_path: Path, scope: str) -> None:
    template_path = template_path.resolve()
    if not template_path.is_file():
        raise ValueError(f"MTO template does not exist: {template_path}")
    if scope not in MTO_SCOPES:
        raise ValueError(f"Unsupported MTO scope: {scope!r}")
    workbook = load_workbook(template_path, read_only=True, data_only=False)
    try:
        sheet_name = str(SCOPE_CONFIG[scope]["sheet"])
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"MTO template workbook is missing worksheet {sheet_name!r}"
            )
        _validate_template(workbook[sheet_name], SCOPE_CONFIG[scope])
    finally:
        workbook.close()


def _clear_existing_data(worksheet, columns: Sequence[str]) -> None:
    for row in range(FIRST_DATA_ROW, worksheet.max_row + 1):
        for column in columns:
            worksheet[f"{column}{row}"] = None


def _remove_legacy_totals(worksheet, last_column: int) -> None:
    if worksheet.max_column > last_column:
        worksheet.delete_cols(last_column + 1, worksheet.max_column - last_column)


def _trim_unused_rows(worksheet, last_row: int) -> None:
    if worksheet.max_row > last_row:
        worksheet.delete_rows(last_row + 1, worksheet.max_row - last_row)
    for row_number in tuple(worksheet.row_dimensions):
        if row_number > last_row:
            del worksheet.row_dimensions[row_number]


def _compact_inactive_sheets(workbook, active_sheet: str) -> None:
    sheet_last_columns = {
        COMBINED_SHEET: 11,
        "BLINDS": 13,
        "PIPE & FITTINGS": 13,
        BOLTS_GASKETS_SHEET: 13,
    }
    for sheet_name, last_column in sheet_last_columns.items():
        if sheet_name == active_sheet or sheet_name not in workbook.sheetnames:
            continue
        worksheet = workbook[sheet_name]
        _clear_existing_data(
            worksheet,
            tuple(get_column_letter(column) for column in range(1, last_column + 1)),
        )
        _remove_legacy_totals(worksheet, last_column)
        _trim_unused_rows(worksheet, 1)
        for table in worksheet.tables.values():
            minimum_column, _minimum_row, maximum_column, _maximum_row = (
                range_boundaries(table.ref)
            )
            table.ref = (
                f"{get_column_letter(minimum_column)}1:"
                f"{get_column_letter(maximum_column)}1"
            )
            if table.autoFilter is not None:
                table.autoFilter.ref = table.ref
        worksheet.freeze_panes = "A2"
        worksheet.sheet_view.showGridLines = False
        worksheet.print_area = f"A1:{get_column_letter(last_column)}1"


def _table_style() -> TableStyleInfo:
    return TableStyleInfo(
        name="TableStyleMedium4",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )


def _add_native_table(worksheet, name: str, reference: str) -> None:
    if name in worksheet.tables:
        del worksheet.tables[name]
    table = Table(displayName=name, ref=reference)
    table.tableStyleInfo = _table_style()
    worksheet.add_table(table)


def _resize_or_add_table(worksheet, name: str, reference: str) -> None:
    existing = list(worksheet.tables.values())
    if existing:
        existing[0].ref = reference
        if existing[0].autoFilter is not None:
            existing[0].autoFilter.ref = reference
        for table in existing[1:]:
            del worksheet.tables[table.name]
        return
    _add_native_table(worksheet, name, reference)


def _ensure_consolidated_category_sheets(workbook) -> None:
    if BOLTS_GASKETS_SHEET not in workbook.sheetnames:
        raise ValueError(
            f"MTO template workbook is missing worksheet {BOLTS_GASKETS_SHEET!r}"
        )
    style_source = workbook[BOLTS_GASKETS_SHEET]
    for sheet_name in CATEGORY_SHEETS:
        if sheet_name in workbook.sheetnames:
            continue
        worksheet = workbook.copy_worksheet(style_source)
        worksheet.title = sheet_name


def _populate_category_sheet(
    worksheet,
    rows: Sequence[Tuple[List[object], str, bool]],
    table_name: str,
) -> None:
    _validate_template(worksheet, CATEGORY_CONFIG)
    worksheet["B1"] = "IWP"
    _clear_existing_data(worksheet, CATEGORY_CONFIG["clear_columns"])
    _remove_legacy_totals(worksheet, int(CATEGORY_CONFIG["last_column"]))
    worksheet.column_dimensions["B"].width = max(
        worksheet.column_dimensions["B"].width or 0,
        IWP_COLUMN_WIDTH,
    )

    required_last_row = FIRST_DATA_ROW + len(rows) - 1
    for row_number in range(worksheet.max_row + 1, required_last_row + 1):
        _copy_row_style(worksheet, FIRST_DATA_ROW, row_number)

    text_columns = {"A", "B", "C", "H", "J"}
    for row_number, (values, _item_type, _partial) in enumerate(rows, FIRST_DATA_ROW):
        for column, value in zip(CATEGORY_CONFIG["columns"], values):
            cell = worksheet[f"{column}{row_number}"]
            cell.value = value
            cell.number_format = "@" if column in text_columns else "General"

    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False
    last_row = max(1, required_last_row)
    _trim_unused_rows(worksheet, last_row)
    _resize_or_add_table(worksheet, table_name, f"A1:M{last_row}")
    worksheet.print_area = f"A1:M{last_row}"


def _replace_combined_tables(worksheet, last_row: int) -> None:
    for table_name in tuple(worksheet.tables):
        del worksheet.tables[table_name]
    _add_native_table(worksheet, "CombinedMtoDetailsTable", f"A1:G{last_row}")
    _add_native_table(worksheet, "CombinedMtoCodesTable", f"I1:K{last_row}")


def _totals_groups(
    source_sheet: str,
    rows: Sequence[Tuple[List[object], str, bool]],
) -> List[Tuple[str, str, object, str, str]]:
    groups = set()
    for values, item_type, _partial in rows:
        description = str(values[5] or "").strip()
        size = values[6]
        code = str(values[7] or "").strip()
        quantity = values[8]
        uom = str(values[9] or "").strip()
        if not description or size in (None, "") or not code or not uom:
            continue
        if not isinstance(quantity, (int, float)):
            continue
        groups.add((item_type, description, size, code, uom))
    return sorted(groups, key=lambda item: tuple(str(value).upper() for value in item))


def _build_totals_sheet(
    workbook,
    config: Dict[str, object],
    rows: Sequence[Tuple[List[object], str, bool]],
    categorized_rows: Optional[
        Dict[str, List[Tuple[List[object], str, bool]]]
    ] = None,
) -> None:
    if TOTALS_SHEET in workbook.sheetnames:
        del workbook[TOTALS_SHEET]
    worksheet = workbook.create_sheet(TOTALS_SHEET)
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A2"

    for column, header in enumerate(TOTAL_HEADERS, 1):
        worksheet.cell(1, column, header)
    worksheet["K1"] = "PIPE STOCK LENGTH (FT)"
    worksheet["K2"] = 20

    header_fill = PatternFill("solid", fgColor="A8D08D")
    header_font = Font(name="Calibri", size=11, bold=True, color="000000")
    thin = Side(style="thin", color="808080")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in worksheet[1]:
        if cell.column > 11:
            continue
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.row_dimensions[1].height = 34
    worksheet["K2"].number_format = "0"
    worksheet["K2"].alignment = Alignment(horizontal="center")
    worksheet["K2"].border = border

    if categorized_rows is None:
        sources = [(str(config["sheet"]), config, rows)]
    else:
        sources = [
            (sheet_name, CATEGORY_CONFIG, categorized_rows[sheet_name])
            for sheet_name in CATEGORY_SHEETS
            if categorized_rows[sheet_name]
        ]
    total_entries = [
        (source_sheet, source_config, source_rows, group)
        for source_sheet, source_config, source_rows in sources
        for group in _totals_groups(source_sheet, source_rows)
    ]

    for row_number, entry in enumerate(total_entries, FIRST_DATA_ROW):
        source_sheet, source_config, source_rows, group = entry
        description_column = str(source_config["description_column"])
        size_column = str(source_config["size_column"])
        code_column = str(source_config["code_column"])
        quantity_column = str(source_config["quantity_column"])
        uom_column = str(source_config["uom_column"])
        source_last_row = max(
            FIRST_DATA_ROW,
            FIRST_DATA_ROW + len(source_rows) - 1,
        )
        item_type, description, size, code, uom = group
        worksheet.cell(row_number, 1, source_sheet)
        worksheet.cell(row_number, 2, item_type)
        worksheet.cell(row_number, 3, description)
        worksheet.cell(row_number, 4, size)
        worksheet.cell(row_number, 5, code)
        worksheet.cell(row_number, 6, uom)
        worksheet.cell(row_number, 7, (
            f"=SUMIFS('{source_sheet}'!${quantity_column}$2:${quantity_column}${source_last_row},"
            f"'{source_sheet}'!${description_column}$2:${description_column}${source_last_row},C{row_number},"
            f"'{source_sheet}'!${size_column}$2:${size_column}${source_last_row},D{row_number},"
            f"'{source_sheet}'!${code_column}$2:${code_column}${source_last_row},E{row_number},"
            f"'{source_sheet}'!${uom_column}$2:${uom_column}${source_last_row},F{row_number})"
        ))
        worksheet.cell(row_number, 8, f'=IF(F{row_number}="LF",CEILING(G{row_number},$K$2),G{row_number})')
        worksheet.cell(row_number, 9, f'=IF(F{row_number}="LF",TEXT($K$2,"0")&" FT STOCK","EXACT QTY")')
        for cell in worksheet[row_number][:9]:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=cell.column == 3)
        worksheet.cell(row_number, 7).number_format = "General"
        worksheet.cell(row_number, 8).number_format = "General"

    total_last_row = max(1, 1 + len(total_entries))
    _add_native_table(worksheet, "MtoTotalsTable", f"A1:I{total_last_row}")
    worksheet.print_area = f"A1:K{max(2, total_last_row)}"
    widths = {"A": 22, "B": 15, "C": 58, "D": 13, "E": 24, "F": 10, "G": 14, "H": 14, "I": 18, "K": 24}
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width


def _build_mto_workbook_from_rows(
    template_path: Path,
    output_path: Path,
    rows: Sequence[Tuple[List[object], str, bool]],
    overwrite: bool = False,
    scope: str = MTO_SCOPE_COMBINED,
    categorize: bool = False,
) -> Path:
    template_path = template_path.resolve()
    output_path = output_path.resolve()
    if not template_path.is_file():
        raise ValueError(f"MTO template does not exist: {template_path}")
    if scope not in MTO_SCOPES:
        raise ValueError(f"Unsupported MTO scope: {scope!r}")
    if output_path.exists() and not overwrite:
        raise FileExistsError(f"Output workbook already exists: {output_path}")

    workbook = load_workbook(template_path)
    config = SCOPE_CONFIG[scope]
    sheet_name = str(config["sheet"])
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"MTO template workbook is missing worksheet {sheet_name!r}")
    worksheet = workbook[sheet_name]
    _validate_template(worksheet, config)
    if categorize:
        _ensure_consolidated_category_sheets(workbook)
    else:
        _compact_inactive_sheets(workbook, sheet_name)
    if scope == MTO_SCOPE_BOLTS_GASKETS:
        worksheet["B1"] = "IWP"
    _clear_existing_data(worksheet, config["clear_columns"])
    _remove_legacy_totals(worksheet, int(config["last_column"]))
    worksheet.column_dimensions["B"].width = max(
        worksheet.column_dimensions["B"].width or 0,
        IWP_COLUMN_WIDTH,
    )

    required_last_row = FIRST_DATA_ROW + len(rows) - 1
    for row_number in range(worksheet.max_row + 1, required_last_row + 1):
        _copy_row_style(worksheet, FIRST_DATA_ROW, row_number)

    text_columns = {"A", "B", "C", str(config["code_column"]), str(config["uom_column"])}
    for row_number, (values, _item_type, _partial) in enumerate(rows, FIRST_DATA_ROW):
        for column, value in zip(config["columns"], values):
            cell = worksheet[f"{column}{row_number}"]
            cell.value = value
            if column in text_columns:
                cell.number_format = "@"
            else:
                cell.number_format = "General"

    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False
    last_row = max(1, required_last_row)
    _trim_unused_rows(worksheet, last_row)
    last_column = int(config["last_column"])
    last_column_letter = get_column_letter(last_column)
    if scope in (MTO_SCOPE_COMBINED, MTO_SCOPE_ALL_MATERIALS):
        _replace_combined_tables(worksheet, last_row)
    else:
        _resize_or_add_table(
            worksheet,
            str(config["table_name"]),
            f"A1:{last_column_letter}{last_row}",
        )
    worksheet.print_area = f"A1:{last_column_letter}{last_row}"
    categorized_rows = _categorized_mto_rows(rows) if categorize else None
    if categorized_rows is not None:
        for category_sheet in CATEGORY_SHEETS:
            _populate_category_sheet(
                workbook[category_sheet],
                categorized_rows[category_sheet],
                CATEGORY_TABLE_NAMES[category_sheet],
            )
    _build_totals_sheet(
        workbook,
        config,
        rows,
        categorized_rows=categorized_rows,
    )
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.active = workbook.sheetnames.index(sheet_name)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=".mto-", suffix=".xlsx", dir=str(output_path.parent)
    )
    os.close(descriptor)
    temporary_path = Path(temporary_name)
    try:
        workbook.save(temporary_path)
        if output_path.exists() and not overwrite:
            raise FileExistsError(f"Output workbook already exists: {output_path}")
        os.replace(temporary_path, output_path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()
    return output_path


def build_mto_workbook(
    template_path: Path,
    output_path: Path,
    iwp_number: str,
    cwa: str,
    iso_pages: Sequence[MtoIsoPage],
    overwrite: bool = False,
    scope: str = MTO_SCOPE_COMBINED,
) -> Path:
    if not iso_pages:
        raise ValueError("Cannot create an MTO workbook without accepted ISO pages")
    return _build_mto_workbook_from_rows(
        template_path,
        output_path,
        _mto_rows(cwa, iwp_number, iso_pages),
        overwrite=overwrite,
        scope=scope,
    )


def build_consolidated_mto_workbook(
    template_path: Path,
    output_path: Path,
    cwa: str,
    packages: Sequence[Tuple[str, Sequence[MtoIsoPage]]],
    overwrite: bool = False,
    scope: str = MTO_SCOPE_COMBINED,
) -> Path:
    if not packages or not any(iso_pages for _iwp, iso_pages in packages):
        raise ValueError(
            "Cannot create a consolidated MTO workbook without accepted ISO pages"
        )
    return _build_mto_workbook_from_rows(
        template_path,
        output_path,
        _consolidated_mto_rows(cwa, packages),
        overwrite=overwrite,
        scope=scope,
        categorize=scope in (MTO_SCOPE_COMBINED, MTO_SCOPE_ALL_MATERIALS),
    )
