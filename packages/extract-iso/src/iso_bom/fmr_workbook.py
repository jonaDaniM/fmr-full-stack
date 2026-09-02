import os
import re
import tempfile
from copy import copy, deepcopy
from io import BytesIO
from pathlib import Path
from typing import List, Optional, Sequence, Tuple

from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, PatternFill

from .fmr_model import IsoPage, material_review_reasons


TEMPLATE_SHEET = "blankFMR"
MATERIAL_FIRST_ROW = 8
MATERIAL_LAST_ROW = 30
MATERIAL_CAPACITY = MATERIAL_LAST_ROW - MATERIAL_FIRST_ROW + 1
INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")
INVALID_FILENAME_CHARS = re.compile(r"[<>:\"/\\|?*]")
REVIEW_FILL = PatternFill("solid", fgColor="FFF2CC")
REVIEW_FONT_COLOR = "9C6500"


def sheet_names(iwp_number: str, count: int) -> List[str]:
    names = [f"{iwp_number}({index:02d})" for index in range(count)]
    for name in names:
        if len(name) > 31:
            raise ValueError(f"Generated worksheet name exceeds 31 characters: {name}")
        if INVALID_SHEET_CHARS.search(name):
            raise ValueError(f"Generated worksheet name contains an invalid Excel character: {name}")
    return names


def workbook_filename(iwp_number: str) -> str:
    if not iwp_number or INVALID_FILENAME_CHARS.search(iwp_number):
        raise ValueError(f"IWP number cannot be used in an output filename: {iwp_number!r}")
    return f"{iwp_number}_FMR.xlsx"


def _template_images(worksheet) -> List[Tuple[bytes, object, float, float]]:
    images = []
    for image in worksheet._images:
        images.append((image._data(), deepcopy(image.anchor), image.width, image.height))
    return images


def _shift_anchor(anchor, start_row: int, row_offset: int):
    shifted = deepcopy(anchor)
    if row_offset <= 0:
        return shifted
    zero_based_start = start_row - 1
    if hasattr(shifted, "_from") and shifted._from.row >= zero_based_start:
        shifted._from.row += row_offset
    if hasattr(shifted, "to") and shifted.to.row >= zero_based_start:
        shifted.to.row += row_offset
    return shifted


def _add_images(
    worksheet,
    images: Sequence[Tuple[bytes, object, float, float]],
    *,
    row_offset_start: int = MATERIAL_LAST_ROW + 1,
    row_offset: int = 0,
) -> None:
    for data, anchor, width, height in images:
        image = Image(BytesIO(data))
        image.anchor = _shift_anchor(anchor, row_offset_start, row_offset)
        image.width = width
        image.height = height
        worksheet.add_image(image)


def _normalize_material_description_cells(
    worksheet,
    material_last_row: int = MATERIAL_LAST_ROW,
) -> None:
    merged = {str(cell_range) for cell_range in worksheet.merged_cells.ranges}
    reference_alignment = copy(worksheet["E23"].alignment)
    reference_alignment.shrink_to_fit = True
    for row in range(MATERIAL_FIRST_ROW, material_last_row + 1):
        target = f"E{row}:H{row}"
        if target not in merged:
            worksheet.merge_cells(target)
        worksheet[f"E{row}"].alignment = copy(reference_alignment)


def _copy_material_row_format(worksheet, source_row: int, target_row: int) -> None:
    source_dimension = worksheet.row_dimensions[source_row]
    target_dimension = worksheet.row_dimensions[target_row]
    target_dimension.height = source_dimension.height
    target_dimension.hidden = source_dimension.hidden
    target_dimension.outlineLevel = source_dimension.outlineLevel
    for column in range(1, worksheet.max_column + 1):
        source = worksheet.cell(source_row, column)
        target = worksheet.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
        target.number_format = source.number_format
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)


def _shift_merged_ranges_down(worksheet, start_row: int, row_offset: int) -> None:
    shifted_ranges = []
    for cell_range in list(worksheet.merged_cells.ranges):
        if cell_range.min_row >= start_row:
            shifted = copy(cell_range)
            worksheet.unmerge_cells(str(cell_range))
            shifted.shift(row_shift=row_offset)
            shifted_ranges.append(shifted)
    for cell_range in shifted_ranges:
        worksheet.merge_cells(str(cell_range))


def _ensure_material_rows(worksheet, required_rows: int) -> Tuple[int, int]:
    required_rows = max(0, int(required_rows))
    if required_rows <= MATERIAL_CAPACITY:
        _normalize_material_description_cells(worksheet)
        return MATERIAL_LAST_ROW, 0

    extra_rows = required_rows - MATERIAL_CAPACITY
    insert_at = MATERIAL_LAST_ROW + 1
    _shift_merged_ranges_down(worksheet, insert_at, extra_rows)
    worksheet.insert_rows(insert_at, amount=extra_rows)
    for target_row in range(insert_at, insert_at + extra_rows):
        _copy_material_row_format(worksheet, MATERIAL_LAST_ROW, target_row)
    material_last_row = MATERIAL_LAST_ROW + extra_rows
    _normalize_material_description_cells(worksheet, material_last_row)
    return material_last_row, extra_rows


def _quantity_value(value: str):
    value = value.strip()
    if re.fullmatch(r"\d+", value):
        return int(value)
    if re.fullmatch(r"\d+\.\d+", value):
        return float(value)
    return value


def _description_font_size(description: str, template_size: float) -> float:
    if len(description) > 82:
        return min(template_size, 5)
    if len(description) > 68:
        return min(template_size, 5.5)
    if len(description) > 52:
        return min(template_size, 6)
    if len(description) > 45:
        return min(template_size, 7)
    return template_size


def _mark_material_review(worksheet, row_number: int, material) -> None:
    reasons = material_review_reasons(material)
    if not reasons:
        return

    columns = set()
    for reason in reasons:
        if reason in {"missing_commodity_code", "multiple_commodity_code_candidates"}:
            columns.add("B")
        elif reason == "missing_nominal_size":
            columns.add("C")
        elif reason in {"missing_quantity", "multiple_quantity_candidates"}:
            columns.add("D")
        elif reason == "missing_description":
            columns.update("EFGH")
        else:
            columns.update("BCDEFGH")

    point = str(material.point_number or "?").strip() or "?"
    readable_reasons = ", ".join(reason.replace("_", " ") for reason in reasons)
    note = (
        f"Review BOM point {point}: {readable_reasons}. "
        "Confirm the highlighted field against the source ISO."
    )
    for column in columns:
        worksheet[f"{column}{row_number}"].fill = copy(REVIEW_FILL)
    for column in {"B", "C", "D", "E"}.intersection(columns):
        worksheet[f"{column}{row_number}"].comment = Comment(note, "ISO FMR")

    review_cell = worksheet[f"K{row_number}"]
    review_cell.value = "REVIEW"
    review_cell.fill = copy(REVIEW_FILL)
    review_font = copy(review_cell.font)
    review_font.bold = True
    review_font.color = REVIEW_FONT_COLOR
    review_cell.font = review_font
    review_cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        shrink_to_fit=True,
    )
    review_cell.comment = Comment(note, "ISO FMR")


def _blank_label(value, label: str):
    """Keep a template rich-text label run while removing its populated value."""
    if isinstance(value, CellRichText) and value:
        label_run = deepcopy(value[0])
        if hasattr(label_run, "text"):
            label_run.text = f"{label}\n"
            return CellRichText([label_run])
    return f"{label}\n"


def _labeled_value(value, label: str, field_value: Optional[str]):
    cleaned_value = (field_value or "").strip()
    if not cleaned_value:
        return _blank_label(value, label)
    if isinstance(value, CellRichText) and value:
        label_run = deepcopy(value[0])
        value_run = deepcopy(value[-1])
        if hasattr(label_run, "text") and hasattr(value_run, "text"):
            label_run.text = f"{label}\n"
            value_run.text = cleaned_value
            return CellRichText([label_run, value_run])
    return f"{label}\n{cleaned_value}"


def _populate_sheet(
    worksheet,
    iwp_number: str,
    iso: IsoPage,
    destination: Optional[str] = None,
    requested_by: Optional[str] = None,
    deliver_to: Optional[str] = None,
) -> Tuple[int, int]:
    material_last_row, row_offset = _ensure_material_rows(
        worksheet, len(iso.spool_numbers) + len(iso.materials)
    )
    if destination is not None:
        worksheet["B4"] = _labeled_value(
            worksheet["B4"].value, "DESTINATION:", destination
        )
    worksheet["B5"] = _labeled_value(
        worksheet["B5"].value, "REQUESTED BY:", requested_by
    )
    worksheet["B6"] = _labeled_value(
        worksheet["B6"].value, "DELIVER TO:", deliver_to
    )
    worksheet["H5"] = f"IWP: {iwp_number}"
    worksheet["H6"] = f"LINE NO:\n{iso.drawing_number}"
    worksheet["K6"] = f"REV:\n{iso.revision}"

    for row in range(MATERIAL_FIRST_ROW, material_last_row + 1):
        for column in ("B", "C", "D", "E", "I", "J", "K"):
            worksheet[f"{column}{row}"] = None

    for row_number, spool_number in enumerate(
        iso.spool_numbers, MATERIAL_FIRST_ROW
    ):
        worksheet[f"B{row_number}"] = f"SPOOL {spool_number}"
        worksheet[f"B{row_number}"].number_format = "@"
        worksheet[f"D{row_number}"] = 1

    material_first_row = MATERIAL_FIRST_ROW + len(iso.spool_numbers)
    for row_number, material in enumerate(iso.materials, material_first_row):
        worksheet[f"B{row_number}"] = material.commodity_code
        worksheet[f"B{row_number}"].number_format = "@"
        worksheet[f"C{row_number}"] = material.nominal_size
        worksheet[f"C{row_number}"].number_format = "@"
        worksheet[f"D{row_number}"] = _quantity_value(material.quantity)
        description_cell = worksheet[f"E{row_number}"]
        description_cell.value = material.description
        description_font = copy(description_cell.font)
        template_size = float(description_font.sz or 8)
        description_font.sz = _description_font_size(material.description, template_size)
        description_cell.font = description_font
        _mark_material_review(worksheet, row_number, material)
    return material_last_row, row_offset


def build_fmr_workbook(
    template_path: Path,
    output_path: Path,
    iwp_number: str,
    iso_pages: Sequence[IsoPage],
    overwrite: bool = False,
    destination: Optional[str] = None,
    requested_by: Optional[str] = None,
    deliver_to: Optional[str] = None,
) -> Path:
    template_path = template_path.resolve()
    output_path = output_path.resolve()
    if not template_path.is_file():
        raise ValueError(f"FMR template does not exist: {template_path}")
    if not iso_pages:
        raise ValueError("Cannot create an FMR workbook without accepted ISO pages")
    if output_path.exists() and not overwrite:
        raise FileExistsError(f"Output workbook already exists: {output_path}")
    names = sheet_names(iwp_number, len(iso_pages))

    workbook = load_workbook(template_path, rich_text=True)
    if TEMPLATE_SHEET not in workbook.sheetnames:
        raise ValueError(f"FMR template workbook is missing worksheet {TEMPLATE_SHEET!r}")
    source = workbook[TEMPLATE_SHEET]
    _normalize_material_description_cells(source)
    images = _template_images(source)

    for worksheet in list(workbook.worksheets):
        if worksheet is not source:
            workbook.remove(worksheet)

    generated = []
    for name, iso in zip(names, iso_pages):
        worksheet = workbook.copy_worksheet(source)
        worksheet.title = name
        _, row_offset = _populate_sheet(
            worksheet,
            iwp_number,
            iso,
            destination=destination,
            requested_by=requested_by,
            deliver_to=deliver_to,
        )
        _add_images(worksheet, images, row_offset=row_offset)
        generated.append(worksheet)

    workbook.remove(source)
    workbook.active = 0
    output_path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=".fmr-", suffix=".xlsx", dir=str(output_path.parent)
    )
    os.close(descriptor)
    temporary_path = Path(temporary_name)
    try:
        workbook.save(temporary_path)
        if output_path.exists() and not overwrite:
            raise FileExistsError(f"Output workbook already exists: {output_path}")
        os.replace(temporary_path, output_path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()
    return output_path
