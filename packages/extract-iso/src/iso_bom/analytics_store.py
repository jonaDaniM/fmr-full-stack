import csv
import hashlib
import json
import re
import sqlite3
import uuid
from collections import defaultdict
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Dict, Iterable, List, Mapping, Optional, Sequence, Tuple

from docx import Document
from openpyxl import load_workbook

from .time_savings import (
    TIME_MODEL_VERSION,
    fmr_iso_time_saved,
    fmr_time_estimate,
    is_overflow_description,
    pipe_iso_time_saved,
    pipe_time_estimate,
    unavailable_estimate,
)


SCHEMA_VERSION = 3
SUCCESS_STATUSES = {"complete", "complete_no_workbook"}
WORKFLOW_LABELS = {
    "pipe_footage": "Pipe Footage",
    "fmr": "Field Material Requests",
    "mto": "Material Takeoffs",
    "bolt_gasket": "Bolt & Gasket",
}


def connect_database(path: Path) -> sqlite3.Connection:
    path = path.resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(path, timeout=10)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    connection.execute("PRAGMA journal_mode = WAL")
    connection.execute("PRAGMA busy_timeout = 10000")
    _migrate(connection)
    return connection


def _migrate(connection: sqlite3.Connection) -> None:
    current = int(connection.execute("PRAGMA user_version").fetchone()[0])
    if current > SCHEMA_VERSION:
        raise RuntimeError(
            f"Analytics database schema {current} is newer than supported version {SCHEMA_VERSION}"
        )
    if current < 1:
        connection.executescript(
            """
            CREATE TABLE IF NOT EXISTS runs (
                run_id TEXT PRIMARY KEY,
                fingerprint TEXT NOT NULL UNIQUE,
                workflow TEXT NOT NULL,
                generated_at TEXT NOT NULL,
                status TEXT NOT NULL,
                package_number TEXT NOT NULL DEFAULT '',
                input_folder TEXT NOT NULL DEFAULT '',
                output_folder TEXT NOT NULL DEFAULT '',
                package_total_lf TEXT,
                fmr_sheets INTEGER NOT NULL DEFAULT 0,
                accepted_records INTEGER NOT NULL DEFAULT 0,
                iso_count INTEGER NOT NULL DEFAULT 0,
                review_count INTEGER NOT NULL DEFAULT 0,
                summary_json TEXT NOT NULL,
                manifest_path TEXT NOT NULL,
                imported_legacy INTEGER NOT NULL DEFAULT 0,
                indexed_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS isos (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                run_id TEXT NOT NULL REFERENCES runs(run_id) ON DELETE CASCADE,
                drawing_number TEXT NOT NULL,
                revision TEXT NOT NULL DEFAULT '',
                status TEXT NOT NULL DEFAULT 'accepted',
                linear_feet TEXT,
                pipe_rows INTEGER NOT NULL DEFAULT 0,
                material_rows INTEGER NOT NULL DEFAULT 0,
                source_pdf TEXT NOT NULL DEFAULT '',
                source_path TEXT NOT NULL DEFAULT '',
                page INTEGER NOT NULL DEFAULT 0,
                UNIQUE(run_id, drawing_number, revision, page)
            );
            CREATE TABLE IF NOT EXISTS artifacts (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                run_id TEXT NOT NULL REFERENCES runs(run_id) ON DELETE CASCADE,
                kind TEXT NOT NULL,
                name TEXT NOT NULL,
                path TEXT NOT NULL,
                exists_flag INTEGER NOT NULL DEFAULT 0,
                size INTEGER,
                modified_at TEXT,
                UNIQUE(run_id, path)
            );
            CREATE TABLE IF NOT EXISTS issues (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                run_id TEXT NOT NULL REFERENCES runs(run_id) ON DELETE CASCADE,
                source_pdf TEXT NOT NULL DEFAULT '',
                source_path TEXT NOT NULL DEFAULT '',
                page INTEGER NOT NULL DEFAULT 0,
                drawing_number TEXT NOT NULL DEFAULT '',
                revision TEXT NOT NULL DEFAULT '',
                reason_code TEXT NOT NULL DEFAULT '',
                reason_detail TEXT NOT NULL DEFAULT ''
            );
            CREATE INDEX IF NOT EXISTS runs_package_idx ON runs(package_number, generated_at DESC);
            CREATE INDEX IF NOT EXISTS runs_workflow_idx ON runs(workflow, generated_at DESC);
            CREATE INDEX IF NOT EXISTS runs_status_idx ON runs(status, generated_at DESC);
            CREATE INDEX IF NOT EXISTS isos_drawing_idx ON isos(drawing_number);
            CREATE INDEX IF NOT EXISTS isos_run_idx ON isos(run_id);
            CREATE INDEX IF NOT EXISTS artifacts_run_idx ON artifacts(run_id);
            CREATE INDEX IF NOT EXISTS issues_run_idx ON issues(run_id);
            PRAGMA user_version = 1;
            """
        )
        connection.commit()
        current = 1
    if current < 2:
        connection.executescript(
            """
            ALTER TABLE runs ADD COLUMN processing_seconds TEXT;
            ALTER TABLE runs ADD COLUMN estimated_time_saved_seconds TEXT;
            ALTER TABLE runs ADD COLUMN time_model_version TEXT NOT NULL DEFAULT '';
            ALTER TABLE runs ADD COLUMN time_estimate_available INTEGER NOT NULL DEFAULT 0;
            ALTER TABLE runs ADD COLUMN time_saved_breakdown_json TEXT NOT NULL DEFAULT '{}';
            ALTER TABLE isos ADD COLUMN overflow_material_rows INTEGER NOT NULL DEFAULT 0;
            ALTER TABLE isos ADD COLUMN estimated_time_saved_seconds TEXT;
            PRAGMA user_version = 2;
            """
        )
        connection.commit()
        current = 2
    if current < 3:
        connection.executescript(
            """
            ALTER TABLE runs ADD COLUMN cwa TEXT NOT NULL DEFAULT '';
            ALTER TABLE runs ADD COLUMN mto_scope TEXT NOT NULL DEFAULT '';
            ALTER TABLE runs ADD COLUMN mto_rows INTEGER NOT NULL DEFAULT 0;
            ALTER TABLE runs ADD COLUMN partial_rows INTEGER NOT NULL DEFAULT 0;
            ALTER TABLE isos ADD COLUMN partial_rows INTEGER NOT NULL DEFAULT 0;
            CREATE INDEX IF NOT EXISTS runs_cwa_idx ON runs(cwa, generated_at DESC);
            CREATE INDEX IF NOT EXISTS runs_mto_scope_idx ON runs(mto_scope, generated_at DESC);
            PRAGMA user_version = 3;
            """
        )
        connection.commit()


def reset_database(connection: sqlite3.Connection) -> None:
    with connection:
        connection.execute("DELETE FROM issues")
        connection.execute("DELETE FROM artifacts")
        connection.execute("DELETE FROM isos")
        connection.execute("DELETE FROM runs")


def _text(value) -> str:
    return "" if value is None else str(value)


def _int(value) -> int:
    try:
        return int(value or 0)
    except (TypeError, ValueError):
        return 0


def _bool(value) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    return _text(value).strip().lower() in {"1", "true", "yes"}


def _find_fmr_workbook(
    metrics: Mapping[str, object],
    artifacts: Sequence[Mapping[str, object]],
    output_folder: str,
) -> Optional[Path]:
    candidates: List[Path] = []
    workbook_value = _text(metrics.get("workbook"))
    if workbook_value:
        candidates.append(Path(workbook_value))
    for artifact in artifacts:
        path_value = _text(artifact.get("path"))
        if path_value and (
            _text(artifact.get("kind")) == "fmr_workbook"
            or path_value.upper().endswith("_FMR.XLSX")
        ):
            candidates.append(Path(path_value))
    folder = Path(output_folder) if output_folder else None
    if folder and folder.is_dir():
        candidates.extend(sorted(folder.glob("*_FMR.xlsx")))
    return next((path.resolve() for path in candidates if path.is_file()), None)


def _enrich_timing_evidence(
    workflow: str,
    status: str,
    metrics: Mapping[str, object],
    isos: Sequence[Mapping[str, object]],
    artifacts: Sequence[Mapping[str, object]],
    output_folder: str,
) -> Tuple[Dict[str, object], List[Dict[str, object]]]:
    enriched_metrics = dict(metrics)
    enriched_isos = [dict(iso) for iso in isos]
    if "time_estimate_available" in enriched_metrics:
        return enriched_metrics, enriched_isos

    if workflow == "fmr":
        evidence_complete = bool(enriched_isos) and all(
            "overflow_material_rows" in iso for iso in enriched_isos
        )
        if not evidence_complete:
            workbook = _find_fmr_workbook(enriched_metrics, artifacts, output_folder)
            parsed = _fmr_workbook_isos(workbook) if workbook else []
            if parsed:
                parsed_by_identity = {
                    (_text(iso.get("drawing_number")), _text(iso.get("revision"))): iso
                    for iso in parsed
                }
                if enriched_isos:
                    for iso in enriched_isos:
                        match = parsed_by_identity.get((
                            _text(iso.get("drawing_number")), _text(iso.get("revision"))
                        ))
                        if match:
                            iso.update({
                                "material_rows": match["material_rows"],
                                "overflow_material_rows": match["overflow_material_rows"],
                                "estimated_time_saved_seconds": match[
                                    "estimated_time_saved_seconds"
                                ],
                            })
                else:
                    enriched_isos = parsed
                evidence_complete = bool(enriched_isos) and all(
                    "overflow_material_rows" in iso for iso in enriched_isos
                )

        selected = len(enriched_isos)
        material_rows = sum(_int(iso.get("material_rows")) for iso in enriched_isos)
        overflow_rows = sum(
            _int(iso.get("overflow_material_rows")) for iso in enriched_isos
        )
        if evidence_complete:
            estimate = fmr_time_estimate(
                selected,
                material_rows,
                overflow_rows,
                eligible=status == "complete",
            )
        else:
            estimate = unavailable_estimate(
                "missing_timing_evidence",
                selected_iso_count=selected,
                material_rows=material_rows,
                overflow_material_rows=overflow_rows,
            )
        enriched_metrics.update(estimate)
    elif workflow == "pipe_footage":
        estimate = pipe_time_estimate(
            len(enriched_isos), eligible=status == "complete" and bool(enriched_isos)
        )
        enriched_metrics.update(estimate)
        if _bool(estimate["time_estimate_available"]):
            for iso in enriched_isos:
                iso.setdefault("estimated_time_saved_seconds", pipe_iso_time_saved())
                iso.setdefault("overflow_material_rows", 0)
    else:
        estimate = unavailable_estimate("workflow_not_measured")
        estimate["time_model_version"] = ""
        enriched_metrics.update(estimate)
    return enriched_metrics, enriched_isos


def _fingerprint(manifest: Mapping[str, object], manifest_path: Path) -> str:
    stable = "|".join((
        _text(manifest.get("run_id")),
        _text(manifest.get("workflow")),
        _text(manifest.get("generated_at")),
        _text(manifest.get("output_folder")),
        str(manifest_path.resolve()),
    ))
    return hashlib.sha256(stable.encode("utf-8")).hexdigest()


def ingest_manifest(
    connection: sqlite3.Connection,
    manifest: Mapping[str, object],
    manifest_path: Path,
    *,
    imported_legacy: bool = False,
) -> str:
    workflow = _text(manifest.get("workflow"))
    if workflow not in WORKFLOW_LABELS:
        raise ValueError(f"Unsupported analytics workflow: {workflow!r}")
    generated_at = _text(manifest.get("generated_at"))
    if not generated_at:
        raise ValueError("Analytics manifest is missing generated_at")
    metrics = dict(manifest.get("metrics") or {})
    isos = list(manifest.get("isos") or [])
    issues = list(manifest.get("issues") or [])
    artifacts = list(manifest.get("artifacts") or [])
    status = _text(manifest.get("status")) or "complete"
    metrics, isos = _enrich_timing_evidence(
        workflow,
        status,
        metrics,
        artifacts=artifacts,
        isos=isos,
        output_folder=_text(manifest.get("output_folder")),
    )
    mto_scope = _text(manifest.get("mto_scope") or metrics.get("mto_scope"))
    if workflow == "mto" and not mto_scope:
        mto_scope = "combined"
    run_id = _text(manifest.get("run_id")) or str(uuid.uuid4())
    fingerprint = _fingerprint(manifest, manifest_path)
    existing = connection.execute(
        "SELECT run_id FROM runs WHERE fingerprint = ?", (fingerprint,)
    ).fetchone()
    if existing:
        run_id = str(existing["run_id"])

    with connection:
        connection.execute(
            """
            INSERT INTO runs (
                run_id, fingerprint, workflow, generated_at, status, package_number,
                input_folder, output_folder, package_total_lf, fmr_sheets,
                accepted_records, iso_count, review_count, summary_json,
                manifest_path, imported_legacy, indexed_at, processing_seconds,
                estimated_time_saved_seconds, time_model_version,
                time_estimate_available, time_saved_breakdown_json, cwa,
                mto_scope, mto_rows, partial_rows
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(fingerprint) DO UPDATE SET
                status=excluded.status,
                package_number=excluded.package_number,
                input_folder=excluded.input_folder,
                output_folder=excluded.output_folder,
                package_total_lf=excluded.package_total_lf,
                fmr_sheets=excluded.fmr_sheets,
                accepted_records=excluded.accepted_records,
                iso_count=excluded.iso_count,
                review_count=excluded.review_count,
                summary_json=excluded.summary_json,
                manifest_path=excluded.manifest_path,
                imported_legacy=excluded.imported_legacy,
                indexed_at=excluded.indexed_at,
                processing_seconds=excluded.processing_seconds,
                estimated_time_saved_seconds=excluded.estimated_time_saved_seconds,
                time_model_version=excluded.time_model_version,
                time_estimate_available=excluded.time_estimate_available,
                time_saved_breakdown_json=excluded.time_saved_breakdown_json,
                cwa=excluded.cwa,
                mto_scope=excluded.mto_scope,
                mto_rows=excluded.mto_rows,
                partial_rows=excluded.partial_rows
            """,
            (
                run_id, fingerprint, workflow, generated_at,
                status,
                _text(manifest.get("package_number")),
                _text(manifest.get("input_folder")),
                _text(manifest.get("output_folder")),
                _text(metrics.get("package_total_linear_feet")) or None,
                _int(metrics.get("fmr_sheets_created")),
                _int(metrics.get("accepted_records")),
                len(isos), len(issues), json.dumps(metrics, sort_keys=True),
                str(manifest_path.resolve()), int(imported_legacy),
                datetime.now(timezone.utc).isoformat(),
                _text(metrics.get("processing_seconds")) or None,
                _text(metrics.get("estimated_time_saved_seconds")) or None,
                _text(metrics.get("time_model_version")),
                int(_bool(metrics.get("time_estimate_available"))),
                json.dumps(metrics.get("time_saved_breakdown") or {}, sort_keys=True),
                _text(manifest.get("cwa") or metrics.get("cwa")),
                mto_scope,
                _int(metrics.get("mto_rows") or metrics.get("mto_rows_created")),
                _int(metrics.get("partial_rows")),
            ),
        )
        connection.execute("DELETE FROM isos WHERE run_id = ?", (run_id,))
        connection.execute("DELETE FROM artifacts WHERE run_id = ?", (run_id,))
        connection.execute("DELETE FROM issues WHERE run_id = ?", (run_id,))
        for iso in isos:
            drawing = _text(iso.get("drawing_number"))
            if not drawing:
                continue
            connection.execute(
                """
                INSERT OR REPLACE INTO isos (
                    run_id, drawing_number, revision, status, linear_feet,
                    pipe_rows, material_rows, source_pdf, source_path, page,
                    overflow_material_rows, estimated_time_saved_seconds,
                    partial_rows
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    run_id, drawing, _text(iso.get("revision")),
                    _text(iso.get("status")) or "accepted",
                    _text(iso.get("linear_feet")) or None,
                    _int(iso.get("pipe_rows")), _int(iso.get("material_rows")),
                    _text(iso.get("source_pdf")), _text(iso.get("source_path")),
                    _int(iso.get("page")),
                    _int(iso.get("overflow_material_rows")),
                    _text(iso.get("estimated_time_saved_seconds")) or None,
                    _int(iso.get("partial_rows")),
                ),
            )
        for artifact in artifacts:
            path = Path(_text(artifact.get("path"))).expanduser()
            exists = path.is_file()
            stat = path.stat() if exists else None
            connection.execute(
                """
                INSERT OR REPLACE INTO artifacts (
                    run_id, kind, name, path, exists_flag, size, modified_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    run_id, _text(artifact.get("kind")) or "artifact",
                    _text(artifact.get("name")) or path.name, str(path.resolve()),
                    int(exists), stat.st_size if stat else None,
                    datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat()
                    if stat else None,
                ),
            )
        for issue in issues:
            connection.execute(
                """
                INSERT INTO issues (
                    run_id, source_pdf, source_path, page, drawing_number,
                    revision, reason_code, reason_detail
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    run_id, _text(issue.get("source_pdf")),
                    _text(issue.get("source_path")), _int(issue.get("page")),
                    _text(issue.get("isometric_drawing_number") or issue.get("drawing_number")),
                    _text(issue.get("revision")),
                    _text(issue.get("reason_code") or issue.get("review_type")),
                    _text(issue.get("reason_detail") or issue.get("review_reasons")),
                ),
            )
    return run_id


def _read_csv(path: Path) -> List[Dict[str, str]]:
    if not path.is_file():
        return []
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def _legacy_pipe_isos(summary: Mapping[str, object], folder: Path) -> List[Dict[str, object]]:
    report_value = _text(summary.get("report"))
    candidates = [Path(report_value)] if report_value else list(folder.glob("*_PIPE_FOOTAGE.docx"))
    report = next((path for path in candidates if path.is_file()), None)
    if not report:
        return []
    result = []
    for table in Document(report).tables:
        for row in table.rows[1:]:
            values = [cell.text.strip() for cell in row.cells]
            if not values or values[0].upper() == "PACKAGE TOTAL" or len(values) < 3:
                continue
            match = re.search(r"-?\d+(?:\.\d+)?", values[2])
            result.append({
                "drawing_number": values[0], "revision": values[1],
                "status": "review" if "REVIEW" in values[2].upper() else "accepted",
                "linear_feet": match.group(0) if match else "",
                "pipe_rows": 0, "material_rows": 0,
                "overflow_material_rows": 0,
                "estimated_time_saved_seconds": pipe_iso_time_saved(),
                "source_pdf": "", "source_path": "", "page": 0,
            })
    return result


def _fmr_workbook_isos(workbook_path: Optional[Path]) -> List[Dict[str, object]]:
    if not workbook_path:
        return []
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    result = []
    try:
        for sheet in workbook.worksheets:
            drawing = _text(sheet["H6"].value).replace("LINE NO:", "").strip()
            revision = _text(sheet["K6"].value).replace("REV:", "").strip()
            if drawing:
                material_rows = 0
                overflow_rows = 0
                for row_number in range(8, 31):
                    values = [sheet[f"{column}{row_number}"].value for column in "BCDE"]
                    if any(value not in (None, "") for value in values):
                        material_rows += 1
                        overflow_rows += int(is_overflow_description(values[-1]))
                result.append({
                    "drawing_number": drawing, "revision": revision,
                    "status": "accepted", "linear_feet": "",
                    "pipe_rows": 0, "material_rows": material_rows,
                    "overflow_material_rows": overflow_rows,
                    "estimated_time_saved_seconds": fmr_iso_time_saved(
                        material_rows, overflow_rows
                    ),
                    "source_pdf": "", "source_path": "", "page": 0,
                })
    finally:
        workbook.close()
    return result


def _legacy_fmr_isos(summary: Mapping[str, object], folder: Path) -> List[Dict[str, object]]:
    workbook_value = _text(summary.get("workbook"))
    candidates = [Path(workbook_value)] if workbook_value else list(folder.glob("*_FMR.xlsx"))
    workbook_path = next((path for path in candidates if path.is_file()), None)
    return _fmr_workbook_isos(workbook_path)


def _legacy_takeoff_isos(folder: Path) -> List[Dict[str, object]]:
    grouped: Dict[str, List[Dict[str, str]]] = defaultdict(list)
    for row in _read_csv(folder / "combined_takeoff.csv"):
        grouped[row.get("source_path") or row.get("source_pdf") or ""].append(row)
    result = []
    for source_path, rows in grouped.items():
        source_pdf = rows[0].get("source_pdf", "")
        revision_match = re.search(r"(?:^|[_-])R([A-Za-z0-9]+)(?:\.[^.]+)?$", source_pdf)
        result.append({
            "drawing_number": Path(source_pdf).stem,
            "revision": revision_match.group(1) if revision_match else "",
            "status": "accepted", "linear_feet": "", "pipe_rows": 0,
            "material_rows": len(rows), "source_pdf": source_pdf,
            "source_path": source_path, "page": _int(rows[0].get("page")),
        })
    return result


def _legacy_mto_isos(
    summary: Mapping[str, object],
    folder: Path,
) -> List[Dict[str, object]]:
    workbook_value = _text(summary.get("workbook"))
    candidates = [Path(workbook_value)] if workbook_value else sorted(folder.glob("MTO *.xlsx"))
    workbook_path = next((path for path in candidates if path.is_file()), None)
    if not workbook_path:
        return []

    scope = _text(summary.get("mto_scope")) or (
        "bolts-gaskets" if "BOLTS & GASKETS" in workbook_path.name.upper() else "combined"
    )
    sheet_name = "BOLTS & GASKETS" if scope == "bolts-gaskets" else "COMBINED"
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    grouped: Dict[str, Dict[str, object]] = {}
    try:
        if sheet_name not in workbook.sheetnames:
            return []
        worksheet = workbook[sheet_name]
        for row in worksheet.iter_rows(min_row=2, max_col=11, values_only=True):
            if not any(value not in (None, "") for value in row):
                continue
            drawing = _text(row[2]).strip()
            if not drawing:
                continue
            if sheet_name == "COMBINED":
                material_values = (row[5], row[6], row[8], row[9], row[10])
            else:
                material_values = (row[5], row[6], row[7], row[8], row[9])
            partial = any(value in (None, "") for value in (row[3], row[4], *material_values[:-1]))
            entry = grouped.setdefault(drawing, {
                "drawing_number": drawing,
                "revision": "",
                "status": "accepted",
                "linear_feet": "",
                "pipe_rows": 0,
                "material_rows": 0,
                "partial_rows": 0,
                "source_pdf": "",
                "source_path": "",
                "page": 0,
            })
            entry["material_rows"] = _int(entry["material_rows"]) + 1
            entry["partial_rows"] = _int(entry["partial_rows"]) + int(partial)
            if partial:
                entry["status"] = "partial"
    finally:
        workbook.close()
    return list(grouped.values())


def _legacy_artifacts(folder: Path) -> List[Dict[str, object]]:
    result = []
    for path in sorted(folder.iterdir()) if folder.is_dir() else []:
        if not path.is_file() or path.name.startswith(".") or path.suffix == ".sqlite3":
            continue
        result.append({"kind": "legacy_artifact", "name": path.name, "path": str(path.resolve())})
    return result


def legacy_manifest(summary_path: Path) -> Dict[str, object]:
    summary = json.loads(summary_path.read_text(encoding="utf-8"))
    folder = summary_path.parent.resolve()
    if summary_path.name == "pipe_footage_run_summary.json":
        workflow = "pipe_footage"
        isos = _legacy_pipe_isos(summary, folder)
        issues = _read_csv(folder / "pipe_footage_quarantine.csv")
    elif summary_path.name == "fmr_run_summary.json":
        workflow = "fmr"
        isos = _legacy_fmr_isos(summary, folder)
        issues = _read_csv(folder / "fmr_quarantine.csv")
    elif summary_path.name == "mto_run_summary.json":
        workflow = "mto"
        isos = _legacy_mto_isos(summary, folder)
        issues = _read_csv(folder / "mto_quarantine.csv")
    elif summary_path.name == "run_summary.json":
        workflow = "bolt_gasket"
        isos = _legacy_takeoff_isos(folder)
        issues = _read_csv(folder / "review_required.csv")
    else:
        raise ValueError(f"Unsupported legacy summary: {summary_path.name}")
    generated_at = _text(summary.get("generated_at"))
    identity = f"{workflow}|{generated_at}|{folder}"
    run_id = str(uuid.uuid5(uuid.NAMESPACE_URL, identity))
    status = _text(summary.get("status")) or (
        "review_required" if _int(summary.get("review_entries")) else "complete"
    )
    artifacts = _legacy_artifacts(folder)
    summary, isos = _enrich_timing_evidence(
        workflow,
        status,
        summary,
        isos=isos,
        artifacts=artifacts,
        output_folder=str(folder),
    )
    return {
        "schema_version": "1.2", "run_id": run_id, "workflow": workflow,
        "generated_at": generated_at,
        "status": status,
        "package_number": _text(summary.get("iwp_number")),
        "cwa": _text(summary.get("cwa")),
        "mto_scope": (
            _text(summary.get("mto_scope"))
            or ("combined" if workflow == "mto" else "")
        ),
        "input_folder": _text(summary.get("input_folder")),
        "output_folder": _text(summary.get("output_folder")) or str(folder),
        "metrics": summary, "isos": isos, "issues": issues,
        "artifacts": artifacts,
    }


def index_roots(
    connection: sqlite3.Connection,
    roots: Sequence[Path],
    *,
    reindex: bool = False,
) -> Dict[str, object]:
    if reindex:
        reset_database(connection)
    indexed = 0
    skipped = 0
    errors = []
    manifest_output_folders = set()
    manifests: List[Path] = []
    for root in roots:
        root = root.resolve()
        if not root.is_dir():
            errors.append(f"Output root does not exist: {root}")
            continue
        manifests.extend(root.rglob("analytics_runs/*.json"))
    for path in sorted(set(manifests)):
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
            ingest_manifest(connection, payload, path)
            manifest_output_folders.add(Path(_text(payload.get("output_folder"))).resolve())
            indexed += 1
        except Exception as exc:
            errors.append(f"{path}: {type(exc).__name__}: {exc}")
    summary_names = {
        "pipe_footage_run_summary.json",
        "fmr_run_summary.json",
        "mto_run_summary.json",
        "run_summary.json",
    }
    summaries = []
    for root in roots:
        if root.resolve().is_dir():
            summaries.extend(path for path in root.resolve().rglob("*.json") if path.name in summary_names)
    for path in sorted(set(summaries)):
        if path.parent.resolve() in manifest_output_folders:
            skipped += 1
            continue
        try:
            payload = legacy_manifest(path)
            ingest_manifest(connection, payload, path, imported_legacy=True)
            indexed += 1
        except Exception as exc:
            errors.append(f"{path}: {type(exc).__name__}: {exc}")
    return {"indexed": indexed, "skipped": skipped, "errors": errors}


def _decimal(value) -> Decimal:
    try:
        return Decimal(_text(value) or "0")
    except InvalidOperation:
        return Decimal("0")


def _filtered_runs(connection: sqlite3.Connection, filters: Mapping[str, str]) -> List[Dict[str, object]]:
    clauses = ["1=1"]
    params: List[object] = []
    if filters.get("workflow"):
        clauses.append("r.workflow = ?")
        params.append(filters["workflow"])
    if filters.get("status"):
        clauses.append("r.status = ?")
        params.append(filters["status"])
    if filters.get("cwa"):
        clauses.append("r.cwa = ?")
        params.append(filters["cwa"])
    if filters.get("mto_scope"):
        clauses.append("r.mto_scope = ?")
        params.append(filters["mto_scope"])
    if filters.get("date_from"):
        clauses.append("substr(r.generated_at, 1, 10) >= ?")
        params.append(filters["date_from"])
    if filters.get("date_to"):
        clauses.append("substr(r.generated_at, 1, 10) <= ?")
        params.append(filters["date_to"])
    query = filters.get("q", "").strip()
    if query:
        like = f"%{query}%"
        clauses.append(
            "(r.package_number LIKE ? OR r.cwa LIKE ? OR r.mto_scope LIKE ? "
            "OR r.input_folder LIKE ? OR r.output_folder LIKE ? "
            "OR EXISTS (SELECT 1 FROM isos i WHERE i.run_id=r.run_id AND i.drawing_number LIKE ?))"
        )
        params.extend((like, like, like, like, like, like))
    rows = connection.execute(
        f"SELECT r.* FROM runs r WHERE {' AND '.join(clauses)} ORDER BY r.generated_at DESC",
        params,
    ).fetchall()
    return [dict(row) for row in rows]


def _current_run_ids(runs: Sequence[Mapping[str, object]]) -> set:
    grouped: Dict[Tuple[str, str, str], List[Mapping[str, object]]] = defaultdict(list)
    for run in runs:
        package_key = _text(run.get("package_number")) or _text(run.get("output_folder"))
        workflow = _text(run.get("workflow"))
        scope = _text(run.get("mto_scope")) if workflow == "mto" else ""
        grouped[(package_key, workflow, scope)].append(run)
    current = set()
    for values in grouped.values():
        successful = [run for run in values if _text(run.get("status")) in SUCCESS_STATUSES]
        chosen = successful[0] if successful else values[0]
        current.add(_text(chosen.get("run_id")))
    return current


def dashboard_data(connection: sqlite3.Connection, filters: Mapping[str, str]) -> Dict[str, object]:
    runs = _filtered_runs(connection, filters)
    current_ids = _current_run_ids(runs)
    run_ids = [_text(run["run_id"]) for run in runs]
    iso_rows: Dict[str, List[Dict[str, object]]] = defaultdict(list)
    artifact_rows: Dict[str, List[Dict[str, object]]] = defaultdict(list)
    issue_counts: Dict[str, int] = defaultdict(int)
    if run_ids:
        placeholders = ",".join("?" for _ in run_ids)
        for row in connection.execute(
            f"SELECT * FROM isos WHERE run_id IN ({placeholders})", run_ids
        ):
            iso_rows[str(row["run_id"])].append(dict(row))
        for row in connection.execute(
            f"SELECT run_id, COUNT(*) AS count FROM issues WHERE run_id IN ({placeholders}) GROUP BY run_id",
            run_ids,
        ):
            issue_counts[str(row["run_id"])] = int(row["count"])
        for row in connection.execute(
            f"SELECT * FROM artifacts WHERE run_id IN ({placeholders}) ORDER BY kind, name",
            run_ids,
        ):
            artifact_rows[str(row["run_id"])].append(dict(row))

    package_groups: Dict[str, Dict[str, object]] = {}
    status_counts: Dict[str, int] = defaultdict(int)
    trend: Dict[str, int] = defaultdict(int)
    pipe_by_package: Dict[str, Decimal] = defaultdict(lambda: Decimal("0"))
    time_by_package: Dict[str, Decimal] = defaultdict(lambda: Decimal("0"))
    current_isos = set()
    total_pipe = Decimal("0")
    total_time_saved = Decimal("0")
    total_processing = Decimal("0")
    time_estimate_runs = 0
    processing_runs = 0
    fmr_sheets = 0
    accepted_records = 0
    mto_workbooks = 0
    mto_rows = 0
    partial_rows = 0
    outstanding_reviews = 0
    for run in runs:
        run_id = _text(run["run_id"])
        package = _text(run["package_number"])
        display_package = package or "Standalone ISO Takeoffs"
        status_counts[_text(run["status"])] += 1
        trend[_text(run["generated_at"])[:10]] += 1
        group = package_groups.setdefault(display_package, {
            "package_number": package, "display_package": display_package,
            "latest_status": run["status"], "last_run": run["generated_at"],
            "cwas": set(), "mto_scopes": set(),
            "workflows": set(), "iso_numbers": set(), "pipe_footage": Decimal("0"),
            "fmr_count": 0, "issue_count": 0, "run_count": 0,
            "mto_count": 0, "mto_rows": 0, "partial_rows": 0,
            "estimated_time_saved_seconds": Decimal("0"),
            "time_estimate_count": 0, "processing_seconds": Decimal("0"),
            "processing_run_count": 0,
        })
        group["workflows"].add(_text(run["workflow"]))
        group["run_count"] += 1
        if run_id not in current_ids:
            continue
        if _text(run.get("cwa")):
            group["cwas"].add(_text(run.get("cwa")))
        for iso in iso_rows[run_id]:
            drawing = _text(iso["drawing_number"])
            if drawing:
                group["iso_numbers"].add(drawing)
                current_isos.add((package, drawing))
        group["issue_count"] += issue_counts[run_id]
        outstanding_reviews += issue_counts[run_id]
        if _bool(run.get("time_estimate_available")):
            saved = _decimal(run.get("estimated_time_saved_seconds"))
            group["estimated_time_saved_seconds"] += saved
            group["time_estimate_count"] += 1
            total_time_saved += saved
            time_estimate_runs += 1
            time_by_package[display_package] += saved
        if run.get("processing_seconds") not in (None, ""):
            duration = _decimal(run.get("processing_seconds"))
            group["processing_seconds"] += duration
            group["processing_run_count"] += 1
            total_processing += duration
            processing_runs += 1
        if run["workflow"] == "pipe_footage":
            amount = _decimal(run["package_total_lf"])
            group["pipe_footage"] += amount
            pipe_by_package[display_package] += amount
            total_pipe += amount
        elif run["workflow"] == "fmr":
            group["fmr_count"] += _int(run["fmr_sheets"])
            fmr_sheets += _int(run["fmr_sheets"])
        elif run["workflow"] == "bolt_gasket":
            accepted_records += _int(run["accepted_records"])
        elif run["workflow"] == "mto":
            scope = _text(run.get("mto_scope")) or "combined"
            group["mto_scopes"].add(scope)
            group["mto_count"] += 1
            group["mto_rows"] += _int(run.get("mto_rows"))
            group["partial_rows"] += _int(run.get("partial_rows"))
            mto_workbooks += int(any(
                artifact["kind"] == "mto_workbook" and artifact["exists_flag"]
                for artifact in artifact_rows[run_id]
            ))
            mto_rows += _int(run.get("mto_rows"))
            partial_rows += _int(run.get("partial_rows"))

    packages = []
    for group in package_groups.values():
        iso_numbers = group.pop("iso_numbers")
        cwas = group.pop("cwas")
        mto_scopes = group.pop("mto_scopes")
        packages.append({
            **group,
            "cwas": sorted(cwas),
            "mto_scopes": sorted(mto_scopes),
            "workflows": sorted(group["workflows"]),
            "iso_count": len(iso_numbers),
            "pipe_footage": str(group["pipe_footage"]),
            "estimated_time_saved_seconds": str(group["estimated_time_saved_seconds"]),
            "time_estimate_available": bool(group["time_estimate_count"]),
            "processing_seconds": str(group["processing_seconds"]),
            "processing_available": bool(group["processing_run_count"]),
        })
    packages.sort(key=lambda item: _text(item["last_run"]), reverse=True)
    for run in runs:
        run["is_current"] = _text(run["run_id"]) in current_ids
        run["workflow_label"] = WORKFLOW_LABELS.get(_text(run["workflow"]), _text(run["workflow"]))
        run["issue_count"] = issue_counts[_text(run["run_id"])]
        try:
            run["time_saved_breakdown"] = json.loads(
                _text(run.get("time_saved_breakdown_json")) or "{}"
            )
        except json.JSONDecodeError:
            run["time_saved_breakdown"] = {}

    mto_register = []
    for run in runs:
        run_id = _text(run["run_id"])
        if run["workflow"] != "mto" or run_id not in current_ids:
            continue
        workbook = next(
            (
                artifact for artifact in artifact_rows[run_id]
                if artifact["kind"] == "mto_workbook"
                or _text(artifact["name"]).lower().endswith(".xlsx")
            ),
            None,
        )
        mto_register.append({
            "run_id": run_id,
            "cwa": _text(run.get("cwa")),
            "package_number": _text(run.get("package_number")),
            "mto_scope": _text(run.get("mto_scope")) or "combined",
            "selected_iso_count": len(iso_rows[run_id]),
            "mto_rows": _int(run.get("mto_rows")),
            "partial_rows": _int(run.get("partial_rows")),
            "generated_at": _text(run.get("generated_at")),
            "status": _text(run.get("status")),
            "workbook": workbook,
        })

    return {
        "overview": {
            "packages": len({run["package_number"] for run in runs if run["package_number"]}),
            "runs": len(runs), "isos": len(current_isos),
            "pipe_footage": str(total_pipe), "fmr_sheets": fmr_sheets,
            "accepted_records": accepted_records, "outstanding_reviews": outstanding_reviews,
            "cwas": len({run["cwa"] for run in runs if run.get("cwa")}),
            "mto_workbooks": mto_workbooks,
            "mto_rows": mto_rows,
            "partial_rows": partial_rows,
            "estimated_time_saved_seconds": str(total_time_saved),
            "time_estimate_available": bool(time_estimate_runs),
            "time_estimate_runs": time_estimate_runs,
            "processing_seconds": str(total_processing),
            "processing_available": bool(processing_runs),
            "processing_runs": processing_runs,
            "time_model_version": TIME_MODEL_VERSION,
        },
        "packages": packages,
        "mto_register": mto_register,
        "runs": runs[:100],
        "charts": {
            "status": [{"label": key, "value": value} for key, value in sorted(status_counts.items())],
            "trend": [{"label": key, "value": value} for key, value in sorted(trend.items())],
            "pipe_by_package": [
                {"label": key, "value": str(value)}
                for key, value in sorted(pipe_by_package.items(), key=lambda item: item[1], reverse=True)
            ],
            "time_saved_by_package": [
                {"label": key, "value": str(value)}
                for key, value in sorted(time_by_package.items(), key=lambda item: item[1], reverse=True)
            ],
        },
        "filters": {
            "workflows": [{"value": key, "label": value} for key, value in WORKFLOW_LABELS.items()],
            "statuses": [row[0] for row in connection.execute(
                "SELECT DISTINCT status FROM runs WHERE status <> '' ORDER BY status"
            )],
            "cwas": [row[0] for row in connection.execute(
                "SELECT DISTINCT cwa FROM runs WHERE cwa <> '' ORDER BY cwa"
            )],
            "mto_scopes": [row[0] for row in connection.execute(
                "SELECT DISTINCT mto_scope FROM runs WHERE mto_scope <> '' ORDER BY mto_scope"
            )],
        },
    }


def package_detail(connection: sqlite3.Connection, package_number: str) -> Dict[str, object]:
    if package_number == "__standalone__":
        rows = connection.execute(
            "SELECT * FROM runs WHERE package_number = '' ORDER BY generated_at DESC"
        ).fetchall()
        display = "Standalone ISO Takeoffs"
    else:
        rows = connection.execute(
            "SELECT * FROM runs WHERE package_number = ? ORDER BY generated_at DESC",
            (package_number,),
        ).fetchall()
        display = package_number
    runs = [dict(row) for row in rows]
    current = _current_run_ids(runs)
    for run in runs:
        run_id = _text(run["run_id"])
        run["is_current"] = run_id in current
        run["workflow_label"] = WORKFLOW_LABELS.get(_text(run["workflow"]), _text(run["workflow"]))
        try:
            run["time_saved_breakdown"] = json.loads(
                _text(run.get("time_saved_breakdown_json")) or "{}"
            )
        except json.JSONDecodeError:
            run["time_saved_breakdown"] = {}
        run["isos"] = [dict(row) for row in connection.execute(
            "SELECT * FROM isos WHERE run_id = ? ORDER BY drawing_number, revision", (run_id,)
        )]
        run["artifacts"] = [dict(row) for row in connection.execute(
            "SELECT * FROM artifacts WHERE run_id = ? ORDER BY kind, name", (run_id,)
        )]
        run["issues"] = [dict(row) for row in connection.execute(
            "SELECT * FROM issues WHERE run_id = ? ORDER BY page, reason_code", (run_id,)
        )]
    return {"package_number": package_number, "display_package": display, "runs": runs}


def artifact_record(connection: sqlite3.Connection, artifact_id: int) -> Optional[Dict[str, object]]:
    row = connection.execute("SELECT * FROM artifacts WHERE id = ?", (artifact_id,)).fetchone()
    return dict(row) if row else None
