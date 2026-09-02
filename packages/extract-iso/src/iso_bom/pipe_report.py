import os
import tempfile
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Optional, Sequence

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .fmr_model import QuarantineEntry
from .pipe_model import PipeIsoPage


TURNER_GREEN = "006B54"
DARK_GREEN = "004B3B"
MUTED = "5B6570"
LIGHT_GREEN = "E8F3EF"
LIGHT_GRAY = "F2F4F7"
GRID = "C9D2D9"
WARNING_FILL = "FFF3CD"
WARNING_BORDER = "B7791F"
WHITE = "FFFFFF"
TABLE_WIDTH_DXA = 9360
MAIN_COLUMN_WIDTHS = (6480, 960, 1920)
REVIEW_COLUMN_WIDTHS = (2700, 720, 5940)


def format_linear_feet(value: Decimal, precision: int) -> str:
    return f"{value:.{max(1, precision)}f} LF"


def _humanize_reason(value: str) -> str:
    return value.replace("_", " ").replace(";", "; ").strip()


def _set_run_font(
    run,
    *,
    name: str = "Calibri",
    size: Optional[float] = None,
    color: Optional[str] = None,
    bold: Optional[bool] = None,
    italic: Optional[bool] = None,
) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def _set_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    title = document.styles["Title"]
    title.font.name = "Calibri"
    title._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    title._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    title.font.size = Pt(23)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(TURNER_GREEN)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)
    title_properties = title._element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)

    subtitle = document.styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    subtitle._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    subtitle.font.size = Pt(12)
    subtitle.font.italic = False
    subtitle.font.color.rgb = RGBColor.from_string(MUTED)
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(14)
    subtitle_properties = subtitle._element.get_or_add_pPr()
    subtitle_numbering = subtitle_properties.find(qn("w:numPr"))
    if subtitle_numbering is not None:
        subtitle_properties.remove(subtitle_numbering)

    heading = document.styles["Heading 1"]
    heading.font.name = "Calibri"
    heading._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    heading._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    heading.font.size = Pt(16)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor.from_string(TURNER_GREEN)
    heading.paragraph_format.space_before = Pt(16)
    heading.paragraph_format.space_after = Pt(8)


def _add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    _set_run_font(run, size=9, color=MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    field_run = OxmlElement("w:r")
    run_props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), MUTED)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    run_props.extend((color, size))
    field_run.append(run_props)
    field_run.append(OxmlElement("w:t"))
    field.append(field_run)
    paragraph._p.append(field)


def _set_header_footer(document: Document, iwp_number: str) -> None:
    section = document.sections[0]
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    left = paragraph.add_run("TURNER INDUSTRIES")
    _set_run_font(left, size=9, color=TURNER_GREEN, bold=True)
    right = paragraph.add_run(f"    PIPE FOOTAGE  |  {iwp_number}")
    _set_run_font(right, size=9, color=MUTED)

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.paragraph_format.space_before = Pt(0)
    _add_page_number(footer_paragraph)


def _shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_geometry(table, widths: Sequence[int]) -> None:
    if sum(widths) != TABLE_WIDTH_DXA:
        raise ValueError("DOCX table widths must total 9360 DXA")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    properties = table._tbl.tblPr

    table_width = properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.append(table_width)
    table_width.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    table_width.set(qn("w:type"), "dxa")

    indent = properties.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        properties.append(indent)
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")

    layout = properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_width = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            cell_width.set(qn("w:w"), str(width))
            cell_width.set(qn("w:type"), "dxa")
            _set_cell_margins(cell)
        cannot_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cannot_split)


def _repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def _set_cell_text(
    cell,
    text: str,
    *,
    bold: bool = False,
    color: str = "000000",
    align=WD_ALIGN_PARAGRAPH.LEFT,
    size: float = 10,
) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.10
    run = paragraph.add_run(text)
    _set_run_font(run, size=size, color=color, bold=bold)


def _set_table_borders(table, color: str = GRID, size: str = "6") -> None:
    properties = table._tbl.tblPr
    borders = properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def _main_table_chunks(
    iso_pages: Sequence[PipeIsoPage],
    first_page_capacity: int = 18,
    continuation_capacity: int = 22,
    minimum_last_page: int = 6,
) -> list[list[PipeIsoPage]]:
    pages = list(iso_pages)
    if not pages:
        return [[]]
    if len(pages) <= first_page_capacity:
        return [pages]

    chunks = [pages[:first_page_capacity]]
    remaining = pages[first_page_capacity:]
    while remaining:
        chunks.append(remaining[:continuation_capacity])
        remaining = remaining[continuation_capacity:]
    if len(chunks) > 1 and len(chunks[-1]) < minimum_last_page:
        needed = minimum_last_page - len(chunks[-1])
        available = max(0, len(chunks[-2]) - minimum_last_page)
        moved = min(needed, available)
        if moved:
            chunks[-1] = chunks[-2][-moved:] + chunks[-1]
            chunks[-2] = chunks[-2][:-moved]
    return chunks


def _add_main_table(
    document: Document,
    iso_pages: Sequence[PipeIsoPage],
    global_start_index: int,
    package_total: Optional[Decimal],
    package_precision: int,
    include_total: bool,
) -> None:
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    _set_table_borders(table)
    header = table.rows[0]
    _repeat_header(header)
    for cell, text, align in zip(
        header.cells,
        ("ISOMETRIC DRAWING", "REV", "LINEAR FEET"),
        (WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT),
    ):
        _shade_cell(cell, TURNER_GREEN)
        _set_cell_text(cell, text, bold=True, color=WHITE, align=align, size=9.5)

    for local_index, iso in enumerate(iso_pages):
        row = table.add_row()
        if (global_start_index + local_index) % 2:
            for cell in row.cells:
                _shade_cell(cell, LIGHT_GRAY)
        _set_cell_text(row.cells[0], iso.drawing_number, size=10)
        _set_cell_text(row.cells[1], iso.revision, align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
        total = iso.total_linear_feet
        if total is None or iso.review_reasons:
            shown_total = "REVIEW REQUIRED"
            total_color = WARNING_BORDER
        else:
            shown_total = format_linear_feet(total, iso.precision)
            total_color = DARK_GREEN
        _set_cell_text(
            row.cells[2], shown_total, bold=True, color=total_color,
            align=WD_ALIGN_PARAGRAPH.RIGHT, size=10,
        )

    if not iso_pages:
        row = table.add_row()
        merged = row.cells[0].merge(row.cells[2])
        _set_cell_text(
            merged, "No selectable ISO drawings were available. See Review Items.",
            color=WARNING_BORDER, size=10,
        )

    if include_total and package_total is not None:
        row = table.add_row()
        label_cell = row.cells[0].merge(row.cells[1])
        _shade_cell(label_cell, LIGHT_GREEN)
        _shade_cell(row.cells[2], LIGHT_GREEN)
        _set_cell_text(label_cell, "PACKAGE TOTAL", bold=True, color=DARK_GREEN, size=11)
        _set_cell_text(
            row.cells[2], format_linear_feet(package_total, package_precision),
            bold=True, color=DARK_GREEN, align=WD_ALIGN_PARAGRAPH.RIGHT, size=11,
        )

    _set_table_geometry(table, MAIN_COLUMN_WIDTHS)


def _add_warning(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(12)
    paragraph.paragraph_format.left_indent = Pt(8)
    paragraph.paragraph_format.right_indent = Pt(8)
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), WARNING_FILL)
    properties.append(shading)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "8")
        node.set(qn("w:color"), WARNING_BORDER)
        borders.append(node)
    properties.append(borders)
    lead = paragraph.add_run("REVIEW REQUIRED — PACKAGE TOTAL WITHHELD\n")
    _set_run_font(lead, size=11, color=WARNING_BORDER, bold=True)
    detail = paragraph.add_run(
        "One or more package pages, drawing revisions, BOM structures, or pipe quantities "
        "could not be verified confidently. Resolve every review item before using a package total."
    )
    _set_run_font(detail, size=10, color="4A3A00")


def build_pipe_footage_report(
    output_path: Path,
    iwp_number: str,
    iso_pages: Sequence[PipeIsoPage],
    status: str,
    generated_at: str,
    quarantine: Sequence[QuarantineEntry],
    package_total: Optional[Decimal],
    package_precision: int,
    overwrite: bool = False,
) -> Path:
    output_path = output_path.resolve()
    if output_path.exists() and not overwrite:
        raise FileExistsError(f"Output report already exists: {output_path}")

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    _set_styles(document)
    _set_header_footer(document, iwp_number)

    title = document.add_paragraph(style="Title")
    title.add_run("PIPE LINEAR FOOTAGE REPORT")
    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.add_run(f"Installation Work Package  |  {iwp_number}")

    metadata = document.add_paragraph()
    metadata.paragraph_format.space_after = Pt(12)
    label = metadata.add_run("Generated: ")
    _set_run_font(label, size=9.5, color=MUTED, bold=True)
    try:
        shown_date = datetime.fromisoformat(generated_at.replace("Z", "+00:00")).strftime("%B %d, %Y %H:%M UTC")
    except ValueError:
        shown_date = generated_at
    value = metadata.add_run(shown_date)
    _set_run_font(value, size=9.5, color=MUTED)

    if status == "review_required":
        _add_warning(document)

    chunks = _main_table_chunks(iso_pages)
    global_start_index = 0
    for chunk_index, chunk in enumerate(chunks):
        if chunk_index:
            document.add_page_break()
        _add_main_table(
            document,
            chunk,
            global_start_index,
            package_total,
            package_precision,
            include_total=(
                chunk_index == len(chunks) - 1
                and status == "complete"
                and package_total is not None
            ),
        )
        global_start_index += len(chunk)

    if status == "review_required":
        withheld = document.add_paragraph()
        withheld.paragraph_format.space_before = Pt(10)
        withheld.paragraph_format.space_after = Pt(4)
        run = withheld.add_run("Package Total: WITHHELD")
        _set_run_font(run, size=12, color=WARNING_BORDER, bold=True)

        document.add_paragraph("Review Items", style="Heading 1")
        review_table = document.add_table(rows=1, cols=3)
        review_table.style = "Table Grid"
        _set_table_borders(review_table)
        review_header = review_table.rows[0]
        _repeat_header(review_header)
        for cell, text, align in zip(
            review_header.cells,
            ("SOURCE", "PAGE", "REASON"),
            (WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT),
        ):
            _shade_cell(cell, WARNING_BORDER)
            _set_cell_text(cell, text, bold=True, color=WHITE, align=align, size=9.5)
        for entry in quarantine:
            row = review_table.add_row()
            source = entry.isometric_drawing_number or entry.source_path or entry.source_pdf
            reason = entry.reason_detail or entry.reason_code
            if entry.reason_code and entry.reason_code not in reason:
                reason = f"{entry.reason_code}: {reason}"
            reason = _humanize_reason(reason)
            _set_cell_text(row.cells[0], source, size=9)
            _set_cell_text(row.cells[1], str(entry.page or "—"), align=WD_ALIGN_PARAGRAPH.CENTER, size=9)
            _set_cell_text(row.cells[2], reason, size=9)
        _set_table_geometry(review_table, REVIEW_COLUMN_WIDTHS)

    document.core_properties.title = f"{iwp_number} Pipe Linear Footage Report"
    document.core_properties.subject = "Pipe linear footage by isometric drawing"
    document.core_properties.author = "Turner Industries"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=".pipe-footage-", suffix=".docx", dir=str(output_path.parent)
    )
    os.close(descriptor)
    temporary_path = Path(temporary_name)
    try:
        document.save(temporary_path)
        if output_path.exists() and not overwrite:
            raise FileExistsError(f"Output report already exists: {output_path}")
        os.replace(temporary_path, output_path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()
    return output_path
